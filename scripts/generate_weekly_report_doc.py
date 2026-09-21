from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


def set_font(run, size=None, bold=None):
    run.font.name = "Microsoft JhengHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_bullet(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(12)
    r = p.add_run("• ")
    set_font(r, bold=True)
    r = p.add_run(text)
    set_font(r, 11)


def main():
    out_dir = Path("reports")
    out_dir.mkdir(exist_ok=True)
    out_path = out_dir / "weekly_update_report_2026-07-07_2026-07-10.docx"

    doc = Document()
    for style_name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        style = doc.styles[style_name]
        style.font.name = "Microsoft JhengHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    doc.styles["Normal"].font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("AI 智慧客服系統本週更新報告")
    set_font(r, 18, True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("更新期間：2026/07/07 - 2026/07/10")
    set_font(r, 11)

    doc.add_paragraph("")

    sections = [
        (
            "一、本週更新重點",
            [
                "修正帳務與繳費流程，避免帳單查詢、線上繳費與復線 API 互相誤判。",
                "改善網路、電視與遙控器排錯流程，降低重複詢問與錯誤分流。",
                "強化 RAG 知識庫檢索詞，補上載具歸戶、單品銷售、加值服務等客服常用說法。",
                "修正續約與退租上下文判斷，讓使用者簡短回覆也能接續正確流程。",
                "清理不必要的服務地址要求，僅在申裝、移機、報修等必要情境要求地址。",
            ],
        ),
        (
            "二、帳務與繳費流程修正",
            [
                "修正「詢問帳單內容」被誤判成報修流程的問題。",
                "「帳單內容、帳單明細、本期帳單金額查詢」現在會走帳單查詢，不會出現維修申告。",
                "修正「我要線上繳費，網路費」被誤判成網路復線 API 的問題。",
                "新增線上繳費與線上刷卡判斷，會回答官方網站線上繳費、哈TV行動客服 APP 繳費，以及付款成功後通常會自動開通。",
                "電視/網路復線 API 不再要求服務地址，只要求戶名與聯絡電話，或客戶編號。",
            ],
        ),
        (
            "三、排錯流程改善",
            [
                "修正網路不穩/斷線話術，移除「不是完全不能上網」這類過度判斷。",
                "網路不穩現在會直接詢問影響範圍：所有網站/APP 都不穩，或只有遊戲、影片、特定 APP。",
                "修正電視訊號不良、畫面 lag、播放停頓時重複詢問分類的問題。",
                "修正遙控器控制異常，會切到遙控器檢查流程，確認按鍵紅燈、電池、IR 接收器等。",
            ],
        ),
        (
            "四、RAG / 知識庫檢索強化",
            [
                "新增「單品銷售 = 加值服務」判斷。",
                "「更多熱門單品銷售」現在會走加值服務知識庫，不會只回澄清。",
                "新增載具歸戶檢索強化詞：載具歸戶、發票載具、發票號碼載具、用戶歸戶、財政部。",
                "補強加值服務、熱門單品、加購服務等查詢詞。",
            ],
        ),
        (
            "五、續約 / 退租上下文修正",
            [
                "修正前文在談續約、合約到期時，使用者後續說「就是結束」被當成一般結束對話的問題。",
                "現在會依上下文判斷為不續約或退租流程。",
            ],
        ),
        (
            "六、優惠與申辦意圖修正",
            [
                "修正「申請300M網路優惠」這類句子。",
                "現在會判斷為使用者要辦理，不是再列方案內容。",
                "避免使用者明確要申請時，AI 還持續回答方案介紹。",
            ],
        ),
        (
            "七、地址欄位清理",
            [
                "合約查詢、帳單查詢、復線 API 不再要求服務地址。",
                "清除多處殘留文案，例如「請提供裝機地址」、「請提供服務地址」。",
                "目前只有真正需要地址的功能才會要求，例如申裝、移機、報修。",
            ],
        ),
        (
            "八、測試驗證",
            [
                "本週修正後已執行 Router 架構測試、Flow regression 測試、Customer validation 測試、Tool manager 測試、KB service 測試。",
                "最新核心測試結果：259 passed。",
            ],
        ),
    ]

    for heading, items in sections:
        doc.add_heading(heading, level=1)
        for item in items:
            add_bullet(doc, item)

    doc.add_heading("九、待客服補充 QA 清單", level=1)
    qa_rows = [
        ("繳費方式查詢", "有哪些繳費方式？", "可透過官方網站線上繳費、哈TV行動客服 APP、公司櫃台、便利商店帳單條碼、7-11 ibon / 全家 FamiPort 等方式繳費。"),
        ("線上刷卡是否自動開通", "線上刷卡繳費後會馬上開通嗎？", "付款成功後系統通常會自動開通服務；若仍無法使用，請重啟數據機或機上盒後再確認。"),
        ("載具歸戶設定", "如何設定發票載具歸戶？", "至公司官網的客戶服務發票查詢，輸入用戶帳號密碼，點選用戶歸戶，並連結財政部網站進行歸戶。"),
        ("單品銷售/加值服務", "熱門單品銷售是什麼？有哪些可以加購？", "熱門單品銷售即加值服務，請列出可加購項目、申辦方式、費用與限制條件。"),
        ("不續約/退租流程", "合約到期不續約或想結束服務怎麼辦？", "請說明退租流程、是否需歸還設備、是否有違約金或未結清費用，以及需由客服確認的資料。"),
        ("多台電視/機上盒收費", "單純安裝有線電視，半年繳，3 台電視要收多少？", "需明確列出收視費、裝機費、第 2 台/第 3 台分機費、機上盒押金；第 3 台以上若需押金 1200 元與分機費，也要寫入。"),
    ]
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["項目", "建議問題", "建議答案重點", "狀態"]
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
    for item, q, a in qa_rows:
        row = table.add_row().cells
        row[0].text = item
        row[1].text = q
        row[2].text = a
        row[3].text = "待補充/確認"

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        set_font(run, 10)

    doc.save(out_path)
    print(out_path.resolve())


if __name__ == "__main__":
    main()
