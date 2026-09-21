from __future__ import annotations

import argparse
import html
import json
import re
import time
import uuid
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

import requests


DEFAULT_API_BASE = "http://127.0.0.1:8000"
DEFAULT_TV_CABLE = "tdtv"
DEFAULT_XLSX = "三個月AI對話情境測試案例整理.xlsx"


@dataclass
class Check:
    label: str
    mode: str
    terms: list[str]
    target: str = "final_text"


@dataclass
class Scenario:
    case_id: str
    category: str
    title: str
    source_ref: str
    expected_behavior: str
    turns: list[str]
    checks: list[Check] = field(default_factory=list)


def c(label: str, mode: str, terms: list[str], target: str = "final_text") -> Check:
    return Check(label=label, mode=mode, terms=terms, target=target)


def load_excel_case_map(workbook_path: Path) -> dict[str, dict[str, str]]:
    if not workbook_path.exists():
        return {}

    try:
        from openpyxl import load_workbook
    except Exception:
        return {}

    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    ws = wb["情境案例"] if "情境案例" in wb.sheetnames else wb.worksheets[1]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {}

    headers = [str(value or "") for value in rows[0]]
    case_map: dict[str, dict[str, str]] = {}
    for row in rows[1:]:
        item = {
            headers[index]: "" if value is None else str(value)
            for index, value in enumerate(row)
            if index < len(headers)
        }
        case_id = item.get("案例ID", "")
        if case_id.startswith("TC-"):
            case_map[case_id] = item
    return case_map


def excel_references_for(source_ref: str, case_map: dict[str, dict[str, str]]) -> list[dict[str, str]]:
    refs = []
    for case_id in re.findall(r"TC-\d+", source_ref or ""):
        if case_id in case_map:
            refs.append(case_map[case_id])
    return refs


SCENARIOS = [
    Scenario(
        "API-001",
        "帳務 API",
        "查詢帳單成功",
        "TC-007 / TC-008",
        "提供正確戶名與電話後，應呼叫 mock 帳單 API 並回傳金額與繳費到期日。",
        ["查詢帳單", "王大明 0988555666"],
        [
            c("回覆帳單金額", "contains_all", ["寬頻月租", "799"]),
            c("回覆繳費到期日", "contains_any", ["繳費到期日", "2026-05-31"]),
            c("工具成功", "contains_any", ["True"], "last_tool_success"),
            c("打到 mock 帳單 endpoint", "contains_any", ["getCustBill"], "last_tool_endpoint"),
        ],
    ),
    Scenario(
        "API-002",
        "帳務 API",
        "查詢帳單失敗：電話不一致",
        "TC-007 / TC-008",
        "提供錯誤電話時，應回覆查無資料，不可揭露帳務資訊。",
        ["查詢帳單", "王大明 0988666555"],
        [
            c("回覆查無資料", "contains_any", ["查詢不到", "不一致", "重新輸入"]),
            c("不可揭露帳單金額", "not_contains_any", ["799", "繳費到期日"]),
            c("工具失敗", "contains_any", ["False"], "last_tool_success"),
        ],
    ),
    Scenario(
        "API-003",
        "合約 API",
        "查詢合約成功",
        "TC-011 / TC-017",
        "提供正確身份後，應呼叫 mock 合約 API 並回傳目前方案與合約到期日。",
        ["想請問我的合約到期日", "王大明 0988555666"],
        [
            c("回覆合約資訊", "contains_all", ["目前方案", "合約到期日"]),
            c("回覆 mock 到期日", "contains_any", ["2026-12-31"]),
            c("工具成功", "contains_any", ["True"], "last_tool_success"),
            c("打到 mock 合約 endpoint", "contains_any", ["contractInfo"], "last_tool_endpoint"),
        ],
    ),
    Scenario(
        "API-004",
        "合約 API",
        "查詢合約失敗：電話不一致",
        "TC-011 / TC-017",
        "身份資料不一致時，應要求重新確認，不可回傳合約資訊。",
        ["想請問我的合約到期日", "王大明 0988666555"],
        [
            c("回覆查無資料", "contains_any", ["查詢不到", "不一致", "重新輸入"]),
            c("不可揭露合約日期", "not_contains_any", ["2026-12-31"]),
            c("工具失敗", "contains_any", ["False"], "last_tool_success"),
        ],
    ),
    Scenario(
        "API-005",
        "帳務 API",
        "補發簡訊帳單成功",
        "TC-007 / TC-008",
        "提供正確身份後，應呼叫 mock 簡訊帳單 API，不發送真實簡訊。",
        ["補發簡訊帳單", "王大明 0988555666"],
        [
            c("回覆 mock 簡訊流程", "contains_all", ["模擬簡訊帳單", "尚未發送真實簡訊"]),
            c("工具成功", "contains_any", ["True"], "last_tool_success"),
            c("打到 mock 簡訊 endpoint", "contains_any", ["reBillE"], "last_tool_endpoint"),
        ],
    ),
    Scenario(
        "API-006",
        "復線 API",
        "網路復線成功",
        "TC-003",
        "補齊正確身份後，應呼叫 mock 網路復線 API 並提醒設備重開。",
        ["已經繳費要恢復網路服務", "王大明 0988555666"],
        [
            c("回覆網路復線 mock 成功", "contains_all", ["模擬網路復線", "設備電源關機重開"]),
            c("工具成功", "contains_any", ["True"], "last_tool_success"),
            c("打到 mock 網路復線 endpoint", "contains_any", ["changeReceive"], "last_tool_endpoint"),
        ],
    ),
    Scenario(
        "API-007",
        "復線 API",
        "網路復線失敗：電話不一致",
        "TC-003",
        "身份資料不一致時，應回覆查無資料，不可宣稱已復線。",
        ["已經繳費要恢復網路服務", "王大明 0988666555"],
        [
            c("回覆查無資料", "contains_any", ["查詢不到", "不一致", "重新輸入"]),
            c("不可宣稱復線成功", "not_contains_any", ["已用模擬網路復線 API 送出測試申請"]),
            c("工具失敗", "contains_any", ["False"], "last_tool_success"),
        ],
    ),
    Scenario(
        "API-008",
        "復線 API",
        "電視復線成功",
        "TC-006",
        "補齊正確身份後，應呼叫 mock 電視復線 API 並提醒設備重開。",
        ["已繳費無法收看", "王大明 0988555666"],
        [
            c("回覆電視復線 mock 成功", "contains_all", ["模擬電視復線", "設備電源關機重開"]),
            c("工具成功", "contains_any", ["True"], "last_tool_success"),
            c("打到 mock 電視復線 endpoint", "contains_any", ["dtvChangeReceive"], "last_tool_endpoint"),
        ],
    ),
    Scenario(
        "API-009",
        "復線 API",
        "電視復線失敗：電話不一致",
        "TC-006",
        "身份資料不一致時，應回覆查無資料，不可宣稱已復線。",
        ["已繳費無法收看", "王大明 0988666555"],
        [
            c("回覆查無資料", "contains_any", ["查詢不到", "不一致", "重新輸入"]),
            c("不可宣稱復線成功", "not_contains_any", ["已用模擬電視復線 API 送出測試申請"]),
            c("工具失敗", "contains_any", ["False"], "last_tool_success"),
        ],
    ),
    Scenario(
        "API-010",
        "優惠 API",
        "優惠查詢資料不足",
        "TC-018 / TC-019",
        "缺少服務地區時，應停在 pending tool 並追問地區。",
        ["我要最新優惠"],
        [
            c("追問服務地區", "contains_any", ["服務地區", "大里區", "台中"]),
            c("pending search_promotion", "contains_any", ["search_promotion"], "pending_tool"),
        ],
    ),
    Scenario(
        "API-011",
        "優惠 API",
        "優惠查詢成功",
        "TC-018 / TC-019",
        "補服務地區後，應呼叫 mock 優惠 API 並回傳方案與官方連結。",
        ["我要最新優惠", "台中"],
        [
            c("回覆 mock 優惠 API", "contains_all", ["模擬優惠 API", "活動名稱"]),
            c("工具成功", "contains_any", ["True"], "last_tool_success"),
            c("打到 mock 優惠 endpoint", "contains_any", ["promotions"], "last_tool_endpoint"),
        ],
    ),
    Scenario(
        "API-012",
        "頻道 API",
        "頻道號查詢成功",
        "TC-023",
        "補服務地區後，應呼叫 mock 頻道 API 並回傳愛爾達體育台號。",
        ["愛爾達體育在第幾台", "大屯"],
        [
            c("回覆頻道號", "contains_all", ["愛爾達體育", "168"]),
            c("工具成功", "contains_any", ["True"], "last_tool_success"),
            c("打到 mock 頻道 endpoint", "contains_any", ["channelNo"], "last_tool_endpoint"),
        ],
    ),
    Scenario(
        "API-013",
        "頻道 API",
        "頻道號查詢失敗：查無頻道",
        "TC-023",
        "頻道名稱不在 mock 資料時，應回覆查不到，不可編造台號。",
        ["不存在頻道在第幾台", "大屯"],
        [
            c("回覆查不到頻道", "contains_any", ["查不到", "查無"]),
            c("不可編造已知台號", "not_contains_any", ["168", "99"]),
        ],
    ),
    Scenario(
        "API-014",
        "服務範圍 API",
        "服務範圍查詢成功：可申辦",
        "TC-026",
        "大里區應回傳 mock 可進一步申辦。",
        ["台中市大里區能申辦寬頻上網嗎？"],
        [
            c("回覆可申辦", "contains_any", ["可進一步申辦", "可申辦"]),
            c("工具成功", "contains_any", ["True"], "last_tool_success"),
            c("打到 mock 服務範圍 endpoint", "contains_any", ["serviceAvailability"], "last_tool_endpoint"),
        ],
    ),
    Scenario(
        "API-015",
        "服務範圍 API",
        "服務範圍查詢失敗：不可申辦",
        "TC-026",
        "彰化和美鎮應回傳 mock 不可申辦。",
        ["彰化縣和美鎮能申辦寬頻上網嗎？"],
        [
            c("回覆不可申辦", "contains_any", ["不可申辦", "未服務此區"]),
            c("工具成功", "contains_any", ["True"], "last_tool_success"),
            c("打到 mock 服務範圍 endpoint", "contains_any", ["serviceAvailability"], "last_tool_endpoint"),
        ],
    ),
    Scenario(
        "API-016",
        "加值 API",
        "加值方案查詢成功",
        "TC-024",
        "補服務地區後，應呼叫 mock 加值方案 API 並回傳購買連結。",
        ["哈TV數位套餐加購", "大屯"],
        [
            c("回覆 mock 加值 API", "contains_all", ["模擬加值方案", "購買"]),
            c("工具成功", "contains_any", ["True"], "last_tool_success"),
            c("打到 mock 加值 endpoint", "contains_any", ["addonPlans"], "last_tool_endpoint"),
        ],
    ),
    Scenario(
        "API-017",
        "取消報修 API",
        "取消報修成功",
        "TC-034",
        "提供報修單號與正確電話後，應呼叫 mock 取消報修 API，不異動真實工單。",
        ["取消報修", "NT20260513-ABC123 0988555666"],
        [
            c("回覆 mock 取消報修", "contains_all", ["模擬取消報修", "尚未異動真實工單"]),
            c("工具成功", "contains_any", ["True"], "last_tool_success"),
            c("打到 mock 取消 endpoint", "contains_any", ["cancelRepair"], "last_tool_endpoint"),
        ],
    ),
    Scenario(
        "API-018",
        "取消報修 API",
        "取消報修失敗：電話不一致",
        "TC-034",
        "電話不一致時，應回覆查無資料，不可宣稱已送出取消。",
        ["取消報修", "NT20260513-ABC123 0988666555"],
        [
            c("回覆查無資料", "contains_any", ["查詢不到", "不一致", "重新輸入"]),
            c("不可宣稱取消成功", "not_contains_any", ["已用模擬取消報修 API 送出測試請求"]),
            c("工具失敗", "contains_any", ["False"], "last_tool_success"),
        ],
    ),
]


def status_counts(results: list[dict[str, Any]]) -> dict[str, int]:
    counts: dict[str, int] = {}
    for result in results:
        counts[result["status"]] = counts.get(result["status"], 0) + 1
    return counts


def check_health(api_base: str, timeout_sec: int) -> bool:
    try:
        resp = requests.get(f"{api_base}/health", timeout=timeout_sec)
        return resp.status_code == 200
    except Exception:
        return False


def reset_user(api_base: str, user_id: str) -> None:
    try:
        requests.post(
            f"{api_base}/reset/{user_id}",
            params={"include_tickets": False},
            timeout=15,
        )
    except Exception:
        pass


def send_chat(api_base: str, user_id: str, text: str, tv_cable: str, timeout_sec: int) -> dict[str, Any]:
    resp = requests.post(
        f"{api_base}/chat",
        json={"user_id": user_id, "user_input": text, "tv_cable": tv_cable},
        timeout=timeout_sec,
    )
    resp.raise_for_status()
    return resp.json()


def fetch_state(api_base: str, user_id: str, tv_cable: str) -> dict[str, Any]:
    try:
        resp = requests.get(
            f"{api_base}/state/{user_id}",
            params={"tv_cable": tv_cable},
            timeout=15,
        )
        if resp.status_code == 200:
            return resp.json()
    except Exception:
        pass
    return {}


def last_tool_result(result: dict[str, Any]) -> dict[str, Any]:
    state = result.get("state") or {}
    tool_result = state.get("last_tool_result")
    return tool_result if isinstance(tool_result, dict) else {}


def target_text(result: dict[str, Any], check: Check) -> str:
    state = result.get("state") or {}
    tool = last_tool_result(result)
    raw = tool.get("data", {}).get("raw", {}) if isinstance(tool.get("data"), dict) else {}

    if check.target == "final_text":
        return result.get("final_ai_response", "")
    if check.target == "pending_tool":
        return str(state.get("pending_tool") or "")
    if check.target == "last_tool_name":
        return str(tool.get("tool_name") or "")
    if check.target == "last_tool_success":
        return str(tool.get("success"))
    if check.target == "last_tool_endpoint":
        return str(raw.get("endpoint") or "")
    return ""


def evaluate_check(result: dict[str, Any], check: Check) -> dict[str, Any]:
    text = target_text(result, check)
    if check.mode == "contains_all":
        passed = all(term in text for term in check.terms)
    elif check.mode == "contains_any":
        passed = any(term in text for term in check.terms)
    elif check.mode == "not_contains_any":
        passed = not any(term in text for term in check.terms)
    else:
        passed = False

    return {
        "label": check.label,
        "mode": check.mode,
        "terms": check.terms,
        "target": check.target,
        "observed": text,
        "passed": passed,
    }


def run_scenario(api_base: str, scenario: Scenario, tv_cable: str, timeout_sec: int) -> dict[str, Any]:
    user_id = f"mock-api-{scenario.case_id}-{uuid.uuid4().hex[:8]}"
    result: dict[str, Any] = {
        "case_id": scenario.case_id,
        "category": scenario.category,
        "title": scenario.title,
        "source_ref": scenario.source_ref,
        "expected_behavior": scenario.expected_behavior,
        "excel_references": [],
        "turns": scenario.turns,
        "messages": [],
        "state": {},
        "checks": [],
        "status": "PASS",
        "error": "",
        "duration_sec": 0.0,
    }

    started = time.perf_counter()
    reset_user(api_base, user_id)

    try:
        for turn in scenario.turns:
            result["messages"].append({"role": "user", "content": turn})
            data = send_chat(api_base, user_id, turn, tv_cable, timeout_sec)
            ai_response = str(data.get("ai_response", ""))
            result["messages"].append({"role": "assistant", "content": ai_response})

        result["final_ai_response"] = result["messages"][-1]["content"] if result["messages"] else ""
        result["state"] = fetch_state(api_base, user_id, tv_cable)
        result["checks"] = [evaluate_check(result, check) for check in scenario.checks]
        if not all(item["passed"] for item in result["checks"]):
            result["status"] = "FAIL"
    except Exception as exc:
        result["status"] = "ERROR"
        result["error"] = str(exc)
    finally:
        result["duration_sec"] = round(time.perf_counter() - started, 3)

    return result


def attach_excel_references(
    results: list[dict[str, Any]],
    scenarios: list[Scenario],
    case_map: dict[str, dict[str, str]],
) -> None:
    scenario_map = {scenario.case_id: scenario for scenario in scenarios}
    for result in results:
        scenario = scenario_map.get(result.get("case_id", ""))
        if scenario:
            result["excel_references"] = excel_references_for(scenario.source_ref, case_map)


def render_badge(status: str) -> str:
    css = {"PASS": "pass", "FAIL": "fail", "ERROR": "error"}.get(status, "skip")
    return f'<span class="badge {css}">{html.escape(status)}</span>'


def render_messages(messages: list[dict[str, str]]) -> str:
    parts = []
    for message in messages:
        role = message.get("role", "")
        content = html.escape(message.get("content", "")).replace("\n", "<br>")
        icon = "🙂" if role == "user" else "🤖"
        label = "User" if role == "user" else "AI"
        parts.append(
            f"""
            <div class="chat-row {role}">
              <div class="avatar">{icon}</div>
              <div class="bubble">
                <div class="role">{label}</div>
                <div class="content">{content}</div>
              </div>
            </div>
            """
        )
    return "\n".join(parts) or '<div class="muted">未執行對話。</div>'


def render_checks(checks: list[dict[str, Any]]) -> str:
    rows = []
    for item in checks:
        icon = "✓" if item["passed"] else "✗"
        css = "ok" if item["passed"] else "bad"
        terms = " / ".join(item["terms"])
        observed = str(item.get("observed", ""))
        if len(observed) > 120:
            observed = observed[:117] + "..."
        rows.append(
            f'<li class="{css}"><b>{icon} {html.escape(item["label"])}</b>'
            f'<span>{html.escape(item["target"])} · {html.escape(item["mode"])}：{html.escape(terms)}</span>'
            f'<span>observed：{html.escape(observed)}</span></li>'
        )
    return f'<ul class="checks">{"".join(rows)}</ul>'


def render_excel_references(refs: list[dict[str, str]]) -> str:
    if not refs:
        return '<div class="excel-ref muted">未讀取到對應 Excel 案例。</div>'

    cards = []
    for item in refs:
        cards.append(
            f"""
            <div class="excel-ref">
              <div class="ref-title">{html.escape(item.get('案例ID', ''))} · {html.escape(item.get('類別', ''))}</div>
              <div><b>情境：</b>{html.escape(item.get('情境', ''))}</div>
              <div><b>輸入：</b>{html.escape(item.get('使用者輸入範例（去識別化）', ''))}</div>
              <div><b>期望：</b>{html.escape(item.get('期望AI行為', ''))}</div>
              <div><b>檢核：</b>{html.escape(item.get('測試檢核點', ''))}</div>
            </div>
            """
        )
    return "\n".join(cards)


def render_tool_result(result: dict[str, Any]) -> str:
    state = result.get("state") or {}
    tool = last_tool_result(result)
    raw = tool.get("data", {}).get("raw", {}) if isinstance(tool.get("data"), dict) else {}
    endpoint = raw.get("endpoint")
    if not tool and not state.get("pending_tool"):
        return '<div class="tool muted">本案例尚未呼叫工具。</div>'

    success = tool.get("success") if tool else None
    success_text = "未呼叫" if success is None else str(success)
    css = "ok" if success is True else "bad" if success is False else "pending"
    lines = [
        f"<b>工具狀態</b>",
        f"pending_tool：{html.escape(str(state.get('pending_tool') or ''))}",
        f"last_tool：{html.escape(str(tool.get('tool_name') or ''))}",
        f"success：<span class=\"{css}\">{html.escape(success_text)}</span>",
    ]
    if endpoint:
        lines.append(f"endpoint：{html.escape(str(endpoint))}")
    if tool.get("message"):
        lines.append(f"message：{html.escape(str(tool.get('message'))).replace(chr(10), '<br>')}")
    return '<div class="tool">' + "<br>".join(lines) + "</div>"


def render_html_report(
    results: list[dict[str, Any]],
    output_path: Path,
    workbook_path: Path | None = None,
    excel_case_count: int = 0,
) -> None:
    counts = status_counts(results)
    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    cards = []
    for result in results:
        meta = (
            f"{html.escape(result['case_id'])} · {html.escape(result['category'])} · "
            f"參考 {html.escape(result['source_ref'])}"
        )
        error = ""
        if result.get("error"):
            error = f'<div class="notice error-text">{html.escape(result["error"])}</div>'
        cards.append(
            f"""
            <section class="case {html.escape(result['status'].lower())}">
              <div class="case-head">
                <div>
                  <div class="meta">{meta}</div>
                  <h2>{html.escape(result['title'])}</h2>
                </div>
                {render_badge(result['status'])}
              </div>
              <div class="grid">
                <div>
                  <h3>期望行為</h3>
                  <p>{html.escape(result['expected_behavior'])}</p>
                  <h3>Excel 參考案例</h3>
                  {render_excel_references(result.get('excel_references', []))}
                  {error}
                  {render_tool_result(result)}
                  {render_checks(result.get('checks', []))}
                </div>
                <div class="chatbox">
                  {render_messages(result.get('messages', []))}
                </div>
              </div>
            </section>
            """
        )

    source_text = "reports/20260508/dialog_scenario_report.html"
    if workbook_path:
        source_text = f"{workbook_path}（已讀取 {excel_case_count} 筆） · {source_text}"

    html_text = f"""<!doctype html>
<html lang="zh-Hant">
<head>
  <meta charset="utf-8">
  <title>Mock API 對話情境測試報告</title>
  <style>
    :root {{
      color-scheme: dark;
      --bg: #0f1117;
      --panel: #181b22;
      --panel-2: #20242d;
      --text: #f3f5f8;
      --muted: #a8b0bf;
      --user: #ff0a4f;
      --ai: #ff7b22;
      --ok: #39d98a;
      --bad: #ff5c7a;
      --skip: #f2c94c;
      --line: #2c3240;
    }}
    * {{ box-sizing: border-box; }}
    body {{
      margin: 0;
      background: var(--bg);
      color: var(--text);
      font: 15px/1.6 "Microsoft JhengHei", "Noto Sans TC", system-ui, sans-serif;
    }}
    header {{
      padding: 28px 36px 18px;
      border-bottom: 1px solid var(--line);
      background: #12151c;
      position: sticky;
      top: 0;
      z-index: 10;
    }}
    h1 {{ margin: 0 0 8px; font-size: 28px; }}
    .sub {{ color: var(--muted); }}
    .summary {{
      display: flex;
      gap: 12px;
      margin-top: 18px;
      flex-wrap: wrap;
    }}
    .summary .tile {{
      background: var(--panel);
      border: 1px solid var(--line);
      border-radius: 8px;
      padding: 10px 14px;
      min-width: 120px;
    }}
    .tile b {{ display: block; font-size: 22px; }}
    main {{ padding: 24px 36px 60px; }}
    .case {{
      background: var(--panel);
      border: 1px solid var(--line);
      border-radius: 10px;
      padding: 18px;
      margin-bottom: 18px;
    }}
    .case.fail, .case.error {{ border-color: rgba(255, 92, 122, .65); }}
    .case-head {{
      display: flex;
      justify-content: space-between;
      gap: 12px;
      align-items: flex-start;
      margin-bottom: 12px;
    }}
    .meta {{ color: var(--muted); font-size: 13px; }}
    h2 {{ margin: 2px 0 0; font-size: 20px; }}
    h3 {{ margin: 14px 0 4px; font-size: 14px; color: var(--muted); }}
    p {{ margin: 0; }}
    .grid {{
      display: grid;
      grid-template-columns: minmax(300px, 460px) minmax(420px, 1fr);
      gap: 18px;
    }}
    .badge {{
      border-radius: 999px;
      padding: 4px 10px;
      font-weight: 700;
      font-size: 12px;
      white-space: nowrap;
    }}
    .badge.pass {{ background: rgba(57, 217, 138, .16); color: var(--ok); }}
    .badge.fail, .badge.error {{ background: rgba(255, 92, 122, .16); color: var(--bad); }}
    .chatbox, .tool {{
      background: #0d1016;
      border: 1px solid var(--line);
      border-radius: 8px;
      padding: 12px;
    }}
    .tool {{
      margin-top: 14px;
      font-family: Consolas, "Microsoft JhengHei", monospace;
      font-size: 13px;
    }}
    .excel-ref {{
      margin-top: 8px;
      background: #10141b;
      border: 1px solid var(--line);
      border-radius: 8px;
      padding: 10px 12px;
      color: var(--muted);
      font-size: 13px;
    }}
    .excel-ref b, .ref-title {{
      color: var(--text);
    }}
    .ref-title {{
      font-weight: 700;
      margin-bottom: 4px;
    }}
    .chat-row {{
      display: flex;
      gap: 10px;
      margin: 10px 0;
      align-items: flex-start;
    }}
    .avatar {{
      width: 40px;
      height: 40px;
      border-radius: 8px;
      display: grid;
      place-items: center;
      flex: 0 0 auto;
      background: var(--ai);
    }}
    .chat-row.user .avatar {{ background: var(--user); }}
    .bubble {{
      background: var(--panel-2);
      border-radius: 8px;
      padding: 10px 12px;
      width: 100%;
      min-height: 42px;
    }}
    .role {{
      color: var(--muted);
      font-size: 12px;
      margin-bottom: 2px;
    }}
    .checks {{
      list-style: none;
      padding: 0;
      margin: 14px 0 0;
    }}
    .checks li {{
      border-radius: 8px;
      padding: 8px 10px;
      margin-bottom: 8px;
      background: #10141b;
      border: 1px solid var(--line);
    }}
    .checks li.ok b, .ok {{ color: var(--ok); }}
    .checks li.bad b, .bad {{ color: var(--bad); }}
    .pending {{ color: var(--skip); }}
    .checks span {{
      display: block;
      color: var(--muted);
      font-size: 12px;
    }}
    .notice {{
      margin-top: 12px;
      color: var(--bad);
      background: rgba(255, 92, 122, .08);
      border: 1px solid rgba(255, 92, 122, .25);
      border-radius: 8px;
      padding: 8px 10px;
    }}
    .muted {{ color: var(--muted); }}
    @media (max-width: 920px) {{
      header, main {{ padding-left: 16px; padding-right: 16px; }}
      .grid {{ grid-template-columns: 1fr; }}
    }}
  </style>
</head>
<body>
  <header>
    <h1>Mock API 對話情境測試報告</h1>
    <div class="sub">參考：{html.escape(source_text)} · 產生時間：{generated_at}</div>
    <div class="summary">
      <div class="tile"><span>Total</span><b>{len(results)}</b></div>
      <div class="tile"><span>PASS</span><b>{counts.get('PASS', 0)}</b></div>
      <div class="tile"><span>FAIL</span><b>{counts.get('FAIL', 0)}</b></div>
      <div class="tile"><span>ERROR</span><b>{counts.get('ERROR', 0)}</b></div>
    </div>
  </header>
  <main>
    {''.join(cards)}
  </main>
</body>
</html>
"""
    output_path.write_text(html_text, encoding="utf-8")


def write_json_report(results: list[dict[str, Any]], output_path: Path) -> None:
    payload = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "summary": status_counts(results),
        "results": results,
    }
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--api-base", default=DEFAULT_API_BASE)
    parser.add_argument("--tv-cable", default=DEFAULT_TV_CABLE)
    parser.add_argument("--xlsx", default=DEFAULT_XLSX)
    parser.add_argument("--case", default=None, help="Comma-separated scenario IDs")
    parser.add_argument("--output-dir", default="reports/20260513")
    parser.add_argument("--request-timeout", type=int, default=60)
    args = parser.parse_args()

    workbook_path = Path(args.xlsx)
    case_map = load_excel_case_map(workbook_path)

    scenarios = SCENARIOS
    if args.case:
        allowed = {item.strip() for item in args.case.split(",") if item.strip()}
        scenarios = [item for item in scenarios if item.case_id in allowed]

    if not check_health(args.api_base, args.request_timeout):
        raise RuntimeError(f"Backend is not reachable: {args.api_base}/health")

    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    html_path = output_dir / "dialog_scenario_report.html"
    json_path = output_dir / "dialog_scenario_report.json"

    results = []
    total = len(scenarios)
    for index, scenario in enumerate(scenarios, start=1):
        print(f"[{index}/{total}] {scenario.case_id} {scenario.title}", flush=True)
        result = run_scenario(args.api_base, scenario, args.tv_cable, args.request_timeout)
        result["excel_references"] = excel_references_for(scenario.source_ref, case_map)
        results.append(result)
        print(f"  -> {result['status']} ({result['duration_sec']}s)", flush=True)
        render_html_report(results, html_path, workbook_path=workbook_path, excel_case_count=len(case_map))
        write_json_report(results, json_path)

    print("Summary:", status_counts(results))
    print(f"HTML report: {html_path}")
    print(f"JSON report: {json_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
