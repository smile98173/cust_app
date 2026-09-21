"""
Run conversation scenario tests from the Excel workbook and render a chat-style report.

The runner intentionally avoids scenarios that require a real external account/tool
success response. Those cases are marked SKIP_EXTERNAL so the report can be shared
without triggering customer-facing APIs.
"""

from __future__ import annotations

import argparse
import html
import json
import time
import uuid
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

import requests
from openpyxl import load_workbook


DEFAULT_API_BASE = "http://127.0.0.1:8000"
DEFAULT_TV_CABLE = "tdtv"


@dataclass
class Check:
    label: str
    mode: str
    terms: list[str]
    target: str = "final_text"


@dataclass
class Scenario:
    case_id: str
    category: str
    title: str
    context: str
    excel_input: str
    expected_behavior: str
    slot_intent: str
    checklist: str
    priority: str
    source: str
    turns: list[str] = field(default_factory=list)
    checks: list[Check] = field(default_factory=list)
    skip_reason: str | None = None


def find_workbook(root: Path, explicit: str | None) -> Path:
    if explicit:
        path = Path(explicit)
        if path.exists():
            return path
        raise FileNotFoundError(path)

    matches = list(root.glob("*.xlsx"))
    if not matches:
        raise FileNotFoundError("No .xlsx workbook found in project root")

    return matches[0]


def load_cases(workbook_path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    ws = wb.worksheets[1]
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(v or "") for v in rows[0]]

    cases: list[dict[str, Any]] = []
    for row in rows[1:]:
        if not row or not row[0] or not str(row[0]).startswith("TC-"):
            continue
        item = {
            headers[i]: "" if value is None else str(value)
            for i, value in enumerate(row)
            if i < len(headers)
        }
        cases.append(item)
    return cases


def c(label: str, mode: str, terms: list[str], target: str = "final_text") -> Check:
    return Check(label=label, mode=mode, terms=terms, target=target)


SCENARIO_OVERRIDES: dict[str, dict[str, Any]] = {
    "TC-001": {
        "turns": ["忘了繳費已被斷訊"],
        "checks": [
            c("需釐清電視或網路", "contains_all", ["電視", "網路"]),
            c("不應直接要求個資", "not_contains_any", ["姓名", "電話"]),
        ],
    },
    "TC-002": {
        "turns": ["忘了繳費已被斷訊", "網路"],
        "checks": [
            c("需要求戶名或姓名", "contains_any", ["姓名", "戶名"]),
            c("需要求聯絡電話", "contains_any", ["電話", "聯絡電話"]),
            c("不應重複問電視/網路", "not_contains_any", ["電視訊號或網路訊號"]),
        ],
    },
    "TC-003": {
        "turns": ["忘了繳費已被斷訊", "網路", "王小明 0912345678"],
        "skip_reason": "需要外部復線 API 或 mock 回傳成功；live 測試不觸發實際工具成功。",
    },
    "TC-004": {
        "turns": ["忘了繳費已被斷訊", "網路", "王小明"],
        "checks": [c("只追問缺漏電話", "contains_any", ["電話", "聯絡電話"])],
    },
    "TC-005": {
        "turns": ["忘了繳費已被斷訊", "網路", "0912345678"],
        "checks": [c("只追問缺漏姓名", "contains_any", ["姓名", "戶名"])],
    },
    "TC-006": {
        "turns": ["已繳費無法收看"],
        "checks": [
            c("需辨識電視/收視情境", "contains_any", ["電視", "收看", "訊號"]),
            c("需進一步釐清或要求身份資料", "contains_any", ["姓名", "電話", "是否", "請問"]),
        ],
    },
    "TC-007": {
        "turns": ["本期帳單金額查詢"],
        "checks": [
            c("查帳單前需要求戶名或姓名", "contains_any", ["戶名", "姓名"]),
            c("查帳單前需要求聯絡電話", "contains_any", ["電話", "聯絡電話"]),
        ],
    },
    "TC-008": {
        "turns": ["查詢帳單繳費截止日期"],
        "checks": [
            c("需要求身份資料或說明查詢方式", "contains_any", ["戶名", "姓名", "電話", "登入", "客服"]),
            c("不可導向不相關方案", "not_contains_any", ["網路方案", "優惠方案"]),
        ],
    },
    "TC-009": {
        "turns": ["繳款方式查詢"],
        "checks": [c("需回答繳費方式", "contains_any", ["線上", "APP", "臨櫃", "超商", "ATM", "繳費"])],
    },
    "TC-010": {
        "turns": ["用戶更改扣款帳號"],
        "checks": [c("帳務個資異動需升級正式流程", "contains_any", ["真人客服", "客服", "個資", "帳務", "身份"])],
    },
    "TC-011": {
        "turns": ["想請問我的合約到期日"],
        "checks": [
            c("不可捏造日期，需導向查詢或身份驗證", "contains_any", ["登入", "APP", "官網", "身份", "客服", "查詢"]),
            c("不可直接編造到期日", "not_contains_any", ["2026-", "2027-"]),
        ],
    },
    "TC-012": {
        "turns": ["無法連線"],
        "checks": [c("網路故障需先釐清範圍", "contains_any", ["所有設備", "單一", "手機", "電腦", "數據機"])],
    },
    "TC-013": {
        "turns": ["連不上wifi"],
        "checks": [c("Wi-Fi 故障需提供可操作排查", "contains_any", ["Wi-Fi", "wifi", "重新", "數據機", "分享器", "密碼"])],
    },
    "TC-014": {
        "turns": ["一直閃紅燈"],
        "checks": [c("燈號異常需導向重啟或報修", "contains_any", ["紅燈", "燈號", "重", "報修", "檢查"])],
    },
    "TC-015": {
        "turns": ["網路排除"],
        "checks": [c("選單式短句需先釐清症狀", "contains_any", ["無法連線", "速度慢", "Wi-Fi", "燈號", "請問"])],
    },
    "TC-016": {
        "turns": ["請問我家中原有20M網路 想升級300M 需要更改線路嗎"],
        "checks": [
            c("升速需保守回答並導向確認", "contains_any", ["不一定", "可能", "更換", "客服", "確認", "線路"]),
            c("不可保證免施工", "not_contains_any", ["一定不用", "保證不用"]),
        ],
    },
    "TC-017": {
        "turns": ["我的網路是幾m"],
        "checks": [c("目前速率需導向自助或身份查詢", "contains_any", ["APP", "官網", "身份", "客服", "查詢", "登入"])],
    },
    "TC-018": {
        "turns": ["家用有線電視加網路"],
        "checks": [c("組合方案需收集地區或需求", "contains_any", ["地區", "需求", "速度", "方案", "客服"])],
    },
    "TC-019": {
        "turns": ["新春換新網"],
        "checks": [c("活動資訊需保守或以公告為準", "contains_any", ["活動", "方案", "公告", "客服", "目前沒有"])],
    },
    "TC-020": {
        "turns": ["沒訊號"],
        "checks": [c("電視無訊號需先進排查", "contains_any", ["機上盒", "電源", "訊號源", "重", "畫面"])],
    },
    "TC-021": {
        "turns": ["機上盒排除"],
        "checks": [c("機上盒排除需問狀態或提供步驟", "contains_any", ["機上盒", "電源", "錯誤", "訊號源", "線材", "請問"])],
    },
    "TC-022": {
        "turns": ["遙控器有紅色燈沒有反應"],
        "checks": [c("遙控器問題需給操作檢查", "contains_any", ["電池", "配對", "遙控器", "清潔", "更換"])],
    },
    "TC-023": {
        "turns": ["愛爾達體育在第幾台"],
        "checks": [c("頻道查詢需要求地區或系統台", "contains_any", ["地區", "系統台", "頻道", "服務區"])],
    },
    "TC-024": {
        "turns": ["哈tv數位套餐加購"],
        "checks": [c("加購需保守並導向方案確認", "contains_any", ["地區", "方案", "客服", "官網", "加購"])],
    },
    "TC-025": {
        "turns": ["網路裝機申請"],
        "checks": [c("新申裝需收集必要資料或導向正式申請", "contains_any", ["地址", "姓名", "電話", "方案", "客服", "申請"])],
    },
    "TC-026": {
        "turns": ["彰化縣和美鎮能申辦寬頻上網嗎？"],
        "checks": [c("服務範圍需依系統台地區保守回答", "contains_any", ["服務地區", "服務範圍", "客服", "確認", "彰化"])],
    },
    "TC-027": {
        "turns": ["我要搬家，網路要移機"],
        "checks": [c("移機需正式流程與必要資料", "contains_any", ["原地址", "新地址", "電話", "施工", "客服", "移機"])],
    },
    "TC-028": {
        "turns": ["申請ip"],
        "checks": [c("固定 IP 申請需回答規則或導向客服", "contains_any", ["固定", "IP", "200", "客服", "申請"])],
    },
    "TC-029": {
        "turns": ["光纖進來接的zyxel，是設定為路由器還是橋接器"],
        "checks": [c("進階設定需說明差異並保守", "contains_any", ["路由", "橋接", "設定", "客服", "差異"])],
    },
    "TC-030": {
        "turns": ["忘記密碼"],
        "checks": [c("忘記密碼需導向重設或客服", "contains_any", ["忘記密碼", "重設", "客服", "身份", "登入"])],
    },
    "TC-031": {
        "turns": ["哈TV行動客服怎麼查帳單"],
        "checks": [c("APP 查帳單需步驟化", "contains_any", ["APP", "登入", "帳務", "帳單", "查詢"])],
    },
    "TC-032": {
        "turns": ["真人客服"],
        "checks": [c("明確人工需求需尊重", "contains_any", ["真人客服", "客服", "人工"])],
    },
    "TC-033": {
        "turns": ["退租"],
        "checks": [c("退租需正式流程與注意事項", "contains_any", ["設備", "歸還", "費用", "客服", "臨櫃", "退租"])],
    },
    "TC-034": {
        "turns": ["取消報修"],
        "checks": [
            c("取消報修需客服協助", "contains_any", ["客服", "報修", "取消"]),
            c("不可未查證就承諾取消成功", "not_contains_any", ["已取消", "取消完成"]),
        ],
    },
    "TC-035": {
        "turns": ["0"],
        "checks": [c("未知短句需安全 fallback", "contains_any", ["不太確定", "需要", "協助", "請問"])],
    },
    "TC-036": {
        "turns": ["用戶上傳圖片"],
        "checks": [c("無圖內容需請使用者描述", "contains_any", ["圖片", "描述", "清晰", "重新"])],
    },
}


def build_scenarios(cases: list[dict[str, Any]]) -> list[Scenario]:
    scenarios: list[Scenario] = []
    for item in cases:
        case_id = item.get("案例ID", "")
        override = SCENARIO_OVERRIDES.get(case_id, {})
        turns = override.get("turns") or [item.get("使用者輸入範例（去識別化）", "")]
        checks = override.get("checks") or [
            c("基本檢查：AI 有回覆", "contains_any", ["請", "您好", "目前", "可以", "客服"])
        ]
        scenario = Scenario(
            case_id=case_id,
            category=item.get("類別", ""),
            title=item.get("情境", ""),
            context=item.get("背景/前置條件", ""),
            excel_input=item.get("使用者輸入範例（去識別化）", ""),
            expected_behavior=item.get("期望AI行為", ""),
            slot_intent=item.get("槽位/意圖", ""),
            checklist=item.get("測試檢核點", ""),
            priority=item.get("優先級", ""),
            source=item.get("來源依據", ""),
            turns=turns,
            checks=checks,
            skip_reason=override.get("skip_reason"),
        )
        scenarios.append(scenario)
    return scenarios


def check_health(api_base: str, timeout_sec: int) -> bool:
    try:
        resp = requests.get(f"{api_base}/health", timeout=timeout_sec)
        return resp.status_code == 200
    except Exception:
        return False


def reset_user(api_base: str, user_id: str) -> None:
    try:
        requests.post(
            f"{api_base}/reset/{user_id}",
            params={"include_tickets": False},
            timeout=15,
        )
    except Exception:
        pass


def send_chat(
    api_base: str,
    user_id: str,
    text: str,
    tv_cable: str,
    timeout_sec: int,
) -> dict[str, Any]:
    resp = requests.post(
        f"{api_base}/chat",
        json={"user_id": user_id, "user_input": text, "tv_cable": tv_cable},
        timeout=timeout_sec,
    )
    resp.raise_for_status()
    return resp.json()


def fetch_state(api_base: str, user_id: str, tv_cable: str) -> dict[str, Any]:
    try:
        resp = requests.get(
            f"{api_base}/state/{user_id}",
            params={"tv_cable": tv_cable},
            timeout=15,
        )
        if resp.status_code == 200:
            return resp.json()
    except Exception:
        pass
    return {}


def target_text(result: dict[str, Any], check: Check) -> str:
    if check.target == "final_text":
        return result.get("final_ai_response", "")
    if check.target == "decision_type":
        return str(result.get("state", {}).get("decision_type", ""))
    if check.target == "pending_tool":
        return str(result.get("state", {}).get("pending_tool", ""))
    return ""


def evaluate_check(result: dict[str, Any], check: Check) -> dict[str, Any]:
    text = target_text(result, check)
    if check.mode == "contains_all":
        passed = all(term in text for term in check.terms)
    elif check.mode == "contains_any":
        passed = any(term in text for term in check.terms)
    elif check.mode == "not_contains_any":
        passed = not any(term in text for term in check.terms)
    else:
        passed = False

    return {
        "label": check.label,
        "mode": check.mode,
        "terms": check.terms,
        "target": check.target,
        "passed": passed,
    }


def run_scenario(
    api_base: str,
    scenario: Scenario,
    tv_cable: str,
    request_timeout: int,
) -> dict[str, Any]:
    user_id = f"scenario-{scenario.case_id}-{uuid.uuid4().hex[:8]}"
    result: dict[str, Any] = {
        "case_id": scenario.case_id,
        "category": scenario.category,
        "title": scenario.title,
        "context": scenario.context,
        "excel_input": scenario.excel_input,
        "expected_behavior": scenario.expected_behavior,
        "slot_intent": scenario.slot_intent,
        "checklist": scenario.checklist,
        "priority": scenario.priority,
        "source": scenario.source,
        "turns": scenario.turns,
        "messages": [],
        "state": {},
        "checks": [],
        "status": "PASS",
        "skip_reason": scenario.skip_reason,
        "error": "",
        "duration_sec": 0.0,
    }

    if scenario.skip_reason:
        result["status"] = "SKIP_EXTERNAL"
        return result

    started = time.perf_counter()
    reset_user(api_base, user_id)

    try:
        for turn in scenario.turns:
            result["messages"].append({"role": "user", "content": turn})
            data = send_chat(api_base, user_id, turn, tv_cable, request_timeout)
            ai_response = str(data.get("ai_response", ""))
            result["messages"].append({"role": "assistant", "content": ai_response})

        result["final_ai_response"] = result["messages"][-1]["content"] if result["messages"] else ""
        result["state"] = fetch_state(api_base, user_id, tv_cable)
        result["checks"] = [evaluate_check(result, check) for check in scenario.checks]
        if not all(item["passed"] for item in result["checks"]):
            result["status"] = "FAIL"
    except Exception as exc:
        result["status"] = "ERROR"
        result["error"] = str(exc)
    finally:
        result["duration_sec"] = round(time.perf_counter() - started, 3)

    return result


def status_counts(results: list[dict[str, Any]]) -> dict[str, int]:
    counts: dict[str, int] = {}
    for result in results:
        counts[result["status"]] = counts.get(result["status"], 0) + 1
    return counts


def render_badge(status: str) -> str:
    css = {
        "PASS": "pass",
        "FAIL": "fail",
        "ERROR": "error",
        "SKIP_EXTERNAL": "skip",
    }.get(status, "skip")
    return f'<span class="badge {css}">{html.escape(status)}</span>'


def render_messages(messages: list[dict[str, str]]) -> str:
    if not messages:
        return '<div class="muted">未執行 live 對話。</div>'

    parts = []
    for message in messages:
        role = message.get("role", "")
        content = html.escape(message.get("content", "")).replace("\n", "<br>")
        icon = "🙂" if role == "user" else "🤖"
        label = "User" if role == "user" else "AI"
        parts.append(
            f"""
            <div class="chat-row {role}">
              <div class="avatar">{icon}</div>
              <div class="bubble">
                <div class="role">{label}</div>
                <div class="content">{content}</div>
              </div>
            </div>
            """
        )
    return "\n".join(parts)


def render_checks(checks: list[dict[str, Any]]) -> str:
    if not checks:
        return '<div class="muted">此案例未執行自動檢核。</div>'

    rows = []
    for item in checks:
        icon = "✓" if item["passed"] else "✗"
        css = "ok" if item["passed"] else "bad"
        terms = " / ".join(item["terms"])
        rows.append(
            f'<li class="{css}"><b>{icon} {html.escape(item["label"])}</b>'
            f'<span>{html.escape(item["mode"])}：{html.escape(terms)}</span></li>'
        )
    return f'<ul class="checks">{"".join(rows)}</ul>'


def render_html_report(results: list[dict[str, Any]], workbook_path: Path, output_path: Path) -> None:
    counts = status_counts(results)
    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    cards = []
    for result in results:
        meta = (
            f"{html.escape(result['case_id'])} · {html.escape(result['priority'])} · "
            f"{html.escape(result['category'])}"
        )
        skip = ""
        if result.get("skip_reason"):
            skip = f'<div class="notice">{html.escape(result["skip_reason"])}</div>'
        error = ""
        if result.get("error"):
            error = f'<div class="notice error-text">{html.escape(result["error"])}</div>'
        cards.append(
            f"""
            <section class="case {html.escape(result['status'].lower())}">
              <div class="case-head">
                <div>
                  <div class="meta">{meta}</div>
                  <h2>{html.escape(result['title'])}</h2>
                </div>
                {render_badge(result['status'])}
              </div>
              <div class="grid">
                <div>
                  <h3>Excel 期望</h3>
                  <p>{html.escape(result['expected_behavior'])}</p>
                  <h3>檢核點</h3>
                  <p>{html.escape(result['checklist'])}</p>
                  {skip}
                  {error}
                  {render_checks(result.get('checks', []))}
                </div>
                <div class="chatbox">
                  {render_messages(result.get('messages', []))}
                </div>
              </div>
            </section>
            """
        )

    html_text = f"""<!doctype html>
<html lang="zh-Hant">
<head>
  <meta charset="utf-8">
  <title>AI 對話情境測試報告</title>
  <style>
    :root {{
      color-scheme: dark;
      --bg: #0f1117;
      --panel: #181b22;
      --panel-2: #20242d;
      --text: #f3f5f8;
      --muted: #a8b0bf;
      --user: #ff0a4f;
      --ai: #ff7b22;
      --ok: #39d98a;
      --bad: #ff5c7a;
      --skip: #f2c94c;
      --line: #2c3240;
    }}
    * {{ box-sizing: border-box; }}
    body {{
      margin: 0;
      background: var(--bg);
      color: var(--text);
      font: 15px/1.6 "Microsoft JhengHei", "Noto Sans TC", system-ui, sans-serif;
    }}
    header {{
      padding: 28px 36px 18px;
      border-bottom: 1px solid var(--line);
      background: #12151c;
      position: sticky;
      top: 0;
      z-index: 10;
    }}
    h1 {{ margin: 0 0 8px; font-size: 28px; }}
    .sub {{ color: var(--muted); }}
    .summary {{
      display: flex;
      gap: 12px;
      margin-top: 18px;
      flex-wrap: wrap;
    }}
    .summary .tile {{
      background: var(--panel);
      border: 1px solid var(--line);
      border-radius: 8px;
      padding: 10px 14px;
      min-width: 120px;
    }}
    .tile b {{ display: block; font-size: 22px; }}
    main {{ padding: 24px 36px 60px; }}
    .case {{
      background: var(--panel);
      border: 1px solid var(--line);
      border-radius: 10px;
      padding: 18px;
      margin-bottom: 18px;
    }}
    .case.fail {{ border-color: rgba(255, 92, 122, .45); }}
    .case.error {{ border-color: rgba(255, 92, 122, .75); }}
    .case.skip_external {{ border-color: rgba(242, 201, 76, .45); }}
    .case-head {{
      display: flex;
      justify-content: space-between;
      gap: 12px;
      align-items: flex-start;
      margin-bottom: 12px;
    }}
    .meta {{ color: var(--muted); font-size: 13px; }}
    h2 {{ margin: 2px 0 0; font-size: 20px; }}
    h3 {{ margin: 14px 0 4px; font-size: 14px; color: var(--muted); }}
    p {{ margin: 0; }}
    .grid {{
      display: grid;
      grid-template-columns: minmax(280px, 420px) minmax(420px, 1fr);
      gap: 18px;
    }}
    .badge {{
      border-radius: 999px;
      padding: 4px 10px;
      font-weight: 700;
      font-size: 12px;
      white-space: nowrap;
    }}
    .badge.pass {{ background: rgba(57, 217, 138, .16); color: var(--ok); }}
    .badge.fail, .badge.error {{ background: rgba(255, 92, 122, .16); color: var(--bad); }}
    .badge.skip {{ background: rgba(242, 201, 76, .16); color: var(--skip); }}
    .chatbox {{
      background: #0d1016;
      border: 1px solid var(--line);
      border-radius: 8px;
      padding: 12px;
    }}
    .chat-row {{
      display: flex;
      gap: 10px;
      margin: 10px 0;
      align-items: flex-start;
    }}
    .chat-row.user {{ justify-content: flex-start; }}
    .avatar {{
      width: 40px;
      height: 40px;
      border-radius: 8px;
      display: grid;
      place-items: center;
      flex: 0 0 auto;
      background: var(--ai);
    }}
    .chat-row.user .avatar {{ background: var(--user); }}
    .bubble {{
      background: var(--panel-2);
      border-radius: 8px;
      padding: 10px 12px;
      width: 100%;
      min-height: 42px;
    }}
    .role {{
      color: var(--muted);
      font-size: 12px;
      margin-bottom: 2px;
    }}
    .content {{ white-space: normal; }}
    .checks {{
      list-style: none;
      padding: 0;
      margin: 14px 0 0;
    }}
    .checks li {{
      border-radius: 8px;
      padding: 8px 10px;
      margin-bottom: 8px;
      background: #10141b;
      border: 1px solid var(--line);
    }}
    .checks li.ok b {{ color: var(--ok); }}
    .checks li.bad b {{ color: var(--bad); }}
    .checks span {{
      display: block;
      color: var(--muted);
      font-size: 12px;
    }}
    .notice {{
      margin-top: 12px;
      color: var(--skip);
      background: rgba(242, 201, 76, .08);
      border: 1px solid rgba(242, 201, 76, .25);
      border-radius: 8px;
      padding: 8px 10px;
    }}
    .error-text {{
      color: var(--bad);
      background: rgba(255, 92, 122, .08);
      border-color: rgba(255, 92, 122, .25);
    }}
    .muted {{ color: var(--muted); }}
    @media (max-width: 920px) {{
      header, main {{ padding-left: 16px; padding-right: 16px; }}
      .grid {{ grid-template-columns: 1fr; }}
    }}
  </style>
</head>
<body>
  <header>
    <h1>AI 對話情境測試報告</h1>
    <div class="sub">來源：{html.escape(str(workbook_path))} · 產生時間：{generated_at}</div>
    <div class="summary">
      <div class="tile"><span>Total</span><b>{len(results)}</b></div>
      <div class="tile"><span>PASS</span><b>{counts.get('PASS', 0)}</b></div>
      <div class="tile"><span>FAIL</span><b>{counts.get('FAIL', 0)}</b></div>
      <div class="tile"><span>ERROR</span><b>{counts.get('ERROR', 0)}</b></div>
      <div class="tile"><span>SKIP_EXTERNAL</span><b>{counts.get('SKIP_EXTERNAL', 0)}</b></div>
    </div>
  </header>
  <main>
    {''.join(cards)}
  </main>
</body>
</html>
"""
    output_path.write_text(html_text, encoding="utf-8")


def write_json_report(results: list[dict[str, Any]], output_path: Path) -> None:
    payload = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "summary": status_counts(results),
        "results": results,
    }
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def print_console_summary(results: list[dict[str, Any]]) -> None:
    counts = status_counts(results)
    print("Summary:", counts)
    for result in results:
        print(f"{result['case_id']} {result['status']} {result['title']}")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", default=None)
    parser.add_argument("--api-base", default=DEFAULT_API_BASE)
    parser.add_argument("--tv-cable", default=DEFAULT_TV_CABLE)
    parser.add_argument("--priority", default="P0,P1,P2")
    parser.add_argument("--case", default=None, help="Comma-separated case IDs, e.g. TC-001,TC-002")
    parser.add_argument("--output-dir", default="reports")
    parser.add_argument("--request-timeout", type=int, default=45)
    args = parser.parse_args()

    root = Path.cwd()
    workbook_path = find_workbook(root, args.xlsx)
    cases = load_cases(workbook_path)
    scenarios = build_scenarios(cases)

    allowed_priorities = {item.strip() for item in args.priority.split(",") if item.strip()}
    scenarios = [item for item in scenarios if item.priority in allowed_priorities]

    if args.case:
        allowed_cases = {item.strip() for item in args.case.split(",") if item.strip()}
        scenarios = [item for item in scenarios if item.case_id in allowed_cases]

    if not check_health(args.api_base, args.request_timeout):
        raise RuntimeError(f"Backend is not reachable: {args.api_base}/health")

    output_dir = root / args.output_dir
    output_dir.mkdir(parents=True, exist_ok=True)
    html_path = output_dir / "dialog_scenario_report.html"
    json_path = output_dir / "dialog_scenario_report.json"

    results = []
    total = len(scenarios)
    for index, scenario in enumerate(scenarios, start=1):
        print(f"[{index}/{total}] {scenario.case_id} {scenario.priority} {scenario.title}", flush=True)
        result = run_scenario(
            args.api_base,
            scenario,
            args.tv_cable,
            args.request_timeout,
        )
        results.append(result)
        print(f"  -> {result['status']} ({result['duration_sec']}s)", flush=True)
        render_html_report(results, workbook_path, html_path)
        write_json_report(results, json_path)

    print_console_summary(results)
    print(f"HTML report: {html_path}")
    print(f"JSON report: {json_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
