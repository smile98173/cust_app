from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "customer_service_kb_qa_2026-06-19.md"
OUTPUT = ROOT / "docs" / "customer_service_kb_qa_2026-06-19.docx"


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(32, 42, 54)
MUTED = RGBColor(102, 112, 128)


def set_font(run, name: str = "Calibri", east_asia: str = "Microsoft JhengHei") -> None:
    run.font.name = name
    r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), east_asia)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.10


def add_bottom_border(paragraph, color: str = "D9E2F3") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run("2026-06-19 客服需補充 QA")
    set_font(run)
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = RGBColor(11, 37, 69)
    add_bottom_border(p)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(12)
    r = meta.add_run("整理目的：提供客服與知識庫維護者直接補入 QA 條目的標準稿。")
    set_font(r)
    r.font.size = Pt(10)
    r.font.color.rgb = MUTED


def add_footer(doc: Document) -> None:
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("Customer Service KB QA")
    set_font(r)
    r.font.size = Pt(9)
    r.font.color.rgb = MUTED


def clean(text: str) -> str:
    return text.replace("`", "").replace("**", "")


def add_content(doc: Document, markdown: str) -> None:
    for raw in markdown.splitlines():
        line = raw.rstrip()
        if not line:
            continue
        if line.startswith("# "):
            continue
        if line.startswith("## "):
            doc.add_heading(clean(line[3:]), level=1)
            continue
        if line.startswith("### "):
            doc.add_heading(clean(line[4:]), level=2)
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(clean(line[2:]))
            set_font(r)
            r.font.size = Pt(10.5)
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        text = clean(line)
        if text.startswith(("A：", "需追問：", "需補資料：", "已知客服回饋：", "轉真人條件：")):
            label, _, rest = text.partition("：")
            r1 = p.add_run(label + "：")
            set_font(r1)
            r1.font.bold = True
            r1.font.color.rgb = DARK_BLUE
            r2 = p.add_run(rest)
            set_font(r2)
            r2.font.size = Pt(10.5)
        else:
            r = p.add_run(text)
            set_font(r)
            r.font.size = Pt(10.5)


def main() -> None:
    markdown = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    style_document(doc)
    add_title(doc)
    add_content(doc, markdown)
    add_footer(doc)
    doc.save(OUTPUT)


if __name__ == "__main__":
    main()
