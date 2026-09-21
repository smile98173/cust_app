"""Generate a standalone review page for the August regression scenarios."""

from __future__ import annotations

import csv
import html
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent.parent
DEFAULT_WORKBOOK = ROOT / "outputs" / "20260903_august_feedback_regression" / "八月客服回饋回歸測試題庫.xlsx"
DEFAULT_FEEDBACK_DIR = Path(r"D:\AI智能客服\cust_app_runtime\feedback")
DEFAULT_OUTPUT = ROOT / "outputs" / "20260903_august_feedback_regression" / "八月回歸題庫_泡泡審閱.html"
DEFAULT_RESULTS = Path(
    r"C:\Users\user\.codex\visualizations\2026\09\03\01a06687-96dc-7673-80cd-970b89bc4c0b"
    r"\spreadsheet_work\august_validation_ai_test_after_adjustments.json"
)


def parse_json(value: str, fallback: Any) -> Any:
    try:
        return json.loads(value) if value else fallback
    except (TypeError, json.JSONDecodeError):
        return fallback


def feedback_file_for(feedback_id: str, feedback_dir: Path) -> Path | None:
    match = re.match(r"FB-(\d{4})(\d{2})(\d{2})", feedback_id or "")
    if not match:
        return None
    return feedback_dir / f"ai_feedback_{match.group(1)}-{match.group(2)}-{match.group(3)}.csv"


def load_feedback_record(feedback_id: str, feedback_dir: Path, cache: dict[Path, dict[str, dict[str, str]]]) -> dict[str, str]:
    path = feedback_file_for(feedback_id, feedback_dir)
    if not path or not path.exists():
        return {"feedback_id": feedback_id, "load_error": "找不到來源回饋檔"}
    if path not in cache:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            cache[path] = {row.get("feedback_id", ""): row for row in csv.DictReader(handle)}
    return cache[path].get(feedback_id, {"feedback_id": feedback_id, "load_error": "找不到來源回饋紀錄"})


def conversation_from(record: dict[str, str]) -> list[dict[str, str]]:
    conversation = parse_json(record.get("conversation_json", ""), [])
    if isinstance(conversation, list):
        messages = [
            {"role": str(item.get("role") or ""), "content": str(item.get("content") or "")}
            for item in conversation
            if isinstance(item, dict) and item.get("content")
        ]
        if messages:
            return messages
    messages = []
    if record.get("user_message"):
        messages.append({"role": "user", "content": record["user_message"]})
    if record.get("ai_response"):
        messages.append({"role": "assistant", "content": record["ai_response"]})
    return messages


def script_messages(script: str) -> list[dict[str, str]]:
    messages = []
    for line in str(script or "").splitlines():
        text = line.strip()
        if not text:
            continue
        text = re.sub(r"^U\d+\s*", "", text)
        messages.append({"role": "user", "content": text})
    return messages


def load_test_results(path: Path) -> dict[str, dict[str, Any]]:
    if not path.exists():
        return {}
    results = parse_json(path.read_text(encoding="utf-8"), [])
    if not isinstance(results, list):
        return {}
    return {
        f"{item.get('caseId')}:{item.get('feedbackId')}": item
        for item in results
        if isinstance(item, dict)
    }


def test_messages(result: dict[str, Any]) -> list[dict[str, str]]:
    messages = []
    turns = result.get("turns") or []
    responses = result.get("responses") or []
    for index, user_text in enumerate(turns):
        messages.append({"role": "user", "content": str(user_text or "")})
        if index < len(responses):
            messages.append({"role": "assistant", "content": str(responses[index].get("reply") or "")})
    return [message for message in messages if message["content"]]


def load_cases(workbook_path: Path, feedback_dir: Path, result_path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook["回歸題庫"]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]
    feedback_cache: dict[Path, dict[str, dict[str, str]]] = {}
    test_results = load_test_results(result_path)
    cases = []

    for values in rows:
        row = {headers[index]: values[index] if index < len(values) else "" for index in range(len(headers))}
        case_id = str(row.get("題號") or "").strip()
        if not case_id:
            continue
        source_ids = [value.strip() for value in str(row.get("來源回饋ID") or "").splitlines() if value.strip()]
        sources = []
        for feedback_id in source_ids:
            record = load_feedback_record(feedback_id, feedback_dir, feedback_cache)
            test_result = test_results.get(f"{case_id}:{feedback_id}", {})
            sources.append({
                "feedbackId": feedback_id,
                "company": record.get("company") or record.get("company_code") or row.get("公司代碼") or "",
                "createdAt": record.get("created_at") or "",
                "feedbackType": record.get("feedback_type") or "",
                "suggestion": record.get("suggestion") or "",
                "conversation": conversation_from(record),
                "loadError": record.get("load_error") or "",
                "testSignal": (test_result.get("signal") or {}).get("signal") or "未執行",
                "testDuration": test_result.get("durationSec"),
                "testError": test_result.get("error") or "",
                "testConversation": test_messages(test_result),
            })
        cases.append({
            "id": case_id,
            "suite": str(row.get("測試套件") or ""),
            "priority": str(row.get("優先級") or ""),
            "group": str(row.get("功能群組") or ""),
            "company": str(row.get("公司代碼") or ""),
            "kind": str(row.get("題型") or ""),
            "script": script_messages(str(row.get("對話腳本（依序輸入）") or "")),
            "expected": str(row.get("期望AI行為") or ""),
            "must": str(row.get("必須包含") or ""),
            "mustNot": str(row.get("禁止出現/禁止行為") or ""),
            "dependency": str(row.get("外部依賴") or ""),
            "focus": str(row.get("判定重點") or ""),
            "sources": sources,
        })
    return cases


def page_html(cases: list[dict[str, Any]], workbook_path: Path) -> str:
    payload = json.dumps(cases, ensure_ascii=False).replace("</", "<\\/")
    return f"""<!doctype html>
<html lang="zh-Hant">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>八月回歸題庫泡泡審閱</title>
<style>
:root {{ --page:#f5f7fa; --panel:#fff; --ink:#172033; --muted:#62708a; --line:#d7dfeb; --blue:#1769aa; --blue-soft:#e9f4ff; --orange:#b45309; --orange-soft:#fff1df; --green:#157347; --red:#b42318; --shadow:0 8px 24px rgba(26,43,72,.08); }}
* {{ box-sizing:border-box; }} body {{ margin:0; background:var(--page); color:var(--ink); font-family:"Microsoft JhengHei","Noto Sans TC",Arial,sans-serif; line-height:1.55; }}
header {{ height:74px; display:flex; align-items:center; justify-content:space-between; gap:20px; padding:0 28px; background:#fff; border-bottom:1px solid var(--line); }}
h1 {{ margin:0; font-size:21px; }} .header-sub {{ color:var(--muted); font-size:13px; margin-top:3px; }} .progress {{ color:var(--muted); font-size:13px; white-space:nowrap; }}
.app {{ display:grid; grid-template-columns:310px minmax(0,1fr); min-height:calc(100vh - 74px); }} aside {{ border-right:1px solid var(--line); background:#fff; padding:18px; }}
.filters {{ display:grid; gap:10px; }} input, select, textarea {{ width:100%; font:inherit; border:1px solid var(--line); border-radius:6px; background:#fff; color:var(--ink); padding:9px 10px; }} textarea {{ min-height:82px; resize:vertical; }}
.case-list {{ margin-top:16px; display:grid; gap:6px; overflow:auto; max-height:calc(100vh - 220px); padding-right:3px; }} .case-button {{ text-align:left; border:1px solid transparent; background:#fff; border-radius:6px; padding:10px; cursor:pointer; }} .case-button:hover {{ background:#f4f9ff; }} .case-button.active {{ background:var(--blue-soft); border-color:#9ac7ef; }} .case-id {{ font-size:12px; font-weight:800; color:var(--blue); }} .case-name {{ font-size:14px; font-weight:700; margin-top:2px; }} .case-meta {{ color:var(--muted); font-size:12px; margin-top:3px; }}
main {{ padding:24px 30px 48px; max-width:1440px; width:100%; margin:0 auto; }} .case-head {{ display:flex; justify-content:space-between; gap:18px; align-items:flex-start; }} .eyebrow {{ color:var(--blue); font-size:13px; font-weight:800; }} h2 {{ margin:3px 0 0; font-size:25px; }} .tags {{ display:flex; flex-wrap:wrap; gap:6px; justify-content:flex-end; }} .tag {{ border:1px solid var(--line); border-radius:999px; padding:4px 9px; color:var(--muted); font-size:12px; background:#fff; }} .tag.p0 {{ color:var(--red); border-color:#f3b5b1; background:#fff5f4; }}
.expectation {{ margin:18px 0; border-left:4px solid var(--blue); background:#edf6ff; padding:12px 15px; }} .expectation strong {{ display:block; margin-bottom:3px; }}
.rules {{ display:grid; grid-template-columns:1fr 1fr; gap:12px; margin:14px 0 20px; }} .rule {{ border:1px solid var(--line); background:#fff; padding:11px 13px; }} .rule-title {{ font-size:12px; font-weight:800; color:var(--muted); margin-bottom:4px; }} .rule.must .rule-title {{ color:var(--green); }} .rule.ban .rule-title {{ color:var(--red); }}
.review {{ display:grid; grid-template-columns:1fr 1fr; gap:18px; align-items:start; }} .column-head {{ display:flex; align-items:center; justify-content:space-between; gap:12px; margin-bottom:9px; }} h3 {{ margin:0; font-size:16px; }} .turn-note {{ color:var(--muted); font-size:12px; }} .chat {{ background:#f9fbfd; border:1px solid var(--line); min-height:360px; padding:16px; }} .bubble-row {{ display:flex; margin:11px 0; }} .bubble-row.user {{ justify-content:flex-end; }} .bubble {{ width:fit-content; max-width:82%; border:1px solid; border-radius:8px; padding:9px 11px; white-space:pre-wrap; }} .bubble-row.user .bubble {{ background:var(--blue-soft); border-color:#9ac7ef; }} .bubble-row.assistant .bubble {{ background:var(--orange-soft); border-color:#f1bf82; }} .role {{ display:block; font-size:11px; font-weight:800; color:var(--muted); margin-bottom:3px; }}
.source-meta {{ color:var(--muted); font-size:12px; margin:-2px 0 10px; }} .source-note {{ background:#fff7ed; border-left:3px solid #f59e0b; padding:9px 11px; font-size:13px; margin-bottom:12px; }} .run-head {{ margin:18px 0 8px; display:flex; align-items:center; justify-content:space-between; gap:10px; }} .run-status {{ border:1px solid var(--line); border-radius:999px; font-size:12px; padding:3px 8px; color:var(--muted); background:#fff; }} .empty {{ color:var(--muted); padding:20px; text-align:center; }}
.decision {{ margin-top:24px; border-top:1px solid var(--line); padding-top:18px; display:grid; grid-template-columns:210px minmax(0,1fr); gap:12px; }} .decision label {{ font-size:13px; font-weight:800; }} .save-note {{ color:var(--muted); font-size:12px; grid-column:2; }}
@media (max-width:960px) {{ .app {{ grid-template-columns:1fr; }} aside {{ border-right:0; border-bottom:1px solid var(--line); }} .case-list {{ max-height:210px; }} main {{ padding:20px 16px; }} .review,.rules,.decision {{ grid-template-columns:1fr; }} .save-note {{ grid-column:auto; }} .case-head {{ display:block; }} .tags {{ justify-content:flex-start; margin-top:12px; }} }}
</style>
</head>
<body>
<header><div><h1>八月回歸題庫泡泡審閱</h1><div class="header-sub">測試輸入與原始客服對話並列，先確認題目是否代表原需求。</div></div><div class="progress" id="progress"></div></header>
<div class="app"><aside><div class="filters"><input id="search" placeholder="搜尋題號、群組或關鍵字"><select id="priority"><option value="">全部優先級</option><option value="P0">P0</option><option value="P1">P1</option><option value="P2">P2</option></select></div><div class="case-list" id="caseList"></div></aside>
<main id="detail"></main></div>
<script>
const cases = {payload};
const esc = value => String(value || '').replace(/[&<>"']/g, char => ({{'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}}[char]));
const bubbles = messages => messages?.length ? messages.map(item => `<div class="bubble-row ${{item.role === 'user' ? 'user' : 'assistant'}}"><div class="bubble"><span class="role">${{item.role === 'user' ? '測試使用者' : '原始 AI 回覆'}}</span>${{esc(item.content)}}</div></div>`).join('') : '<div class="empty">沒有可顯示的對話內容。</div>';
let visible = cases, selectedId = cases[0]?.id || '', sourceIndex = 0;
const reviewKey = id => `august-regression-review-${{id}}`;
function renderList() {{
  const search = document.querySelector('#search').value.trim().toLowerCase();
  const priority = document.querySelector('#priority').value;
  visible = cases.filter(item => (!priority || item.priority === priority) && (!search || JSON.stringify(item).toLowerCase().includes(search)));
  if (!visible.some(item => item.id === selectedId)) {{ selectedId = visible[0]?.id || ''; sourceIndex = 0; }}
  document.querySelector('#caseList').innerHTML = visible.map(item => `<button class="case-button ${{item.id === selectedId ? 'active' : ''}}" data-id="${{esc(item.id)}}"><div class="case-id">${{esc(item.id)}} · ${{esc(item.priority)}}</div><div class="case-name">${{esc(item.group)}}</div><div class="case-meta">${{esc(item.kind)}} · ${{item.sources.length}} 筆來源</div></button>`).join('') || '<div class="empty">找不到符合的題目。</div>';
  document.querySelectorAll('.case-button').forEach(button => button.onclick = () => {{ selectedId = button.dataset.id; sourceIndex = 0; render(); }});
}}
function renderDetail() {{
  const item = cases.find(value => value.id === selectedId);
  if (!item) {{ document.querySelector('#detail').innerHTML = '<div class="empty">請從左側選擇題目。</div>'; return; }}
  const source = item.sources[sourceIndex] || item.sources[0] || {{}};
  const saved = JSON.parse(localStorage.getItem(reviewKey(item.id)) || '{{}}');
  const sourceOptions = item.sources.map((value, index) => `<option value="${{index}}" ${{index === sourceIndex ? 'selected' : ''}}>${{esc(value.feedbackId)}}${{value.company ? ' · ' + esc(value.company) : ''}}</option>`).join('');
  document.querySelector('#detail').innerHTML = `
    <div class="case-head"><div><div class="eyebrow">${{esc(item.id)}} · ${{esc(item.suite)}}</div><h2>${{esc(item.group)}}</h2></div><div class="tags"><span class="tag ${{item.priority.toLowerCase()}}">${{esc(item.priority)}}</span><span class="tag">${{esc(item.company)}}</span><span class="tag">${{esc(item.kind)}}</span>${{item.dependency ? `<span class="tag">${{esc(item.dependency)}}</span>` : ''}}</div></div>
    <section class="expectation"><strong>期望 AI 行為</strong>${{esc(item.expected)}}</section>
    <section class="rules"><div class="rule must"><div class="rule-title">必須包含</div>${{esc(item.must || '未指定')}}</div><div class="rule ban"><div class="rule-title">禁止出現／禁止行為</div>${{esc(item.mustNot || '未指定')}}</div><div class="rule"><div class="rule-title">判定重點</div>${{esc(item.focus || '未指定')}}</div><div class="rule"><div class="rule-title">外部依賴</div>${{esc(item.dependency || '無')}}</div></section>
    <section class="review"><div><div class="column-head"><h3>題庫測試情境</h3><span class="turn-note">依序輸入</span></div><div class="chat">${{bubbles(item.script)}}</div><div class="run-head"><h3>修改後 AI 實測對話</h3><span class="run-status">${{esc(source.testSignal)}}${{source.testDuration !== null && source.testDuration !== undefined ? ' · ' + esc(source.testDuration) + ' 秒' : ''}}</span></div>${{source.testError ? `<div class="source-note">測試錯誤：${{esc(source.testError)}}</div>` : ''}}<div class="chat">${{bubbles(source.testConversation)}}</div></div><div><div class="column-head"><h3>來源客服對話</h3><select id="sourceSelect">${{sourceOptions}}</select></div><div class="source-meta">${{esc(source.createdAt)}} · ${{esc(source.feedbackType || '未分類')}}</div>${{source.suggestion ? `<div class="source-note"><strong>原始回饋：</strong>${{esc(source.suggestion)}}</div>` : ''}}${{source.loadError ? `<div class="source-note">${{esc(source.loadError)}}</div>` : ''}}<div class="chat">${{bubbles(source.conversation)}}</div></div></section>
    <section class="decision"><label for="verdict">你的確認結果</label><select id="verdict"><option value="">尚未確認</option><option value="correct">符合原意</option><option value="adjust">需調整題目</option></select><label for="note">調整備註</label><textarea id="note" placeholder="例如：這題漏掉第二輪條件、原始抱怨其實是帳務問題">${{esc(saved.note || '')}}</textarea><div></div><div class="save-note">確認結果會儲存在此瀏覽器，不會寫回原始題庫。<span id="saved"></span></div></section>`;
  const verdict = document.querySelector('#verdict'); verdict.value = saved.verdict || '';
  document.querySelector('#sourceSelect').onchange = event => {{ sourceIndex = Number(event.target.value); renderDetail(); }};
  const save = () => {{ localStorage.setItem(reviewKey(item.id), JSON.stringify({{verdict: verdict.value, note: document.querySelector('#note').value}})); document.querySelector('#saved').textContent = ' 已儲存。'; renderProgress(); }};
  verdict.onchange = save; document.querySelector('#note').onchange = save;
}}
function renderProgress() {{ const reviewed = cases.filter(item => JSON.parse(localStorage.getItem(reviewKey(item.id)) || '{{}}').verdict).length; document.querySelector('#progress').textContent = `已確認 ${{reviewed}} / ${{cases.length}} 題`; }}
function render() {{ renderList(); renderDetail(); renderProgress(); }}
document.querySelector('#search').oninput = render; document.querySelector('#priority').onchange = render; render();
</script></body></html>"""


def main() -> None:
    cases = load_cases(DEFAULT_WORKBOOK, DEFAULT_FEEDBACK_DIR, DEFAULT_RESULTS)
    DEFAULT_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    DEFAULT_OUTPUT.write_text(page_html(cases, DEFAULT_WORKBOOK), encoding="utf-8")
    print(f"Generated {DEFAULT_OUTPUT} with {len(cases)} cases.")


if __name__ == "__main__":
    main()
