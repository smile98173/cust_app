import csv
import hashlib
import json
import re
import shutil
import time
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from app.config.settings import (
    BASE_DIR,
    RAG_CHUNK_OVERLAP,
    RAG_CHUNK_SIZE,
    RAG_LOCAL_COLLECTION,
    RAG_LOCAL_DOCS_DIR,
    RAG_LOCAL_EMBED_DEVICE,
    RAG_LOCAL_EMBED_MODEL,
    RAG_LOCAL_MANIFEST_PATH,
    RAG_LOCAL_PERSIST_DIR,
)
from app.services.knowledge_base_policy import (
    CENTRAL_COMMON_KNOWLEDGE_BASE,
    LEGACY_COMMON_KNOWLEDGE_BASE,
    normalize_knowledge_base_name,
)
from app.services.error_logging import get_error_logger


ALLOWED_EXTENSIONS = {"txt", "md", "html", "htm", "csv", "jsonl", "pdf", "docx"}


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _index_failure_detail(record: Dict[str, Any]) -> Dict[str, Any]:
    return {
        "document_id": str(record.get("id") or ""),
        "title": str(record.get("title") or ""),
        "file_name": str(record.get("file_name") or ""),
        "knowledge_base": str(record.get("knowledge_base") or ""),
        "stage": str(record.get("processing_error_stage") or "unknown"),
        "exception_type": str(record.get("processing_error_type") or ""),
        "error": str(record.get("processing_error") or "未提供錯誤訊息"),
    }


def _report_reindex_failures(failures: List[Dict[str, Any]]) -> None:
    if not failures:
        return

    print(f"[KB] Reindex failed documents ({len(failures)}):")
    for failure in failures:
        print(
            "[KB] Reindex document failed: "
            + json.dumps(failure, ensure_ascii=False, sort_keys=True)
        )

    record = {
        "timestamp": datetime.now().isoformat(),
        "level": "ERROR",
        "operation": "kb_reindex_documents_failed",
        "failed_count": len(failures),
        "failures": failures,
    }
    try:
        get_error_logger().error(json.dumps(record, ensure_ascii=False))
    except Exception:
        pass


def normalize_actor_name(value: Optional[str]) -> str:
    actor = str(value or "").strip()
    return actor or "system"


def sanitize_filename(name: str) -> str:
    raw = Path(name or "document.txt").name
    cleaned = re.sub(r"[^\w.\-\u4e00-\u9fff]+", "_", raw, flags=re.UNICODE).strip("._")
    return cleaned or "document.txt"


def normalize_text(text: str) -> str:
    text = str(text or "").replace("\ufeff", "")
    text = re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F\x7F]", "", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def normalize_preserve_lines(text: str) -> str:
    text = str(text or "").replace("\ufeff", "")
    text = re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F\x7F]", "", text)
    lines = [re.sub(r"[ \t　]+", " ", line).strip() for line in text.splitlines()]
    lines = [line for line in lines if line]
    return "\n".join(lines)


def read_text_with_fallback(path: Path) -> str:
    for encoding in ["utf-8-sig", "utf-8", "cp950", "big5"]:
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="ignore")


def extract_html_text(path: Path) -> str:
    raw = read_text_with_fallback(path)
    try:
        from bs4 import BeautifulSoup

        return BeautifulSoup(raw, "html.parser").get_text("\n")
    except Exception:
        return re.sub(r"<[^>]+>", " ", raw)


def extract_pdf_sections(path: Path) -> List[Dict[str, Any]]:
    try:
        from pypdf import PdfReader
    except Exception as exc:
        raise RuntimeError("PDF 匯入需要安裝 pypdf。") from exc

    reader = PdfReader(str(path))
    sections = []
    for index, page in enumerate(reader.pages, start=1):
        content = extract_pdf_page_text(page)
        if content:
            sections.append({
                "content": content,
                "page_no": index,
                "section": f"page {index}",
            })
    return sections


def extract_pdf_page_text(page: Any) -> str:
    try:
        text = page.extract_text(
            extraction_mode="layout",
            layout_mode_space_vertically=False,
            layout_mode_strip_rotated=False,
        )
    except TypeError:
        text = page.extract_text()
    except Exception:
        text = page.extract_text()

    text = normalize_preserve_lines(text or "")
    if text:
        return text
    return normalize_preserve_lines(page.extract_text() or "")


def extract_docx_text(path: Path) -> str:
    try:
        from docx import Document
    except Exception as exc:
        raise RuntimeError("DOCX 匯入需要安裝 python-docx。") from exc

    document = Document(str(path))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def normalize_docx_table_cells(cells: List[str]) -> List[str]:
    normalized: List[str] = []
    for cell in cells:
        value = normalize_text(cell)
        if not value:
            continue
        if normalized and normalized[-1] == value:
            continue
        normalized.append(value)
    return normalized


def format_docx_table_row(cells: List[str]) -> str:
    return " | ".join(normalize_docx_table_cells(cells))


def detect_docx_table_group_label(content: str) -> str:
    normalized = normalize_text(content)
    compact = normalized.replace(" ", "").replace("　", "").replace("-", "")
    match = re.search(r"哈TV([ABC])套餐", compact, flags=re.IGNORECASE)
    if match:
        return f"哈TV-{match.group(1).upper()}套餐"
    return ""


def build_docx_table_sections(
    table_rows: List[Tuple[int, List[str]]],
    table_index: int,
) -> List[Dict[str, Any]]:
    formatted_rows = [
        (row_index, format_docx_table_row(cells))
        for row_index, cells in table_rows
    ]
    formatted_rows = [(row_index, content) for row_index, content in formatted_rows if content]
    if not formatted_rows:
        return []

    group_starts = [
        (position, row_index, detect_docx_table_group_label(content))
        for position, (row_index, content) in enumerate(formatted_rows)
        if detect_docx_table_group_label(content)
    ]

    groups: List[Tuple[int, int, str]] = []
    if len(group_starts) >= 2:
        for group_position, (start_position, start_row, label) in enumerate(group_starts):
            next_position = (
                group_starts[group_position + 1][0]
                if group_position + 1 < len(group_starts)
                else len(formatted_rows)
            )
            groups.append((start_position, next_position, label))
    else:
        groups.append((0, len(formatted_rows), ""))

    sections: List[Dict[str, Any]] = []
    for start_position, end_position, label in groups:
        rows = formatted_rows[start_position:end_position]
        if not rows:
            continue
        first_row = rows[0][0]
        last_row = rows[-1][0]
        content = normalize_preserve_lines("\n".join(content for _, content in rows))
        if not content:
            continue
        section_label = f"table {table_index} rows {first_row}-{last_row}"
        if label:
            section_label = f"table {table_index} {label} rows {first_row}-{last_row}"
        sections.append({
            "content": content,
            "page_no": None,
            "section": section_label,
            "section_type": "docx_table",
        })
    return sections


def extract_docx_sections(path: Path) -> List[Dict[str, Any]]:
    try:
        from docx import Document
    except Exception as exc:
        raise RuntimeError("DOCX 匯入需要安裝 python-docx。") from exc

    document = Document(str(path))
    sections: List[Dict[str, Any]] = []
    buffer: List[str] = []
    start_index = 1

    def flush(end_index: int) -> None:
        nonlocal buffer, start_index
        content = normalize_preserve_lines("\n".join(buffer))
        if content:
            label = f"paragraph {start_index}"
            if end_index > start_index:
                label = f"paragraphs {start_index}-{end_index}"
            sections.append({
                "content": content,
                "page_no": None,
                "section": label,
                "section_type": "docx",
            })
        buffer = []

    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = normalize_preserve_lines(paragraph.text)
        if not text:
            flush(index - 1)
            start_index = index + 1
            continue
        if not buffer:
            start_index = index
        buffer.append(text)

    flush(len(document.paragraphs))

    table_index = 0
    for table in document.tables:
        table_index += 1
        table_rows = [
            (row_index, [cell.text for cell in row.cells])
            for row_index, row in enumerate(table.rows, start=1)
        ]
        sections.extend(build_docx_table_sections(table_rows, table_index))

    if not sections:
        text = normalize_preserve_lines(extract_docx_text(path))
        if text:
            sections.append({
                "content": text,
                "page_no": None,
                "section": "document",
                "section_type": "docx",
            })
    return sections


def extract_csv_sections(path: Path) -> List[Dict[str, Any]]:
    raw = read_text_with_fallback(path)
    rows = list(csv.DictReader(raw.splitlines()))
    if not rows:
        return [{"content": normalize_text(raw), "page_no": None, "section": None}]

    def first_value(row: Dict[str, Any], *keys: str) -> str:
        normalized_row = {
            str(key or "").strip().lower(): value
            for key, value in row.items()
        }
        for key in keys:
            value = normalized_row.get(key.lower())
            if value:
                return str(value).strip()
        return ""

    sections = []
    for index, row in enumerate(rows, start=1):
        question = first_value(row, "question", "問題", "q")
        answer = first_value(row, "answer", "答案", "a")
        company = first_value(row, "company", "公司", "knowledge_base", "知識庫")
        if question or answer:
            content = normalize_text(" ".join(part for part in [question, answer, company] if part))
        else:
            parts = [f"{key}: {value}" for key, value in row.items() if value]
            content = normalize_text("；".join(parts))
        if content:
            sections.append({
                "content": content,
                "question": question,
                "answer": answer,
                "company": company,
                "page_no": None,
                "section": f"row {index}",
            })
    return sections


def extract_jsonl_sections(path: Path) -> List[Dict[str, Any]]:
    sections = []
    for index, line in enumerate(read_text_with_fallback(path).splitlines(), start=1):
        line = line.strip()
        if not line:
            continue
        try:
            item = json.loads(line)
        except json.JSONDecodeError:
            item = {"content": line}

        if isinstance(item, dict):
            title = item.get("title") or item.get("question") or item.get("name") or ""
            content = item.get("content") or item.get("answer") or item.get("summary") or ""
            text = normalize_text(f"{title}\n{content}")
        else:
            text = normalize_text(str(item))

        if text:
            sections.append({
                "content": text,
                "page_no": None,
                "section": f"line {index}",
            })
    return sections


def extract_document_sections(path: Path) -> List[Dict[str, Any]]:
    extension = path.suffix.lower().lstrip(".")

    if extension == "pdf":
        return extract_pdf_sections(path)
    if extension == "docx":
        return extract_docx_sections(path)
    if extension in {"html", "htm"}:
        return [{
            "content": normalize_text(extract_html_text(path)),
            "page_no": None,
            "section": None,
        }]
    if extension == "csv":
        return extract_csv_sections(path)
    if extension == "jsonl":
        return extract_jsonl_sections(path)

    return [{
        "content": normalize_preserve_lines(read_text_with_fallback(path)),
        "page_no": None,
        "section": "document",
        "section_type": "text",
    }]


def chunk_text(
    text: str,
    size: int = RAG_CHUNK_SIZE,
    overlap: int = RAG_CHUNK_OVERLAP,
    preserve_lines: bool = False,
) -> List[str]:
    text = normalize_preserve_lines(text) if preserve_lines else normalize_text(text)
    if not text:
        return []

    size = max(200, int(size or 1200))
    overlap = max(0, min(int(overlap or 0), size - 1))
    chunks = []
    start = 0
    while start < len(text):
        chunk = text[start:start + size].strip()
        if chunk:
            chunks.append(normalize_preserve_lines(chunk) if preserve_lines else chunk)
        next_start = start + size - overlap
        if next_start <= start:
            next_start = start + size
        start = next_start
    return chunks


PDF_BLOCK_START_RE = re.compile(
    r"^(?:"
    r"[一二三四五六七八九十]+、"
    r"|\d+[.、]"
    r"|方案[一二三四五六七八九十\d]+"
    r"|(?:\d+M|1G)\s"
    r"|[A-Za-z0-9]+M/[A-Za-z0-9]+M"
    r"|(?:方案名稱|活動期間|裝機費|寬頻設備押金|機上盒借用|用戶預設類別|繳別|綁約|違約金|舊戶是否可參加|申裝條件|中途換約或升級|合約期間可加價|加值服務|設備名稱|本方案|客新裝|客約|用戶若需)[:：]"
    r"|合約期間可加價購租借$"
    r"|加值服務[（(]"
    r"|下述頻寬不推"
    r"|.+(?:說明|注意事項|拆帳表|備註範例)[:：]?$"
    r")"
)

PDF_SEPARATOR_RE = re.compile(r"^[\-─━_]{8,}$")
PDF_CAMPAIGN_FORM_RE = re.compile(r"(方案名稱\s*[:：].*活動期間\s*[:：]|飆網守護家.*拆帳表)", re.DOTALL)
PDF_CAMPAIGN_PLAN_PAGE_RE = re.compile(r"方案[一二三四五六七八九十]、")
PDF_CAMPAIGN_PLAN_START_RE = re.compile(r"^方案[一二三四五六七八九十]、")
PDF_CAMPAIGN_PLAN_ITEM_CONTINUATION_RE = re.compile(r"^[2-9]、")
PDF_DISPATCH_NOTE_CONTINUATION_RE = re.compile(r"^(?:價[、，]|LITV|LINETV|收費|完工|簽合約)")
PDF_RATE_TABLE_HEADING_RE = re.compile(r".*拆帳表$")
PDF_COLON_PLAN_START_RE = re.compile(r"^方案名稱\s*[:：]", re.MULTILINE)


DOCX_BLOCK_START_RE = re.compile(
    r"^(?:"
    r"[一二三四五六七八九十]+、(?=\s*(?:TV|收視|移機|分機|加值|設備|其他|優惠|申請|限制|方案|費用|裝機|行政|重要|規範))"
    r"|【方案名稱】"
    r"|【(?:適用對象|申請方式|申請條件|優惠內容|限制條件|備註|注意事項)】"
    r"|(?:項目|方案|費用|收費|裝機費|月租|設備|優惠|限制|申請).{0,12}[:：]"
    r")"
)

DOCX_HEADING_RE = re.compile(r"(?:【方案名稱】|方案名稱[:：])\s*([^【\n]+)")
DOCX_PLAN_START_RE = re.compile(r"^(?:【方案名稱】|方案名稱\s*[:：])")
CAMPAIGN_HEADER_RE = re.compile(r"^(?:【方案名稱】|方案名稱\s*[:：])", re.MULTILINE)
PLAN_NAME_RE = re.compile(r"(?:【方案名稱】|方案名稱\s*[:：])\s*([^【\n]+)")
PLAN_BLOCK_LABEL_RE = re.compile(r"^(方案[一二三四五六七八九十\d]+[、:：].+)$", re.MULTILINE)

CAMPAIGN_ACTIVITY_TERMS = (
    "活動期間", "方案一", "方案二", "方案三", "方案四", "綁約", "違約金",
    "贈哈POINTS", "贈LINE TV", "贈家電", "售價", "抽獎", "摸彩", "獎項",
)
SOCIAL_DISCOUNT_TERMS = ("低收入", "中低收入", "身心障礙", "優惠戶")
FEE_TERMS = ("費用", "月繳", "月租", "半年繳", "年繳", "售價", "收費", "裝機費", "押金")
GIFT_TERMS = ("贈", "贈品", "家電", "POINTS", "LINE TV", "LINETV", "LITV", "小冰箱", "投影機", "電視", "大冰箱")
GIFT_DETAIL_TERMS = ("贈", "贈品", "贈送", "加贈", "好禮", "免費送", "二擇一")
LOTTERY_TERMS = (
    "抽獎", "摸彩", "抽出", "中獎", "得獎", "抽獎資格", "抽獎期間",
    "抽獎日期", "抽獎券", "獎項", "獎品",
)
CAMPAIGN_OCCASION_ALIASES = {
    "春節": ("春節", "過年", "新春", "農曆年"),
    "元宵節": ("元宵節", "元宵", "燈會"),
    "清明節": ("清明節", "清明"),
    "兒童節": ("兒童節", "兒童月"),
    "勞動節": ("勞動節", "五一勞動節"),
    "母親節": ("母親節", "媽媽節", "媽咪", "母親", "媽媽"),
    "端午節": ("端午節", "端午", "粽夏", "粽情"),
    "父親節": ("父親節", "爸爸節", "爸氣", "父親", "爸爸"),
    "七夕": ("七夕", "七夕情人節"),
    "情人節": ("情人節", "西洋情人節"),
    "中秋節": ("中秋節", "中秋", "月圓"),
    "國慶日": ("國慶日", "國慶", "雙十"),
    "聖誕節": ("聖誕節", "聖誕", "耶誕節", "耶誕"),
    "跨年": ("跨年", "迎新年"),
    "開學季": ("開學季", "開學"),
    "暑期": ("暑期", "暑假", "夏日", "盛夏"),
    "週年慶": ("週年慶", "周年慶"),
    "年終": ("年終", "歲末", "尾牙"),
}
ELIGIBILITY_TERMS = ("舊戶", "可參加", "申請條件", "適用對象", "限制條件", "資格", "申裝條件")
CONTRACT_TERMS = ("活動期間", "綁約", "合約", "違約金", "到期", "中途換約", "升級")
WIFI_TERMS = ("wifi", "wi-fi", "WIFI", "WiFi", "租借 WIFI", "租借WiFi", "月租50")
EQUIPMENT_TERMS = ("設備", "數據機", "機上盒", "聯網機上盒", "押金", "賠償")
CAMPAIGN_SERVICE_TERMS = {
    "寬頻": ("網路", "寬頻", "光纖", "wifi", "wi-fi", "數據機", "100m", "300m", "500m", "1g"),
    "有線電視": ("有線電視", "第四台", "收視", "頻道", "機上盒", "catv"),
    "電視網路同裝": ("電視+網路", "電視＋網路", "電視網路同裝", "有線電視+網路", "好視成雙"),
    "加值服務": ("加值服務", "linetv", "line tv", "litv", "hbo", "wifi租借", "智慧攝影機"),
}
CAMPAIGN_CUSTOMER_TERMS = (
    "新戶", "新裝", "舊戶", "續約", "換約", "升級", "低收入戶",
    "中低收入戶", "身心障礙", "一般戶", "無合約",
)
CAMPAIGN_PAYMENT_TERMS = ("月繳", "季繳", "半年繳", "年繳", "循環扣款", "信用卡")
CAMPAIGN_SECTION_LABELS = {
    "price": FEE_TERMS,
    "gift": GIFT_TERMS,
    "lottery": LOTTERY_TERMS,
    "eligibility": ELIGIBILITY_TERMS,
    "contract": CONTRACT_TERMS,
    "equipment": EQUIPMENT_TERMS,
    "wifi": WIFI_TERMS,
}
CAMPAIGN_SPEED_RE = re.compile(
    r"(?<![A-Za-z0-9])(\d+(?:\.\d+)?\s*[MG])\s*/\s*(\d+(?:\.\d+)?\s*[MG])(?![A-Za-z0-9])",
    re.IGNORECASE,
)
CAMPAIGN_CONTRACT_MONTH_RE = re.compile(
    r"(?:綁約(?:期限)?|合約(?:期間)?|約期|租期)"
    r"[^。\n；;]{0,20}?(\d{1,3})\s*個?月",
    re.IGNORECASE,
)
CAMPAIGN_PERIOD_RE = re.compile(
    r"活動期間\s*[:：]\s*([^\n。；;]+)",
    re.IGNORECASE,
)
CAMPAIGN_TITLE_TERMS = (
    "優惠", "活動", "專案", "方案", "好視成雙", "飆網守護",
    "哈NET", "低收入", "中低收入", "身心障礙", "獻禮",
)
CAMPAIGN_CHINESE_CONTRACT_MONTHS = {
    "半年": 6,
    "一年": 12,
    "一年度": 12,
    "兩年": 24,
    "二年": 24,
    "三年": 36,
    "四年": 48,
}

VALUE_ADDED_CATALOG_TITLE_TERMS = (
    "各項單品銷售", "單品銷售", "加值服務", "加值產品", "數位加值",
)
VALUE_ADDED_PRODUCT_DEFINITIONS = (
    {
        "key": "line_tv",
        "name": "LINE TV",
        "aliases": ("LINE TV", "LINETV", "LITV", "LINE TV套餐", "影音服務"),
    },
    {
        "key": "wifi_5",
        "name": "WiFi 5 系列分享器",
        "aliases": ("WIFI 5系列分享器", "WIFI5", "WI-FI 5", "WiFi 5分享器"),
    },
    {
        "key": "wifi_6",
        "name": "WiFi 6 系列分享器",
        "aliases": ("WIFI 6系列分享器", "WIFI6", "WI-FI 6", "WiFi 6分享器"),
    },
    {
        "key": "wifi_addon",
        "name": "WiFi 加值服務",
        "aliases": (
            "WIFI加值服務", "WIFI", "WI-FI", "MESH", "分享器", "無線網路設備",
        ),
    },
    {
        "key": "home_camera",
        "name": "居家智慧攝影機",
        "aliases": (
            "居家智慧攝影機", "智慧攝影機", "攝影機", "監視器", "智慧鏡頭",
            "居家監控", "攝影機租借", "租借攝影機",
        ),
    },
    {
        "key": "marpa_user",
        "name": "瑪帛用戶",
        "aliases": ("瑪帛用戶", "瑪帛基本方案"),
    },
    {
        "key": "marpa_friend",
        "name": "瑪帛好友",
        "aliases": ("瑪帛好友", "瑪帛好友方案"),
    },
    {
        "key": "marpa_partner",
        "name": "瑪帛夥伴",
        "aliases": ("瑪帛夥伴", "瑪帛夥伴方案"),
    },
    {
        "key": "bear_care",
        "name": "熊搭心",
        "aliases": ("熊搭心", "熊搭心服務", "電視電話", "家庭相簿", "生活提醒", "瑪帛"),
    },
)
VALUE_ADDED_NUMBERED_HEADING_RE = re.compile(
    r"^(?P<number>[一二三四五六七八九十]+|\d+)\s*[、.．]\s*(?P<body>.+)$"
)
VALUE_ADDED_DETAIL_TERMS = (
    "原價", "特價", "月租", "月繳", "半年繳", "年繳", "費用", "價格",
    "綁約", "租借", "加購", "申請", "申辦", "服務", "功能", "設備",
)


def normalize_product_identity(value: str) -> str:
    return re.sub(r"[\s　_\-－/（）()]+", "", str(value or "")).lower()


def match_value_added_product(value: str) -> Optional[Dict[str, Any]]:
    normalized = normalize_product_identity(value)
    if not normalized:
        return None
    for definition in VALUE_ADDED_PRODUCT_DEFINITIONS:
        for alias in definition["aliases"]:
            normalized_alias = normalize_product_identity(alias)
            if normalized_alias and normalized_alias in normalized:
                return dict(definition)
    return None


def detect_value_added_catalog_document(
    record: Dict[str, Any],
    chunks: List[Dict[str, Any]],
) -> bool:
    title = str(record.get("title") or "")
    file_name = str(record.get("file_name") or "")
    content = "\n".join(str(chunk.get("content") or "") for chunk in chunks)
    identity = "\n".join([title, file_name])
    combined = "\n".join([identity, content])
    if detect_campaign_document(record, chunks):
        return False

    matched_products = {
        definition["key"]
        for definition in VALUE_ADDED_PRODUCT_DEFINITIONS
        if any(
            normalize_product_identity(alias) in normalize_product_identity(combined)
            for alias in definition["aliases"]
        )
    }
    numbered_headings = [
        line for line in normalize_preserve_lines(content).splitlines()
        if VALUE_ADDED_NUMBERED_HEADING_RE.match(line)
    ]
    has_structured_catalog = (
        has_any_term(identity, VALUE_ADDED_CATALOG_TITLE_TERMS)
        and bool(numbered_headings)
        and has_any_term(content, VALUE_ADDED_DETAIL_TERMS)
    )
    return has_structured_catalog or len(matched_products) >= 3


def build_value_added_product_profile(
    record: Dict[str, Any],
    chunks: List[Dict[str, Any]],
) -> Dict[str, Any]:
    if not detect_value_added_catalog_document(record, chunks):
        return {}

    content = normalize_preserve_lines(
        "\n".join(str(chunk.get("content") or "") for chunk in chunks)
    )
    lines = content.splitlines()
    events: List[Dict[str, Any]] = []
    current_parent = ""

    for line_index, line in enumerate(lines):
        match = VALUE_ADDED_NUMBERED_HEADING_RE.match(line)
        if not match:
            continue
        number = match.group("number")
        body = normalize_text(match.group("body"))
        heading = re.split(r"[:：]", body, maxsplit=1)[0].strip()
        definition = match_value_added_product(heading)
        level = "top" if not number.isdigit() else "child"

        if not definition and level == "child" and current_parent not in {"wifi_addon", "bear_care"}:
            continue
        if not definition and level == "top":
            definition = {
                "key": f"item_{line_index + 1}",
                "name": heading,
                "aliases": (heading,),
            }
        if not definition:
            definition = {
                "key": f"{current_parent or 'item'}_{line_index + 1}",
                "name": heading,
                "aliases": (heading,),
            }

        if level == "top":
            current_parent = str(definition["key"])
            parent_key = ""
        else:
            parent_key = current_parent

        events.append({
            "line_index": line_index,
            "level": level,
            "key": str(definition["key"]),
            "name": str(definition["name"]),
            "aliases": unique_normalized_values([
                definition["name"],
                heading,
                *definition.get("aliases", ()),
            ]),
            "parent_key": parent_key,
        })

    products: List[Dict[str, Any]] = []
    for event_index, event in enumerate(events):
        if event["level"] == "top":
            end_index = next(
                (
                    candidate["line_index"]
                    for candidate in events[event_index + 1:]
                    if candidate["level"] == "top"
                ),
                len(lines),
            )
        else:
            end_index = events[event_index + 1]["line_index"] if event_index + 1 < len(events) else len(lines)

        product_content = normalize_preserve_lines(
            "\n".join(lines[event["line_index"]:end_index])
        )
        if not product_content:
            continue
        parent = next(
            (
                item for item in events
                if item["key"] == event["parent_key"] and item["level"] == "top"
            ),
            None,
        )
        aliases = list(event["aliases"])
        if parent:
            aliases = unique_normalized_values([
                *aliases,
                parent["name"],
                *parent["aliases"],
            ])
        products.append({
            **event,
            "parent_name": str(parent["name"] if parent else ""),
            "aliases": aliases,
            "content": product_content,
        })

    if not products:
        return {}
    return {
        "schema_version": "1.1",
        "document_type": "product_service_catalog",
        "catalog_name": str(record.get("title") or strip_file_extension(record.get("file_name") or "")),
        "aliases": unique_normalized_values(
            [alias for product in products for alias in product.get("aliases", [])]
        ),
        "products": products,
    }


def value_added_product_metadata(
    profile: Dict[str, Any],
    product: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any]:
    product = product or {}
    return {
        "product_catalog": str(profile.get("catalog_name") or ""),
        "product_name": str(product.get("name") or ""),
        "product_aliases": " | ".join(product.get("aliases") or profile.get("aliases") or []),
        "product_parent": str(product.get("parent_name") or ""),
    }


def build_value_added_product_index_records(
    record: Dict[str, Any],
    profile: Dict[str, Any],
) -> List[Dict[str, Any]]:
    if not profile:
        return []

    output: List[Dict[str, Any]] = []
    products = profile.get("products") or []
    for index, product in enumerate(products, start=1):
        name = str(product.get("name") or "")
        aliases = " ".join(product.get("aliases") or [])
        parent = str(product.get("parent_name") or "")
        content = normalize_preserve_lines(product.get("content") or "")
        query_hints = " ".join(
            value for value in (
                name,
                aliases,
                parent,
                "是什麼 服務內容 有哪些 費用 價格 月租 半年繳 年繳 加購 申請 租借 綁約 限制",
            ) if value
        )
        output.append({
            "id": f"{record['id']}:product:{index}",
            "document_id": record["id"],
            "chunk_index": index,
            "question": normalize_text(query_hints),
            "answer": content,
            "content": content,
            "company": record["knowledge_base"],
            "knowledge_base": record["knowledge_base"],
            "category": "value_added_service",
            "title": record["title"],
            "source": record["file_name"],
            "page_no": "",
            "section": f"product {index} {name}",
            "record_type": "product_service",
            "document_type": profile["document_type"],
            **value_added_product_metadata(profile, product),
        })

    all_content = normalize_preserve_lines(
        "\n".join(
            str(product.get("content") or "")
            for product in products
            if product.get("level") == "top"
        )
    )
    output.append({
        "id": f"{record['id']}:product:catalog",
        "document_id": record["id"],
        "chunk_index": 0,
        "question": normalize_text(
            f"{profile.get('catalog_name') or ''} {' '.join(profile.get('aliases') or [])} "
            "有哪些加值服務 加值產品 單品銷售 熱門單品 清單 費用"
        ),
        "answer": all_content,
        "content": all_content,
        "company": record["knowledge_base"],
        "knowledge_base": record["knowledge_base"],
        "category": "value_added_service",
        "title": record["title"],
        "source": record["file_name"],
        "page_no": "",
        "section": "product catalog summary",
        "record_type": "product_service_catalog",
        "document_type": profile["document_type"],
        **value_added_product_metadata(profile),
    })
    return output


def split_pdf_section_blocks(section: Dict[str, Any], max_chars: int = 650) -> List[Dict[str, Any]]:
    content = str(section.get("content") or "")
    lines = [line.strip() for line in content.splitlines() if line.strip()]
    if len(lines) <= 1:
        return [section]

    is_campaign_form = bool(PDF_CAMPAIGN_FORM_RE.search(content))
    is_campaign_plan_page = bool(PDF_CAMPAIGN_PLAN_PAGE_RE.search(content))
    if is_campaign_form and is_campaign_plan_page:
        return split_campaign_summary_and_plan_blocks(section)
    if is_campaign_plan_page and not is_campaign_form:
        return split_pdf_campaign_plan_blocks(section)
    if PDF_COLON_PLAN_START_RE.search(content) and not is_campaign_form:
        return split_pdf_colon_plan_blocks(section)
    use_coarse_campaign_split = is_campaign_form or is_campaign_plan_page
    effective_max_chars = 1400 if use_coarse_campaign_split else max_chars
    blocks: List[Dict[str, Any]] = []
    current: List[str] = []
    current_heading = ""

    def flush() -> None:
        if not current:
            return
        block_content = normalize_preserve_lines("\n".join(current))
        if block_content:
            if current_heading and current_heading not in block_content:
                block_content = normalize_preserve_lines(f"{current_heading}\n{block_content}")
            block = dict(section)
            block["content"] = block_content
            block["section"] = f"{section.get('section') or 'page'} block {len(blocks) + 1}"
            blocks.append(block)
        current.clear()

    for line in lines:
        if PDF_SEPARATOR_RE.match(line):
            flush()
            current_heading = ""
            continue
        starts_new = False if use_coarse_campaign_split else bool(PDF_BLOCK_START_RE.search(line))
        would_be_long = len(" ".join(current + [line])) > effective_max_chars
        if current and (starts_new or would_be_long):
            flush()
        if starts_new and any(term in line for term in ("說明", "注意事項", "拆帳表", "備註範例")):
            current_heading = line
        current.append(line)

    flush()
    return blocks or [section]


def split_campaign_summary_and_plan_blocks(section: Dict[str, Any]) -> List[Dict[str, Any]]:
    lines = [line.strip() for line in str(section.get("content") or "").splitlines() if line.strip()]
    blocks: List[Dict[str, Any]] = []
    current: List[str] = []

    def flush() -> None:
        if not current:
            return
        block_content = normalize_preserve_lines("\n".join(current))
        if block_content:
            block = dict(section)
            block["content"] = block_content
            block["section"] = f"{section.get('section') or 'document'} block {len(blocks) + 1}"
            blocks.append(block)
        current.clear()

    for line in lines:
        if PDF_SEPARATOR_RE.match(line):
            flush()
            continue

        starts_plan = bool(PDF_CAMPAIGN_PLAN_START_RE.search(line))
        if current and starts_plan:
            flush()
        current.append(line)

    flush()
    return blocks or [section]


def split_pdf_colon_plan_blocks(section: Dict[str, Any], max_chars: int = 1400) -> List[Dict[str, Any]]:
    lines = [line.strip() for line in str(section.get("content") or "").splitlines() if line.strip()]
    blocks: List[Dict[str, Any]] = []
    current: List[str] = []

    def flush() -> None:
        if not current:
            return
        block_content = normalize_preserve_lines("\n".join(current))
        if block_content:
            block = dict(section)
            block["content"] = block_content
            block["section"] = f"{section.get('section') or 'page'} block {len(blocks) + 1}"
            blocks.append(block)
        current.clear()

    for line in lines:
        if PDF_SEPARATOR_RE.match(line):
            flush()
            continue

        starts_plan = bool(PDF_COLON_PLAN_START_RE.search(line))
        current_has_plan = any(PDF_COLON_PLAN_START_RE.search(item) for item in current)
        would_be_long = len(" ".join(current + [line])) > max_chars

        if current and ((starts_plan and current_has_plan) or would_be_long):
            flush()
        current.append(line)

    flush()
    return blocks or [section]


def split_pdf_campaign_plan_blocks(section: Dict[str, Any], max_chars: int = 1400) -> List[Dict[str, Any]]:
    lines = [line.strip() for line in str(section.get("content") or "").splitlines() if line.strip()]
    blocks: List[Dict[str, Any]] = []
    current: List[str] = []

    def flush() -> None:
        if not current:
            return
        block_content = normalize_preserve_lines("\n".join(current))
        if block_content:
            block = dict(section)
            block["content"] = block_content
            block["section"] = f"{section.get('section') or 'page'} block {len(blocks) + 1}"
            blocks.append(block)
        current.clear()

    for line in lines:
        if PDF_SEPARATOR_RE.match(line):
            flush()
            continue

        starts_plan = bool(PDF_CAMPAIGN_PLAN_START_RE.search(line))
        starts_rate_table = bool(PDF_RATE_TABLE_HEADING_RE.search(line))
        would_be_long = len(" ".join(current + [line])) > max_chars

        if current and (starts_plan or starts_rate_table or would_be_long):
            flush()
        current.append(line)

    flush()
    return blocks or [section]


def split_text_section_blocks(section: Dict[str, Any]) -> List[Dict[str, Any]]:
    content = str(section.get("content") or "")
    has_campaign_header = bool(CAMPAIGN_HEADER_RE.search(content))
    has_campaign_plans = bool(PDF_CAMPAIGN_PLAN_PAGE_RE.search(content))

    if has_campaign_header and has_campaign_plans:
        return split_campaign_summary_and_plan_blocks(section)
    if has_campaign_plans:
        return split_pdf_campaign_plan_blocks(section)
    return [section]


def normalize_docx_logical_lines(content: str) -> List[str]:
    text = normalize_preserve_lines(content)
    if not text:
        return []
    text = re.sub(
        r"(?<!^)([一二三四五六七八九十]+、(?=\s*(?:TV|收視|移機|分機|加值|設備|其他|優惠|申請|限制|方案|費用|裝機|行政|重要|規範)))",
        r"\n\1",
        text,
    )
    text = re.sub(r"(?<!^)(【方案名稱】)", r"\n\1", text)
    text = re.sub(r"(?<!^)(方案名稱\s*[:：])", r"\n\1", text)
    text = re.sub(r"(?<!^)([-]{10,})", r"\n\1", text)
    lines: List[str] = []
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if PDF_SEPARATOR_RE.match(line):
            lines.append(line)
            continue
        parts = [
            part.strip()
            for part in re.split(r"(?=【(?:適用對象|申請方式|申請條件|優惠內容|限制條件|備註|注意事項)】)", line)
            if part.strip()
        ]
        lines.extend(parts or [line])
    return lines


def split_docx_section_blocks(section: Dict[str, Any], max_chars: int = 1400) -> List[Dict[str, Any]]:
    content = str(section.get("content") or "")
    has_campaign_header = bool(CAMPAIGN_HEADER_RE.search(content))
    has_campaign_plans = bool(PDF_CAMPAIGN_PLAN_PAGE_RE.search(content))
    if has_campaign_header and has_campaign_plans:
        return split_campaign_summary_and_plan_blocks(section)
    if has_campaign_plans:
        return split_pdf_campaign_plan_blocks(section)

    lines = normalize_docx_logical_lines(content)
    if len(lines) <= 1:
        return [section]

    blocks: List[Dict[str, Any]] = []
    current: List[str] = []
    current_heading = ""

    def flush() -> None:
        if not current:
            return
        block_content = normalize_preserve_lines("\n".join(current))
        if block_content:
            if current_heading and current_heading not in block_content:
                block_content = normalize_preserve_lines(f"{current_heading}\n{block_content}")
            block = dict(section)
            block["content"] = block_content
            block["section"] = f"{section.get('section') or 'docx'} block {len(blocks) + 1}"
            blocks.append(block)
        current.clear()

    for line in lines:
        if PDF_SEPARATOR_RE.match(line):
            flush()
            current_heading = ""
            continue
        if DOCX_PLAN_START_RE.search(line):
            heading_match = DOCX_HEADING_RE.search(line)
            current_heading = normalize_text(heading_match.group(0) if heading_match else line)
            current_has_plan = any(DOCX_PLAN_START_RE.search(item) for item in current)
            if current and current_has_plan:
                flush()
            current.append(line)
            continue
        heading_match = DOCX_HEADING_RE.search(line)
        if heading_match:
            current_heading = normalize_text(heading_match.group(0))
            if current:
                flush()
            current.append(line)
            continue
        starts_new = bool(DOCX_BLOCK_START_RE.search(line)) and not current_heading
        would_be_long = len(" ".join(current + [line])) > max_chars
        if current and (starts_new or would_be_long):
            flush()
        current.append(line)

    flush()
    return blocks or [section]


def is_docx_paragraph_section(section: Dict[str, Any]) -> bool:
    return (
        section.get("section_type") == "docx"
        and str(section.get("section") or "").startswith("paragraph")
    )


def merge_docx_paragraph_sections(sections: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    merged: List[Dict[str, Any]] = []
    buffer: List[Dict[str, Any]] = []

    def flush() -> None:
        if not buffer:
            return
        content = normalize_preserve_lines("\n".join(str(item.get("content") or "") for item in buffer))
        if content:
            first = str(buffer[0].get("section") or "paragraph")
            last = str(buffer[-1].get("section") or first)
            section = dict(buffer[0])
            section["content"] = content
            section["section"] = first if first == last else f"{first}-{last}"
            merged.append(section)
        buffer.clear()

    for section in sections:
        if is_docx_paragraph_section(section):
            buffer.append(section)
            continue
        flush()
        merged.append(section)

    flush()
    return merged


def merge_pdf_campaign_plan_continuations(sections: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    merged: List[Dict[str, Any]] = []

    for section in sections:
        content = normalize_preserve_lines(str(section.get("content") or ""))
        previous_content = str(merged[-1].get("content") or "") if merged else ""
        if (
            merged
            and content
            and PDF_CAMPAIGN_PLAN_ITEM_CONTINUATION_RE.search(content)
            and PDF_CAMPAIGN_PLAN_START_RE.search(previous_content)
        ):
            merged[-1]["content"] = normalize_preserve_lines(
                f"{merged[-1].get('content')}\n{content}"
            )
            continue
        if (
            merged
            and content
            and "派工備註範例" in previous_content
            and PDF_DISPATCH_NOTE_CONTINUATION_RE.search(content)
        ):
            merged[-1]["content"] = normalize_preserve_lines(
                f"{merged[-1].get('content')}\n{content}"
            )
            continue
        merged.append(section)

    return merged


def chunk_sections(sections: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    sections = merge_docx_paragraph_sections(sections)
    block_sections: List[Dict[str, Any]] = []
    for section in sections:
        if section.get("page_no"):
            block_sections.extend(split_pdf_section_blocks(section))
        elif section.get("section_type") == "docx_table":
            block_sections.append(section)
        elif section.get("section_type") == "docx":
            block_sections.extend(split_docx_section_blocks(section))
        elif section.get("section_type") == "text":
            block_sections.extend(split_text_section_blocks(section))
        else:
            block_sections.append(section)

    block_sections = merge_pdf_campaign_plan_continuations(block_sections)

    chunks = []
    for block_section in block_sections:
        preserve_lines = bool(
            block_section.get("page_no")
            or block_section.get("section_type") in {"docx", "docx_table", "text"}
        )
        for content in chunk_text(block_section.get("content", ""), preserve_lines=preserve_lines):
            chunks.append({
                "content": content,
                "question": block_section.get("question") or "",
                "answer": block_section.get("answer") or "",
                "company": block_section.get("company") or "",
                "page_no": block_section.get("page_no"),
                "section": block_section.get("section"),
            })
    return chunks


def infer_category_from_text(text: str) -> str:
    normalized = (text or "").lower()
    if any(keyword in normalized for keyword in [
        "帳單", "繳費", "退費", "退租", "發票", "載具", "月租", "違約金", "費用"
    ]):
        return "billing"
    if any(keyword in normalized for keyword in [
        "網路", "寬頻", "wifi", "wi-fi", "數據機", "固定ip", "固定 ip", "光纖", "mesh"
    ]):
        return "network_support"
    if any(keyword in normalized for keyword in [
        "第四台", "有線電視", "機上盒", "遙控器", "頻道", "無訊號", "電視"
    ]):
        return "set_top_box"
    if any(keyword in normalized for keyword in [
        "line tv", "linetv", "hbo", "friday", "加值", "監視器", "攝影機", "ott"
    ]):
        return "value_added_service"
    if any(keyword in normalized for keyword in [
        "規定", "政策", "辦法", "注意事項", "限制", "資格"
    ]):
        return "policy"
    if any(keyword in normalized for keyword in [
        "流程", "步驟", "sop", "作業", "處理方式"
    ]):
        return "sop"
    return ""


def strip_file_extension(value: str) -> str:
    text = normalize_text(value)
    suffix = Path(text).suffix
    if suffix:
        text = text[: -len(suffix)]
    return text.strip(" _-")


def has_any_term(text: str, terms: tuple[str, ...]) -> bool:
    normalized = str(text or "").lower()
    return any(term.lower() in normalized for term in terms)


def compact_context_parts(parts: List[str]) -> List[str]:
    output: List[str] = []
    normalized_seen: set[str] = set()
    for part in parts:
        value = normalize_text(part).strip("。")
        if not value:
            continue
        normalized = re.sub(r"\s+", "", value).lower()
        if normalized in normalized_seen:
            continue
        if any(normalized and normalized in existing for existing in normalized_seen):
            continue
        normalized_seen.add(normalized)
        output.append(value)
    return output


def extract_plan_name_from_content(content: str, fallback: str = "") -> str:
    match = PLAN_NAME_RE.search(content or "")
    if match:
        return normalize_text(match.group(1)).strip("。")
    return strip_file_extension(fallback)


def extract_plan_block_label(content: str) -> str:
    match = PLAN_BLOCK_LABEL_RE.search(content or "")
    if match:
        return normalize_text(match.group(1)).strip("。")

    for line in normalize_preserve_lines(content).splitlines()[:3]:
        value = normalize_text(line)
        if value.endswith("拆帳表") or value in {"售價", "優惠內容", "限制條件"}:
            return value
    return ""


def is_social_discount_content(content: str) -> bool:
    return has_any_term(content, SOCIAL_DISCOUNT_TERMS)


def is_campaign_activity_content(content: str) -> bool:
    return has_any_term(content, CAMPAIGN_ACTIVITY_TERMS) and not is_social_discount_content(content)


def make_chunk_answer(content: str, plan_name: str, block_label: str) -> str:
    answer = normalize_preserve_lines(content)
    prefixes = []
    if plan_name and plan_name not in answer:
        prefixes.append(f"方案名稱：{plan_name}")
    if block_label and block_label not in answer:
        prefixes.append(block_label)
    if prefixes:
        answer = normalize_preserve_lines("\n".join(prefixes + [answer]))
    return answer


def unique_normalized_values(values: List[Any]) -> List[str]:
    output: List[str] = []
    seen: set[str] = set()
    for value in values:
        text = normalize_text(value).strip(" _-。")
        if not text:
            continue
        key = re.sub(r"[\s_.\-()（）]+", "", text).lower()
        if not key or key in seen:
            continue
        seen.add(key)
        output.append(text)
    return output


def extract_campaign_speeds(text: str) -> List[str]:
    source = str(text or "")
    source = re.sub(
        r"(?m)^\s*(?:[-–—]{1,3}\s*)?\d{1,2}\s*[.、．]\s*"
        r"(?=\d+(?:\.\d+)?\s*[MG]\s*/)",
        "",
        source,
        flags=re.IGNORECASE,
    )
    speeds = []
    for download, upload in CAMPAIGN_SPEED_RE.findall(source):
        speeds.append(f"{normalize_text(download).upper()}/{normalize_text(upload).upper()}")
    return unique_normalized_values(speeds)


def extract_campaign_contract_months(text: str) -> List[int]:
    source = str(text or "")
    positioned_months: List[tuple[int, int]] = []
    for match in CAMPAIGN_CONTRACT_MONTH_RE.finditer(source):
        try:
            month = int(match.group(1))
        except (TypeError, ValueError):
            continue
        if 1 <= month <= 120:
            positioned_months.append((match.start(), month))
    for label, month in CAMPAIGN_CHINESE_CONTRACT_MONTHS.items():
        for match in re.finditer(
            rf"(?:綁約(?:期限)?|合約(?:期間)?|約期|租期)[^。\n；;]{{0,20}}?{re.escape(label)}",
            source,
            re.IGNORECASE,
        ):
            positioned_months.append((match.start(), month))
    months = []
    for _, month in sorted(positioned_months):
        if month not in months:
            months.append(month)
    return months


def extract_campaign_period(text: str) -> str:
    match = CAMPAIGN_PERIOD_RE.search(str(text or ""))
    return normalize_text(match.group(1)).strip("。") if match else ""


def extract_campaign_occasion_terms(text: str) -> List[str]:
    """Expand matched holiday/season wording into searchable related terms."""
    normalized = str(text or "").lower()
    values: List[str] = []
    for canonical, aliases in CAMPAIGN_OCCASION_ALIASES.items():
        if any(alias.lower() in normalized for alias in aliases):
            values.extend([canonical, *aliases])
    return unique_normalized_values(values)


def extract_campaign_detail_lines(
    text: str,
    terms: tuple[str, ...],
    max_items: int = 20,
) -> List[str]:
    details: List[str] = []
    for raw_line in normalize_preserve_lines(text).splitlines():
        for segment in re.split(r"(?<=[。；;])", raw_line):
            value = normalize_text(segment).strip("；;。 ")
            if not value or not has_any_term(value, terms):
                continue
            details.append(value[:220])
    return unique_normalized_values(details)[:max_items]


def extract_campaign_gift_items(text: str) -> List[str]:
    return extract_campaign_detail_lines(text, GIFT_DETAIL_TERMS)


def extract_campaign_lottery_details(text: str) -> List[str]:
    return extract_campaign_detail_lines(text, LOTTERY_TERMS)


def extract_campaign_aliases(
    campaign_name: str,
    record_title: str,
    file_name: str,
    text: str,
) -> List[str]:
    values: List[Any] = [campaign_name, record_title, strip_file_extension(file_name)]

    if campaign_name:
        values.append(re.sub(r"\s+", "", campaign_name))
        values.append(re.sub(r"[_（(].*$", "", campaign_name).strip())

    aliases = unique_normalized_values(values)
    return [value for value in aliases if len(re.sub(r"\s+", "", value)) >= 2][:12]


def detect_campaign_document(record: Dict[str, Any], chunks: List[Dict[str, Any]]) -> bool:
    title = str(record.get("title") or "")
    file_name = str(record.get("file_name") or "")
    text = "\n".join(str(chunk.get("content") or "") for chunk in chunks)
    combined = "\n".join([title, file_name, text])
    title_identity = "\n".join([title, file_name])
    leading_text = "\n".join(
        str(chunk.get("content") or "") for chunk in chunks[:3]
    )[:6000]
    title_has_campaign_identity = has_any_term(title_identity, CAMPAIGN_TITLE_TERMS)

    if (
        title_has_campaign_identity
        and is_social_discount_content(combined)
        and has_any_term(combined, FEE_TERMS + ELIGIBILITY_TERMS)
    ):
        return True
    if title_has_campaign_identity and PLAN_NAME_RE.search(leading_text):
        return True

    signal_count = sum(
        1
        for terms in (
            CAMPAIGN_ACTIVITY_TERMS,
            FEE_TERMS,
            CONTRACT_TERMS,
            GIFT_TERMS,
            LOTTERY_TERMS,
            ELIGIBILITY_TERMS,
        )
        if has_any_term(combined, terms)
    )
    if title_has_campaign_identity:
        return signal_count >= 2
    return bool(PLAN_NAME_RE.search(leading_text)) and signal_count >= 3


def classify_campaign_sections(content: str) -> List[str]:
    return [
        label
        for label, terms in CAMPAIGN_SECTION_LABELS.items()
        if has_any_term(content, terms)
    ]


def infer_campaign_service_types(text: str) -> List[str]:
    normalized = str(text or "").lower()
    services = [
        service
        for service, terms in CAMPAIGN_SERVICE_TERMS.items()
        if any(term.lower() in normalized for term in terms)
    ]
    if "電視網路同裝" in services:
        services = ["電視網路同裝"] + [
            service for service in services if service not in {"寬頻", "有線電視", "電視網路同裝"}
        ]
    return unique_normalized_values(services)


def build_campaign_profile(
    record: Dict[str, Any],
    chunks: List[Dict[str, Any]],
) -> Dict[str, Any]:
    if not detect_campaign_document(record, chunks):
        return {}

    title = str(record.get("title") or record.get("file_name") or "")
    file_name = str(record.get("file_name") or "")
    all_content = normalize_preserve_lines(
        "\n".join(str(chunk.get("content") or "") for chunk in chunks)
    )
    campaign_name = ""
    for chunk in chunks:
        campaign_name = extract_plan_name_from_content(str(chunk.get("content") or ""), "")
        if campaign_name:
            break
    campaign_name = campaign_name or strip_file_extension(title) or strip_file_extension(file_name)
    aliases = extract_campaign_aliases(campaign_name, title, file_name, all_content)
    speeds = extract_campaign_speeds(all_content)
    contract_months = extract_campaign_contract_months(all_content)
    payment_terms = [
        term for term in CAMPAIGN_PAYMENT_TERMS if term.lower() in all_content.lower()
    ]
    customer_types = [
        term for term in CAMPAIGN_CUSTOMER_TERMS if term.lower() in all_content.lower()
    ]
    service_types = infer_campaign_service_types(all_content)
    profile_context = "\n".join([title, file_name, all_content])
    occasion_terms = extract_campaign_occasion_terms(profile_context)
    gift_items = extract_campaign_gift_items(all_content)
    lottery_details = extract_campaign_lottery_details(all_content)
    variants = []

    for index, chunk in enumerate(chunks, start=1):
        content = normalize_preserve_lines(chunk.get("content") or "")
        if not content:
            continue
        variant = {
            "chunk_index": index,
            "section": str(chunk.get("section") or ""),
            "label": extract_plan_block_label(content),
            "speeds": extract_campaign_speeds(content),
            "contract_months": extract_campaign_contract_months(content),
            "payment_terms": [
                term for term in CAMPAIGN_PAYMENT_TERMS if term.lower() in content.lower()
            ],
            "sections": classify_campaign_sections(content),
            "occasion_terms": extract_campaign_occasion_terms(content),
            "gift_items": extract_campaign_gift_items(content),
            "lottery_details": extract_campaign_lottery_details(content),
        }
        if variant["label"] or variant["speeds"] or variant["contract_months"] or variant["sections"]:
            variants.append(variant)

    required_status = {
        "campaign_name": bool(campaign_name),
        "service_or_speed": bool(service_types or speeds),
        "price": has_any_term(all_content, FEE_TERMS),
        "eligibility_or_contract": bool(
            has_any_term(all_content, ELIGIBILITY_TERMS)
            or has_any_term(all_content, CONTRACT_TERMS)
        ),
    }
    missing_fields = [name for name, present in required_status.items() if not present]
    return {
        "schema_version": "1.0",
        "document_type": "promotion_campaign",
        "campaign_name": campaign_name,
        "aliases": aliases,
        "service_types": service_types,
        "speeds": speeds,
        "contract_months": contract_months,
        "payment_terms": unique_normalized_values(payment_terms),
        "customer_types": unique_normalized_values(customer_types),
        "valid_period": extract_campaign_period(all_content),
        "occasion_terms": occasion_terms,
        "gift_items": gift_items,
        "lottery_details": lottery_details,
        "section_types": classify_campaign_sections(all_content),
        "variants": variants,
        "validation": {
            "status": "complete" if not missing_fields else "needs_review",
            "missing_fields": missing_fields,
        },
    }


def campaign_metadata(profile: Dict[str, Any], variant: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
    variant = variant or {}
    return {
        "document_type": str(profile.get("document_type") or ""),
        "campaign_name": str(profile.get("campaign_name") or ""),
        "campaign_aliases": " | ".join(profile.get("aliases") or []),
        "campaign_sections": " | ".join(variant.get("sections") or profile.get("section_types") or []),
        "service_types": " | ".join(profile.get("service_types") or []),
        "speeds": " | ".join(variant.get("speeds") or profile.get("speeds") or []),
        "contract_months": " | ".join(
            str(value) for value in (variant.get("contract_months") or profile.get("contract_months") or [])
        ),
        "payment_terms": " | ".join(variant.get("payment_terms") or profile.get("payment_terms") or []),
        "customer_types": " | ".join(profile.get("customer_types") or []),
        "valid_period": str(profile.get("valid_period") or ""),
        "occasion_terms": " | ".join(variant.get("occasion_terms") or profile.get("occasion_terms") or []),
        "gift_items": " | ".join(variant.get("gift_items") or profile.get("gift_items") or []),
        "lottery_details": " | ".join(variant.get("lottery_details") or profile.get("lottery_details") or []),
    }


def build_campaign_index_records(
    record: Dict[str, Any],
    chunks: List[Dict[str, Any]],
    category: str,
    profile: Dict[str, Any],
) -> List[Dict[str, Any]]:
    if not profile:
        return []

    output: List[Dict[str, Any]] = []
    campaign_name = str(profile.get("campaign_name") or record.get("title") or "")
    aliases = " ".join(profile.get("aliases") or [])
    services = " ".join(profile.get("service_types") or [])
    speeds = " ".join(profile.get("speeds") or [])
    months = " ".join(f"{value}個月" for value in profile.get("contract_months") or [])
    occasions = " ".join(profile.get("occasion_terms") or [])
    gifts = " ".join(profile.get("gift_items") or [])
    lottery = " ".join(profile.get("lottery_details") or [])
    summary_terms = " ".join(
        value for value in [campaign_name, aliases, services, speeds, months, occasions, gifts, lottery] if value
    )

    for chunk_index, chunk in enumerate(chunks, start=1):
        content = normalize_preserve_lines(chunk.get("content") or "")
        if not content:
            continue
        variant = next(
            (
                item for item in profile.get("variants") or []
                if int(item.get("chunk_index") or 0) == chunk_index
            ),
            {},
        )
        block_label = str(variant.get("label") or extract_plan_block_label(content) or "")
        section_names = " ".join(variant.get("sections") or classify_campaign_sections(content))
        variant_speeds = " ".join(variant.get("speeds") or [])
        variant_months = " ".join(
            f"{value}個月" for value in variant.get("contract_months") or []
        )
        variant_occasions = " ".join(variant.get("occasion_terms") or [])
        variant_gifts = " ".join(variant.get("gift_items") or [])
        variant_lottery = " ".join(variant.get("lottery_details") or [])
        query_hints = " ".join(
            value
            for value in [
                summary_terms,
                block_label,
                variant_speeds,
                variant_months,
                section_names,
                variant_occasions,
                variant_gifts,
                variant_lottery,
                "優惠內容 月租 費用 贈品 贈品內容 送什麼 節慶 節日 抽獎 中獎 獎項 獎品 抽獎資格 申請資格 限制 綁約",
            ]
            if value
        )
        answer = make_chunk_answer(content, campaign_name, block_label)
        record_type = "campaign_variant" if block_label or variant_speeds or variant_months else "campaign_summary"
        output.append({
            "id": f"{record['id']}:campaign:{chunk_index}",
            "document_id": record["id"],
            "chunk_index": chunk_index,
            "question": normalize_text(query_hints),
            "answer": answer,
            "content": content,
            "company": chunk.get("company") or record["knowledge_base"],
            "knowledge_base": record["knowledge_base"],
            "category": category,
            "title": record["title"],
            "source": record["file_name"],
            "page_no": chunk.get("page_no") or "",
            "section": chunk.get("section") or "",
            "record_type": record_type,
            **campaign_metadata(profile, variant),
        })

    return output


def build_qa_questions_for_chunk(content: str, record_title: str) -> List[str]:
    plan_name = extract_plan_name_from_content(content, record_title)
    block_label = extract_plan_block_label(content)
    context = " ".join(compact_context_parts([plan_name, block_label, record_title]))
    if not context:
        return []

    questions: List[str] = []

    def add(question: str) -> None:
        value = normalize_text(question)
        if value and value not in questions:
            questions.append(value)

    if is_social_discount_content(content):
        add(f"{context} 低收入 中低收入 身心障礙 優惠方案 申請方式 資格 條件是什麼？")
        if has_any_term(content, FEE_TERMS):
            add(f"{context} 優惠內容 費用 裝機費 收視費 怎麼算？")
        if has_any_term(content, ELIGIBILITY_TERMS):
            add(f"{context} 適用對象 申請條件 限制條件有哪些？")
        return questions[:5]

    if is_campaign_activity_content(content):
        add(f"{context} 有哪些推薦優惠方案 活動方案 完整內容？")
        add(f"{context} 方案說明 活動期間 售價 綁約 贈品 條件？")

    if has_any_term(content, FEE_TERMS):
        add(f"{context} 費用 月租 月繳 半年繳 年繳 售價 裝機費 押金 怎麼算？")

    if has_any_term(content, GIFT_TERMS):
        gift_items = " ".join(extract_campaign_gift_items(content))
        add(f"{context} 贈品內容 送什麼 有哪些好禮 家電 POINTS LINE TV LITV？ {gift_items}")

    occasion_terms = extract_campaign_occasion_terms(f"{record_title}\n{content}")
    if occasion_terms:
        add(f"{context} {' '.join(occasion_terms)} 節慶 節日 有什麼優惠活動？")

    if has_any_term(content, LOTTERY_TERMS):
        lottery_details = " ".join(extract_campaign_lottery_details(content))
        add(f"{context} 抽獎活動 怎麼參加 抽獎資格 抽獎日期 中獎方式 獎項 獎品有哪些？ {lottery_details}")

    if has_any_term(content, WIFI_TERMS):
        add(f"{context} WiFi WIFI 租借 月租50元 可以只付一個月嗎 繳費期間？")

    if has_any_term(content, ELIGIBILITY_TERMS):
        add(f"{context} 舊戶 可以申辦嗎 申請條件 適用資格 限制條件？")

    if has_any_term(content, CONTRACT_TERMS):
        add(f"{context} 綁約 違約金 活動期間 到期 中途換約 升級 規定？")

    if has_any_term(content, EQUIPMENT_TERMS):
        add(f"{context} 設備押金 數據機 機上盒 聯網機上盒 遺失損壞賠償？")

    return questions[:9]


def build_qa_index_records(
    record: Dict[str, Any],
    chunks: List[Dict[str, Any]],
    category: str,
) -> List[Dict[str, Any]]:
    title = str(record.get("title") or record.get("file_name") or "")
    output: List[Dict[str, Any]] = []

    for chunk_index, chunk in enumerate(chunks, start=1):
        content = normalize_preserve_lines(chunk.get("content") or "")
        if not content:
            continue

        # CSV rows already contain the customer-facing question. Generating
        # broad template questions for them can drown out that exact FAQ.
        if normalize_text(chunk.get("question") or ""):
            continue

        questions = build_qa_questions_for_chunk(content, title)
        if not questions:
            continue

        plan_name = extract_plan_name_from_content(content, title)
        block_label = extract_plan_block_label(content)
        answer = make_chunk_answer(content, plan_name, block_label)

        for question_index, question in enumerate(questions, start=1):
            output.append({
                "id": f"{record['id']}:qa:{chunk_index}:{question_index}",
                "document_id": record["id"],
                "chunk_index": chunk_index,
                "question": question,
                "answer": answer,
                "content": content,
                "company": chunk.get("company") or record["knowledge_base"],
                "knowledge_base": record["knowledge_base"],
                "category": category,
                "title": record["title"],
                "source": record["file_name"],
                "page_no": chunk.get("page_no") or "",
                "section": f"{chunk.get('section') or ''} qa {question_index}".strip(),
                "record_type": "qa",
            })

    return output


def manifest_path() -> Path:
    return Path(RAG_LOCAL_MANIFEST_PATH)


def docs_dir() -> Path:
    return Path(RAG_LOCAL_DOCS_DIR)


def resolve_document_file_path(raw_path: str | Path) -> Path:
    path = Path(raw_path or "")
    current_docs_candidate = docs_dir() / path.name
    if path.name and current_docs_candidate.exists():
        return current_docs_candidate

    if path.is_absolute():
        return path

    candidates = [
        BASE_DIR / path,
        Path.cwd() / path,
        docs_dir() / path.name,
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return candidates[0]


def load_manifest() -> Dict[str, Any]:
    path = manifest_path()
    if not path.exists():
        return {"documents": []}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return {"documents": []}
    if not isinstance(data, dict):
        return {"documents": []}
    data.setdefault("documents", [])
    return data


def save_manifest(data: Dict[str, Any]) -> None:
    path = manifest_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def list_documents(
    status: str = "active",
    knowledge_base: Optional[str] = None,
) -> List[Dict[str, Any]]:
    docs = load_manifest().get("documents", [])
    output = []
    for doc in docs:
        if status != "all" and doc.get("status", "active") != status:
            continue
        if knowledge_base and doc.get("knowledge_base") != knowledge_base:
            continue
        output.append(doc)
    return sorted(output, key=lambda item: item.get("created_at", ""), reverse=True)


def get_document(document_id: str) -> Optional[Dict[str, Any]]:
    for doc in load_manifest().get("documents", []):
        if str(doc.get("id")) == str(document_id):
            return doc
    return None


def upsert_document_record(record: Dict[str, Any]) -> Dict[str, Any]:
    data = load_manifest()
    docs = data.setdefault("documents", [])
    for index, current in enumerate(docs):
        if current.get("id") == record.get("id"):
            docs[index] = record
            save_manifest(data)
            return record
    docs.append(record)
    save_manifest(data)
    return record


def migrate_legacy_common_documents(
    indexed_by: Optional[str] = None,
) -> Dict[str, Any]:
    """Move legacy common documents to the central regional common knowledge base."""
    data = load_manifest()
    documents = data.setdefault("documents", [])
    migrated_ids: List[str] = []
    migrated_record_count = 0
    migrated_at = now_iso()

    for record in documents:
        current_base = str(record.get("knowledge_base") or "").strip()
        if current_base != LEGACY_COMMON_KNOWLEDGE_BASE:
            continue
        record["knowledge_base"] = CENTRAL_COMMON_KNOWLEDGE_BASE
        record["knowledge_base_migrated_from"] = LEGACY_COMMON_KNOWLEDGE_BASE
        record["knowledge_base_migrated_at"] = migrated_at
        record["updated_at"] = migrated_at
        migrated_record_count += 1
        if record.get("status", "active") == "active":
            migrated_ids.append(str(record.get("id") or ""))

    if not migrated_record_count:
        return {
            "migrated_count": 0,
            "indexed_count": 0,
            "failed_count": 0,
            "documents": [],
        }

    save_manifest(data)

    migrated_documents = []
    for document_id in migrated_ids:
        record = get_document(document_id)
        if not record:
            continue
        migrated_documents.append(
            index_document(dict(record), indexed_by=indexed_by or "system")
        )

    return {
        "migrated_count": migrated_record_count,
        "indexed_count": sum(
            1
            for record in migrated_documents
            if record.get("processing_status") == "indexed"
        ),
        "failed_count": sum(
            1
            for record in migrated_documents
            if record.get("processing_status") == "failed"
        ),
        "documents": migrated_documents,
    }


def cleanup_deleted_file_records_by_name(file_name: str) -> None:
    safe_name = sanitize_filename(file_name)
    if not safe_name:
        return

    for doc in load_manifest().get("documents", []):
        if doc.get("status") != "deleted":
            continue
        if sanitize_filename(str(doc.get("file_name") or "")) != safe_name:
            continue

        try:
            file_path = resolve_document_file_path(doc.get("file_path", ""))
            if file_path.exists():
                file_path.unlink()
        except Exception:
            pass


def load_searcher():
    from app.services.kb_core import load_kb_searcher

    return load_kb_searcher(
        persist_dir=RAG_LOCAL_PERSIST_DIR,
        collection_name=RAG_LOCAL_COLLECTION,
        model_name=RAG_LOCAL_EMBED_MODEL,
        device=RAG_LOCAL_EMBED_DEVICE,
    )


def is_hnsw_index_error(exc: Exception) -> bool:
    message = str(exc).lower()
    return (
        "hnsw" in message
        or "error in compaction" in message
        or "error loading hnsw index" in message
        or "constructing hnsw segment reader" in message
    )


def upsert_records_with_index_repair(
    searcher,
    records: List[Dict[str, Any]],
    document_id: str,
    *,
    full_rebuild: bool = False,
) -> int:
    try:
        return searcher.upsert_records(records, delete_document_id=document_id)
    except Exception as exc:
        if not is_hnsw_index_error(exc):
            raise
        if full_rebuild:
            raise RuntimeError(
                "Chroma HNSW 在完整重建期間仍無法寫入；"
                "目前重建目錄可能不是全新索引，或仍被其他程序占用。"
            ) from exc
        raise RuntimeError(
            "Chroma HNSW 實體索引損壞，已停止單筆寫入；"
            "必須完整重建索引，避免只重建目前文件而清除其他文件向量。"
        ) from exc


def kb_index_lock_path() -> Path:
    return Path(RAG_LOCAL_MANIFEST_PATH).parent / ".chroma-index.write.lock"


def kb_index_write_lock():
    from app.services.kb_index_lock import interprocess_index_lock

    return interprocess_index_lock(kb_index_lock_path())


def _close_searcher(searcher: Any) -> None:
    close = getattr(searcher, "close", None)
    if callable(close):
        close()


def reset_runtime_search_cache(*, clear_degraded: bool = False) -> None:
    try:
        from app.services.kb_service import reset_local_searcher_cache

        reset_local_searcher_cache(clear_degraded=clear_degraded)
    except Exception:
        pass


def release_chroma_runtime_handles() -> None:
    reset_runtime_search_cache()
    try:
        from app.services.kb_core import clear_chroma_system_cache

        clear_chroma_system_cache()
    except Exception:
        pass


def chroma_backup_dir() -> Path:
    persist_dir = Path(RAG_LOCAL_PERSIST_DIR)
    return persist_dir.with_name(f"{persist_dir.name}.backup")


def should_auto_backup_chroma_store() -> bool:
    try:
        persist_dir = Path(RAG_LOCAL_PERSIST_DIR).resolve()
        docs_path = Path(RAG_LOCAL_DOCS_DIR).resolve()
    except Exception:
        return False
    return persist_dir.parent == docs_path.parent


def is_path_within_directory(path: Path, directory: Path) -> bool:
    try:
        resolved_path = path.resolve()
        resolved_directory = directory.resolve()
    except Exception:
        return False
    return resolved_path == resolved_directory or resolved_directory in resolved_path.parents


def build_chroma_manifest_snapshot() -> Dict[str, Any]:
    documents = []
    for record in load_manifest().get("documents", []):
        if not isinstance(record, dict):
            continue
        if record.get("status", "active") != "active":
            continue
        if record.get("processing_status") != "indexed":
            continue
        documents.append({
            "id": str(record.get("id") or ""),
            "file_name": str(record.get("file_name") or ""),
            "file_path": str(record.get("file_path") or ""),
            "file_size": int(record.get("file_size") or 0),
            "knowledge_base": str(record.get("knowledge_base") or ""),
            "category": str(record.get("category") or ""),
            "updated_at": str(record.get("updated_at") or ""),
            "processed_at": str(record.get("processed_at") or ""),
            "indexed_vector_count": int(record.get("indexed_vector_count") or 0),
        })
    documents.sort(key=lambda item: item["id"])
    return {
        "schema_version": 1,
        "documents": documents,
    }


def chroma_manifest_signature(snapshot: Optional[Dict[str, Any]] = None) -> str:
    data = snapshot if isinstance(snapshot, dict) else build_chroma_manifest_snapshot()
    payload = json.dumps(data, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def diff_chroma_manifest_snapshots(
    backup_snapshot: Optional[Dict[str, Any]],
    current_snapshot: Optional[Dict[str, Any]] = None,
) -> Dict[str, List[str]]:
    backup_docs = (
        backup_snapshot.get("documents")
        if isinstance(backup_snapshot, dict)
        else []
    )
    current = current_snapshot if isinstance(current_snapshot, dict) else build_chroma_manifest_snapshot()
    current_docs = current.get("documents") if isinstance(current, dict) else []
    backup_by_id = {
        str(item.get("id") or ""): item
        for item in backup_docs
        if isinstance(item, dict) and item.get("id")
    }
    current_by_id = {
        str(item.get("id") or ""): item
        for item in current_docs
        if isinstance(item, dict) and item.get("id")
    }
    changed_document_ids = [
        document_id
        for document_id, current_record in current_by_id.items()
        if backup_by_id.get(document_id) != current_record
    ]
    deleted_document_ids = [
        document_id
        for document_id in backup_by_id
        if document_id not in current_by_id
    ]
    return {
        "changed_document_ids": sorted(changed_document_ids),
        "deleted_document_ids": sorted(deleted_document_ids),
    }


def read_chroma_backup_metadata(backup_dir: Path) -> Dict[str, Any]:
    metadata_path = backup_dir / "_backup_metadata.json"
    if not metadata_path.exists():
        return {}
    try:
        data = json.loads(metadata_path.read_text(encoding="utf-8"))
    except Exception:
        return {}
    return data if isinstance(data, dict) else {}


def backup_active_chroma_store(*, _lock_held: bool = False) -> Dict[str, Any]:
    if not _lock_held:
        with kb_index_write_lock():
            return backup_active_chroma_store(_lock_held=True)

    persist_dir = Path(RAG_LOCAL_PERSIST_DIR)
    backup_dir = chroma_backup_dir()
    if not persist_dir.exists():
        return {
            "backup_created": False,
            "reason": "persist_dir_missing",
            "persist_dir": str(persist_dir),
            "backup_dir": str(backup_dir),
        }
    if not any(persist_dir.iterdir()):
        return {
            "backup_created": False,
            "reason": "persist_dir_empty",
            "persist_dir": str(persist_dir),
            "backup_dir": str(backup_dir),
        }

    release_chroma_runtime_handles()
    backup_dir.parent.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    tmp_dir = backup_dir.with_name(f"{backup_dir.name}.tmp-{stamp}-{uuid.uuid4().hex[:8]}")
    previous_dir = backup_dir.with_name(f"{backup_dir.name}.previous-{stamp}-{uuid.uuid4().hex[:8]}")
    try:
        manifest_snapshot = build_chroma_manifest_snapshot()
        shutil.copytree(persist_dir, tmp_dir)
        metadata = {
            "created_at": now_iso(),
            "source_dir": str(persist_dir),
            "backup_dir": str(backup_dir),
            "manifest_signature": chroma_manifest_signature(manifest_snapshot),
            "manifest_snapshot": manifest_snapshot,
        }
        try:
            (tmp_dir / "_backup_metadata.json").write_text(
                json.dumps(metadata, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
        except Exception:
            pass

        had_previous = backup_dir.exists()
        if had_previous:
            backup_dir.replace(previous_dir)
        tmp_dir.replace(backup_dir)
        if previous_dir.exists():
            shutil.rmtree(previous_dir, ignore_errors=True)
        return {
            "backup_created": True,
            "persist_dir": str(persist_dir),
            "backup_dir": str(backup_dir),
            "replaced_previous": had_previous,
            "manifest_signature": metadata["manifest_signature"],
        }
    except Exception:
        if tmp_dir.exists():
            shutil.rmtree(tmp_dir, ignore_errors=True)
        if previous_dir.exists() and not backup_dir.exists():
            try:
                previous_dir.replace(backup_dir)
            except Exception:
                pass
        raise


def backup_active_chroma_store_best_effort(
    reason: str = "",
    *,
    source_file_path: Optional[str] = None,
    _lock_held: bool = False,
) -> Dict[str, Any]:
    if not should_auto_backup_chroma_store():
        return {"backup_created": False, "reason": "auto_backup_not_in_runtime"}
    if source_file_path and not is_path_within_directory(Path(source_file_path), docs_dir()):
        return {"backup_created": False, "reason": "source_file_outside_docs_dir"}
    try:
        summary = backup_active_chroma_store(_lock_held=_lock_held)
        if summary.get("backup_created"):
            print(
                "[KB] Chroma backup updated"
                + (f" after {reason}" if reason else "")
                + f": {summary.get('backup_dir')}"
            )
        return summary
    except Exception as exc:
        print(f"[KB] Chroma backup skipped after {reason or 'index update'}: {exc}")
        return {"backup_created": False, "reason": str(exc)}


def restore_chroma_store_from_backup(*, _lock_held: bool = False) -> Dict[str, Any]:
    if not _lock_held:
        with kb_index_write_lock():
            return restore_chroma_store_from_backup(_lock_held=True)

    persist_dir = Path(RAG_LOCAL_PERSIST_DIR)
    backup_dir = chroma_backup_dir()
    if not backup_dir.exists():
        return {
            "restored": False,
            "reason": "backup_missing",
            "persist_dir": str(persist_dir),
            "backup_dir": str(backup_dir),
        }
    if not any(backup_dir.iterdir()):
        return {
            "restored": False,
            "reason": "backup_empty",
            "persist_dir": str(persist_dir),
            "backup_dir": str(backup_dir),
        }

    metadata = read_chroma_backup_metadata(backup_dir)
    backup_snapshot = metadata.get("manifest_snapshot") if isinstance(metadata, dict) else None
    current_snapshot = build_chroma_manifest_snapshot()
    current_signature = chroma_manifest_signature(current_snapshot)
    backup_signature = str(metadata.get("manifest_signature") or "")
    manifest_matches = bool(backup_signature and backup_signature == current_signature)
    manifest_diff = diff_chroma_manifest_snapshots(backup_snapshot, current_snapshot)

    release_chroma_runtime_handles()
    persist_dir.parent.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    replaced_dir: Optional[Path] = None
    tmp_restore = persist_dir.with_name(f"{persist_dir.name}.restore-tmp-{stamp}-{uuid.uuid4().hex[:8]}")
    try:
        shutil.copytree(backup_dir, tmp_restore, ignore=shutil.ignore_patterns("_backup_metadata.json"))
        if persist_dir.exists():
            replaced_dir = persist_dir.with_name(
                f"{persist_dir.name}.bad-before-restore-{stamp}-{uuid.uuid4().hex[:8]}"
            )
            persist_dir.replace(replaced_dir)
        tmp_restore.replace(persist_dir)
        release_chroma_runtime_handles()
        return {
            "restored": True,
            "persist_dir": str(persist_dir),
            "backup_dir": str(backup_dir),
            "replaced_persist_store": str(replaced_dir) if replaced_dir else None,
            "manifest_matches": manifest_matches,
            "backup_manifest_signature": backup_signature,
            "current_manifest_signature": current_signature,
            "backup_manifest_snapshot": backup_snapshot,
            **manifest_diff,
        }
    except Exception:
        if tmp_restore.exists():
            shutil.rmtree(tmp_restore, ignore_errors=True)
        if replaced_dir and replaced_dir.exists() and not persist_dir.exists():
            try:
                replaced_dir.replace(persist_dir)
            except Exception:
                pass
        release_chroma_runtime_handles()
        raise


def refresh_restored_chroma_store_from_manifest_delta(
    restore_summary: Dict[str, Any],
    *,
    _lock_held: bool = False,
) -> Dict[str, Any]:
    if not _lock_held:
        with kb_index_write_lock():
            return refresh_restored_chroma_store_from_manifest_delta(
                restore_summary,
                _lock_held=True,
            )

    changed_document_ids = [
        str(item)
        for item in restore_summary.get("changed_document_ids") or []
        if str(item)
    ]
    deleted_document_ids = [
        str(item)
        for item in restore_summary.get("deleted_document_ids") or []
        if str(item)
    ]
    if not changed_document_ids and not deleted_document_ids:
        return {
            "changed_document_ids": [],
            "deleted_document_ids": [],
            "indexed_count": 0,
            "failed_count": 0,
            "deleted_count": 0,
            "backup": backup_active_chroma_store_best_effort(
                "backup restore verification",
                _lock_held=True,
            ),
        }

    records_by_id = {
        str(record.get("id") or ""): record
        for record in list_documents(status="active")
        if isinstance(record, dict)
    }
    searcher = load_searcher()
    documents = []
    try:
        if deleted_document_ids:
            searcher.create_or_load_collection()
            for document_id in deleted_document_ids:
                searcher.collection.delete(where={"document_id": document_id})

        for document_id in changed_document_ids:
            record = records_by_id.get(document_id)
            if not record or record.get("processing_status") != "indexed":
                continue
            documents.append(index_document(
                dict(record),
                indexed_by="system",
                _lock_held=True,
                _searcher=searcher,
            ))
    finally:
        _close_searcher(searcher)
        reset_runtime_search_cache(clear_degraded=False)

    failed_count = sum(1 for item in documents if item.get("processing_status") == "failed")
    failures = [
        _index_failure_detail(item)
        for item in documents
        if item.get("processing_status") == "failed"
    ]
    if failed_count:
        _report_reindex_failures(failures)
        raise RuntimeError(
            f"備份還原後有 {failed_count} 份異動文件補索引失敗；"
            "將改走完整重建。"
        )

    backup_summary = backup_active_chroma_store_best_effort(
        "backup restore delta refresh",
        _lock_held=True,
    )
    return {
        "changed_document_ids": changed_document_ids,
        "deleted_document_ids": deleted_document_ids,
        "indexed_count": len(documents),
        "failed_count": failed_count,
        "deleted_count": len(deleted_document_ids),
        "backup": backup_summary,
    }


def archive_corrupted_chroma_store() -> Optional[Path]:
    persist_dir = Path(RAG_LOCAL_PERSIST_DIR)
    persist_dir.parent.mkdir(parents=True, exist_ok=True)
    release_chroma_runtime_handles()

    backup_dir: Optional[Path] = None
    if persist_dir.exists():
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        backup_dir = persist_dir.with_name(
            f"{persist_dir.name}.corrupt-{stamp}-{uuid.uuid4().hex[:8]}"
        )
        last_error: Optional[Exception] = None
        for attempt in range(6):
            try:
                persist_dir.replace(backup_dir)
                last_error = None
                break
            except PermissionError as exc:
                last_error = exc
                release_chroma_runtime_handles()
                time.sleep(0.2 * (attempt + 1))
        if last_error is not None:
            raise PermissionError(
                "Chroma 索引目錄仍被程序占用，無法封存損壞索引。"
                "請確認沒有第二個 API 程序使用同一個 cust_app_runtime。"
            ) from last_error

    persist_dir.mkdir(parents=True, exist_ok=False)
    return backup_dir


def restore_archived_chroma_store(backup_dir: Optional[Path]) -> None:
    persist_dir = Path(RAG_LOCAL_PERSIST_DIR)
    release_chroma_runtime_handles()

    if persist_dir.exists():
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        failed_dir = persist_dir.with_name(
            f"{persist_dir.name}.failed-{stamp}-{uuid.uuid4().hex[:8]}"
        )
        persist_dir.replace(failed_dir)

    if backup_dir and backup_dir.exists():
        backup_dir.replace(persist_dir)
    else:
        persist_dir.mkdir(parents=True, exist_ok=True)
    release_chroma_runtime_handles()


def rebuild_active_index(*, _lock_held: bool = False) -> Dict[str, Any]:
    if not _lock_held:
        with kb_index_write_lock():
            return rebuild_active_index(_lock_held=True)

    manifest_path = Path(RAG_LOCAL_MANIFEST_PATH)
    manifest_snapshot = manifest_path.read_bytes() if manifest_path.exists() else None
    backup_dir: Optional[Path] = None
    try:
        backup_dir = archive_corrupted_chroma_store()
    except PermissionError as exc:
        reset_runtime_search_cache(clear_degraded=False)
        raise RuntimeError(
            "無法封存已損壞的 Chroma 索引，已停止完整重建，"
            "避免在損壞索引上覆寫。請確認只有一個 API 程序使用同一個 "
            "cust_app_runtime，停止占用後再重試。"
        ) from exc
    try:
        summary = reindex_all_documents(
            indexed_by="system",
            reset_collection=False,
            _lock_held=True,
        )
        summary["repair_mode"] = "replace_persist_store"
        summary["replaced_persist_store"] = True
        if int(summary.get("failed_count") or 0):
            failures = list(summary.get("failures") or [])
            _report_reindex_failures(failures)
            raise RuntimeError(
                f"全新索引仍有 {summary.get('failed_count')} 份文件建立失敗；"
                "請查看上方逐筆 [KB] Reindex document failed 訊息，"
                "或每日 app_errors 記錄。"
            )
        summary["corrupt_store_backup"] = str(backup_dir) if backup_dir else None
        summary["healthy_store_backup"] = backup_active_chroma_store_best_effort(
            "full rebuild",
            _lock_held=True,
        )
        reset_runtime_search_cache(clear_degraded=True)
        return summary
    except Exception:
        restore_archived_chroma_store(backup_dir)
        if manifest_snapshot is not None:
            manifest_path.parent.mkdir(parents=True, exist_ok=True)
            manifest_path.write_bytes(manifest_snapshot)
        reset_runtime_search_cache(clear_degraded=False)
        raise


def delete_document_vectors_from_index(document_id: str, *, _lock_held: bool = False) -> None:
    if not _lock_held:
        with kb_index_write_lock():
            delete_document_vectors_from_index(document_id, _lock_held=True)
            return
    searcher = load_searcher()
    try:
        searcher.delete_by_document_id(str(document_id))
        collection = getattr(searcher, "collection", None)
        verify_get = getattr(collection, "get", None)
        if callable(verify_get):
            remaining = verify_get(where={"document_id": str(document_id)}, limit=1)
            remaining_ids = remaining.get("ids") if isinstance(remaining, dict) else []
            if remaining_ids:
                raise RuntimeError(f"Chroma 仍有文件 {document_id} 的殘留向量。")
    finally:
        _close_searcher(searcher)
        reset_runtime_search_cache(clear_degraded=False)


def reindex_all_documents(
    indexed_by: Optional[str] = None,
    *,
    reset_collection: bool = True,
    _lock_held: bool = False,
) -> Dict[str, Any]:
    if not _lock_held:
        with kb_index_write_lock():
            return reindex_all_documents(
                indexed_by=indexed_by,
                reset_collection=reset_collection,
                _lock_held=True,
            )

    searcher = load_searcher()
    try:
        if reset_collection:
            searcher.delete_collection()
            searcher.collection = None

        documents = []
        for active_record in list_documents(status="active"):
            if active_record.get("processing_status") == "deleted":
                continue
            documents.append(index_document(
                dict(active_record),
                indexed_by=indexed_by,
                _lock_held=True,
                _searcher=searcher,
            ))
    finally:
        _close_searcher(searcher)
        reset_runtime_search_cache(clear_degraded=False)
    indexed_count = sum(1 for item in documents if item.get("processing_status") == "indexed")
    failed_count = sum(1 for item in documents if item.get("processing_status") == "failed")
    failures = [
        _index_failure_detail(item)
        for item in documents
        if item.get("processing_status") == "failed"
    ]
    return {
        "total_count": len(documents),
        "indexed_count": indexed_count,
        "failed_count": failed_count,
        "chunk_count": sum(int(item.get("indexed_chunk_count") or 0) for item in documents),
        "failures": failures,
        "documents": documents,
    }


def index_document(
    record: Dict[str, Any],
    indexed_by: Optional[str] = None,
    *,
    _lock_held: bool = False,
    _searcher: Any = None,
) -> Dict[str, Any]:
    if not _lock_held:
        with kb_index_write_lock():
            return index_document(
                record,
                indexed_by=indexed_by,
                _lock_held=True,
                _searcher=_searcher,
            )

    actor = normalize_actor_name(indexed_by) if indexed_by else ""
    record["processing_status"] = "processing"
    record["processing_error"] = None
    record["processing_error_stage"] = None
    record["processing_error_type"] = None
    record["updated_at"] = now_iso()
    if actor:
        record["last_indexed_by"] = actor
    upsert_document_record(record)

    owns_searcher = _searcher is None
    searcher = _searcher
    failure_stage = "extract_document"
    try:
        file_path = resolve_document_file_path(record["file_path"])
        sections = extract_document_sections(file_path)
        failure_stage = "chunk_document"
        chunks = chunk_sections(sections)

        if not chunks:
            raise RuntimeError("文件沒有可建立索引的文字內容。")

        failure_stage = "build_metadata"
        campaign_profile = build_campaign_profile(record, chunks)
        product_service_profile = build_value_added_product_profile(record, chunks)
        if campaign_profile:
            record["document_type"] = campaign_profile["document_type"]
            record["campaign_profile"] = campaign_profile
            record["campaign_validation_status"] = campaign_profile["validation"]["status"]
            record.pop("product_service_profile", None)
        elif product_service_profile:
            record["document_type"] = product_service_profile["document_type"]
            record["product_service_profile"] = product_service_profile
            record.pop("campaign_profile", None)
            record.pop("campaign_validation_status", None)
        else:
            record["document_type"] = "general"
            record.pop("campaign_profile", None)
            record.pop("campaign_validation_status", None)
            record.pop("product_service_profile", None)

        inferred_category = record.get("category") or infer_category_from_text(
            " ".join([
                str(record.get("title") or ""),
                str(record.get("file_name") or ""),
                " ".join(chunk.get("content", "") for chunk in chunks[:5]),
            ])
        )
        record["category"] = inferred_category

        failure_stage = "build_index_records"
        records = []
        for index, chunk in enumerate(chunks, start=1):
            chunk_id = f"{record['id']}:{index}"
            chunk_question = chunk.get("question") or record["title"]
            chunk_answer = chunk.get("answer") or chunk["content"]
            chunk_company = chunk.get("company") or record["knowledge_base"]
            variant = next(
                (
                    item for item in campaign_profile.get("variants", [])
                    if int(item.get("chunk_index") or 0) == index
                ),
                {},
            )
            records.append({
                "id": chunk_id,
                "document_id": record["id"],
                "chunk_index": index,
                "question": chunk_question,
                "answer": chunk_answer,
                "content": chunk["content"],
                "company": chunk_company,
                "knowledge_base": record["knowledge_base"],
                "category": record.get("category") or "",
                "title": record["title"],
                "source": record["file_name"],
                "page_no": chunk.get("page_no") or "",
                "section": chunk.get("section") or "",
                "record_type": "chunk",
                **campaign_metadata(campaign_profile, variant),
                **value_added_product_metadata(product_service_profile),
            })

        qa_records = build_qa_index_records(record, chunks, inferred_category)
        for qa_record in qa_records:
            chunk_index = int(qa_record.get("chunk_index") or 0)
            variant = next(
                (
                    item for item in campaign_profile.get("variants", [])
                    if int(item.get("chunk_index") or 0) == chunk_index
                ),
                {},
            )
            qa_record.update(campaign_metadata(campaign_profile, variant))
            qa_record.update(value_added_product_metadata(product_service_profile))
        records.extend(qa_records)

        campaign_records = build_campaign_index_records(
            record,
            chunks,
            inferred_category,
            campaign_profile,
        )
        records.extend(campaign_records)

        product_service_records = build_value_added_product_index_records(
            record,
            product_service_profile,
        )
        records.extend(product_service_records)

        failure_stage = "load_vector_store"
        if searcher is None:
            searcher = load_searcher()
        failure_stage = "write_vectors"
        indexed = upsert_records_with_index_repair(
            searcher,
            records,
            record["id"],
            full_rebuild=not owns_searcher,
        )

        record["chunk_count"] = len(chunks)
        record["indexed_chunk_count"] = len(chunks)
        record["indexed_vector_count"] = indexed
        record["qa_chunk_count"] = len(qa_records)
        record["campaign_index_count"] = len(campaign_records)
        record["product_service_index_count"] = len(product_service_records)
        record["processing_status"] = "indexed"
        record["processing_error"] = None
        record["processing_error_stage"] = None
        record["processing_error_type"] = None
        record["processed_at"] = now_iso()
    except Exception as exc:
        record["chunk_count"] = 0
        record["indexed_chunk_count"] = 0
        record["indexed_vector_count"] = 0
        record["qa_chunk_count"] = 0
        record["campaign_index_count"] = 0
        record["product_service_index_count"] = 0
        record["processing_status"] = "failed"
        record["processing_error"] = str(exc)
        record["processing_error_stage"] = failure_stage
        record["processing_error_type"] = type(exc).__name__
    finally:
        if owns_searcher and searcher is not None:
            _close_searcher(searcher)
        if owns_searcher:
            reset_runtime_search_cache(clear_degraded=False)

    record["updated_at"] = now_iso()
    if actor:
        record["last_indexed_by"] = actor
        record["last_indexed_at"] = record["updated_at"]
    saved_record = upsert_document_record(record)
    if saved_record.get("processing_status") == "indexed" and owns_searcher:
        backup_active_chroma_store_best_effort(
            "document index update",
            source_file_path=str(saved_record.get("file_path") or ""),
            _lock_held=True,
        )
    return saved_record


def create_document(
    file_name: str,
    content: bytes,
    title: Optional[str] = None,
    knowledge_base: str = "通用",
    category: Optional[str] = None,
    status: str = "active",
    uploaded_by: Optional[str] = None,
) -> Dict[str, Any]:
    knowledge_base = (
        normalize_knowledge_base_name(knowledge_base)
        or CENTRAL_COMMON_KNOWLEDGE_BASE
    )
    safe_name = sanitize_filename(file_name)
    extension = Path(safe_name).suffix.lower().lstrip(".")
    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError(f"不支援的文件格式：{extension}")

    document_id = uuid.uuid4().hex
    actor = normalize_actor_name(uploaded_by)
    created_at = now_iso()
    target_dir = docs_dir()
    target_dir.mkdir(parents=True, exist_ok=True)
    cleanup_deleted_file_records_by_name(safe_name)
    file_path = target_dir / f"{document_id}_{safe_name}"
    file_path.write_bytes(content)

    record = {
        "id": document_id,
        "title": title.strip() if title and title.strip() else Path(safe_name).stem,
        "file_name": safe_name,
        "file_path": str(file_path),
        "file_type": extension,
        "file_size": len(content),
        "knowledge_base": knowledge_base,
        "category": category or "",
        "status": status or "active",
        "processing_status": "queued",
        "processing_error": None,
        "chunk_count": 0,
        "indexed_chunk_count": 0,
        "created_at": created_at,
        "updated_at": created_at,
        "uploaded_by": actor,
        "uploaded_at": created_at,
    }
    upsert_document_record(record)
    return index_document(record, indexed_by=actor)


def reindex_document(document_id: str, indexed_by: Optional[str] = None) -> Dict[str, Any]:
    record = get_document(document_id)
    if not record:
        raise ValueError("找不到指定文件。")
    if record.get("status") == "deleted":
        raise ValueError("文件已刪除。")
    return index_document(record, indexed_by=indexed_by)


def delete_document(document_id: str, deleted_by: Optional[str] = None) -> Dict[str, Any]:
    record = get_document(document_id)
    if not record:
        raise ValueError("找不到指定文件。")

    actor = normalize_actor_name(deleted_by)
    deleted_at = now_iso()
    file_path = resolve_document_file_path(record.get("file_path", ""))
    if file_path.exists():
        file_path.unlink()

    record["status"] = "deleted"
    record["processing_status"] = "deleted"
    record["deleted_by"] = actor
    record["deleted_at"] = deleted_at
    record["updated_at"] = deleted_at
    deleted_record = upsert_document_record(record)

    try:
        try:
            delete_document_vectors_from_index(document_id)
        except Exception as delete_exc:
            print(f"[KB] delete document vectors failed verification, rebuilding active index: {delete_exc}")
            reindex_all_documents(indexed_by=actor)
        backup_active_chroma_store_best_effort(
            "document delete",
            source_file_path=str(record.get("file_path") or ""),
        )
    except Exception as exc:
        reset_runtime_search_cache(clear_degraded=False)
        raise RuntimeError(
            "文件已標記為刪除，但知識庫索引刷新失敗；"
            "請重新執行「全部重建索引」以避免客服對話查到舊資料。"
        ) from exc

    return deleted_record


def get_document_file_path(document_id: str) -> Path:
    record = get_document(document_id)
    if not record:
        raise ValueError("找不到指定文件。")
    if record.get("status") == "deleted":
        raise ValueError("文件已刪除。")

    file_path = resolve_document_file_path(record.get("file_path", ""))
    if not file_path.exists():
        raise FileNotFoundError("原始檔案不存在。")

    resolved_file = file_path.resolve()
    resolved_docs_dir = docs_dir().resolve()
    if resolved_docs_dir not in resolved_file.parents and resolved_file != resolved_docs_dir:
        raise ValueError("原始檔案路徑不合法。")

    return resolved_file


def get_document_chunks(document_id: str) -> List[Dict[str, Any]]:
    record = get_document(document_id)
    if not record:
        raise ValueError("找不到指定文件。")

    sections = extract_document_sections(resolve_document_file_path(record["file_path"]))
    chunks = chunk_sections(sections)
    return [
        {
            "document_id": document_id,
            "chunk_index": index,
            "content": chunk["content"],
            "page_no": chunk.get("page_no"),
            "section": chunk.get("section"),
        }
        for index, chunk in enumerate(chunks, start=1)
    ]


def reset_local_kb_files() -> None:
    target_dir = docs_dir()
    if target_dir.exists():
        shutil.rmtree(target_dir)
    manifest_path().parent.mkdir(parents=True, exist_ok=True)
    save_manifest({"documents": []})
