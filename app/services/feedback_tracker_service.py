"""Persistent tracking state for feedback and regression-case review."""

from __future__ import annotations

import csv
import hashlib
import json
import sqlite3
import uuid
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
import requests

from app.config.settings import (
    BASE_DIR,
    CHAT_LATENCY_LOG_PATH,
    DB_PATH,
    FEEDBACK_DIR,
    FEEDBACK_TRACKER_DB_PATH,
    FEEDBACK_TRACKER_ONLINE_API_BASE_URL,
    FEEDBACK_TRACKER_ONLINE_RUNTIME_DIR,
    FEEDBACK_TRACKER_SYNC_TOKEN,
    RUNTIME_DIR,
)
from app.services.memory_service import list_chat_logs_since


DEFAULT_START_DATE = "2026-08-24"
DEFAULT_FEEDBACK_STATUS = "pending"
DEFAULT_CASE_STATUS = "ready_for_review"
DEFAULT_REVIEW_STATUS = "pending"
LEGACY_FEEDBACK_STATUS_MAP = {
    "passed": "processed",
    "closed": "processed",
    "in_progress": "processed",
    "pending_verification": "processed",
    "failed": "pending",
    "needs_rework": "pending",
    "ready_for_review": "pending",
}
LEGACY_REVIEW_STATUS_MAP = {
    "needs_adjustment": "failed",
}
REGRESSION_WORKBOOK = (
    BASE_DIR / "outputs" / "20260903_august_feedback_regression" / "八月客服回饋回歸測試題庫.xlsx"
)
REGRESSION_RESULTS_FILE = (
    BASE_DIR / "outputs" / "20260903_august_feedback_regression" / "八月回歸題庫_AI測試結果_修改後.json"
)
TRACKER_DB_FILE = Path(FEEDBACK_TRACKER_DB_PATH)
TRACKER_SYNC_INBOX = Path(RUNTIME_DIR) / "feedback_tracker_sync" / "inbox"
TRACKER_SYNC_HEADER = "X-Feedback-Tracker-Sync-Token"
SYNC_CURSOR_KEYS = {
    "feedback": "online_feedback_cursor_v1",
    "conversation": "online_conversation_cursor_v1",
    "latency": "online_latency_cursor_v1",
}


def now_iso() -> str:
    return datetime.now().isoformat(timespec="seconds")


def tracker_db_conn() -> sqlite3.Connection:
    TRACKER_DB_FILE.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(TRACKER_DB_FILE, check_same_thread=False)
    conn.row_factory = sqlite3.Row
    return conn


def normalize_feedback_status(value: Any) -> str:
    status = str(value or "").strip()
    return LEGACY_FEEDBACK_STATUS_MAP.get(status, status or DEFAULT_FEEDBACK_STATUS)


def normalize_review_status(value: Any) -> str:
    status = str(value or "").strip()
    return LEGACY_REVIEW_STATUS_MAP.get(status, status or DEFAULT_REVIEW_STATUS)


def init_tracker_schema() -> None:
    conn = tracker_db_conn()
    try:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS feedback_records (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                feedback_id TEXT UNIQUE NOT NULL,
                user_id TEXT NOT NULL,
                company_code TEXT,
                company TEXT,
                tv_cable TEXT,
                feedback_type TEXT,
                suggestion TEXT NOT NULL,
                user_message TEXT,
                ai_response TEXT,
                conversation_json TEXT,
                memory_json TEXT,
                csv_path TEXT,
                created_at TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS feedback_tracker_state (
                feedback_id TEXT PRIMARY KEY,
                status TEXT NOT NULL DEFAULT 'pending',
                priority TEXT,
                owner TEXT,
                review_status TEXT NOT NULL DEFAULT 'pending',
                review_note TEXT,
                acceptance_issue TEXT,
                acceptance_feedback TEXT,
                adjusted_conversation_json TEXT,
                updated_at TEXT NOT NULL,
                updated_by TEXT
            );

            CREATE TABLE IF NOT EXISTS regression_cases (
                case_id TEXT PRIMARY KEY,
                suite TEXT,
                priority TEXT,
                group_name TEXT,
                company_codes TEXT,
                case_kind TEXT,
                script TEXT,
                expected_behavior TEXT,
                must_include TEXT,
                must_not_include TEXT,
                dependency TEXT,
                focus TEXT,
                source_feedback_ids_json TEXT NOT NULL,
                test_results_json TEXT NOT NULL DEFAULT '[]',
                status TEXT NOT NULL DEFAULT 'ready_for_review',
                owner TEXT,
                review_status TEXT NOT NULL DEFAULT 'pending',
                review_note TEXT,
                updated_at TEXT NOT NULL,
                updated_by TEXT
            );

            CREATE TABLE IF NOT EXISTS tracker_sync_meta (
                meta_key TEXT PRIMARY KEY,
                meta_value TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS tracker_conversation_logs (
                log_id TEXT PRIMARY KEY,
                user_id TEXT NOT NULL,
                customer_key TEXT,
                role TEXT NOT NULL,
                message TEXT NOT NULL,
                created_at TEXT NOT NULL,
                imported_at TEXT NOT NULL
            );

            CREATE INDEX IF NOT EXISTS idx_tracker_conversation_logs_created_at
            ON tracker_conversation_logs(created_at);

            CREATE TABLE IF NOT EXISTS tracker_latency_logs (
                log_id TEXT PRIMARY KEY,
                timestamp TEXT NOT NULL,
                request_id TEXT,
                user_id TEXT,
                channel TEXT,
                company_code TEXT,
                source_file TEXT NOT NULL,
                payload_json TEXT NOT NULL,
                imported_at TEXT NOT NULL
            );

            CREATE INDEX IF NOT EXISTS idx_tracker_latency_logs_timestamp
            ON tracker_latency_logs(timestamp);
            """
        )
        columns = {row[1] for row in conn.execute("PRAGMA table_info(regression_cases)")}
        if "test_results_json" not in columns:
            conn.execute(
                "ALTER TABLE regression_cases ADD COLUMN test_results_json TEXT NOT NULL DEFAULT '[]'"
            )
        tracker_columns = {row[1] for row in conn.execute("PRAGMA table_info(feedback_tracker_state)")}
        if "adjusted_conversation_json" not in tracker_columns:
            conn.execute(
                "ALTER TABLE feedback_tracker_state ADD COLUMN adjusted_conversation_json TEXT"
            )
        if "acceptance_issue" not in tracker_columns:
            conn.execute("ALTER TABLE feedback_tracker_state ADD COLUMN acceptance_issue TEXT")
        if "acceptance_feedback" not in tracker_columns:
            conn.execute("ALTER TABLE feedback_tracker_state ADD COLUMN acceptance_feedback TEXT")
        # Early batches used pass/fail for AI evaluation. Those values belong to
        # processing progress, while review_status remains reserved for CSR review.
        migration_time = now_iso()
        conn.execute(
            """
            UPDATE feedback_tracker_state
            SET status = CASE status
                WHEN 'passed' THEN 'processed'
                WHEN 'closed' THEN 'processed'
                WHEN 'in_progress' THEN 'processed'
                WHEN 'pending_verification' THEN 'processed'
                WHEN 'failed' THEN 'pending'
                WHEN 'needs_rework' THEN 'pending'
                WHEN 'ready_for_review' THEN 'pending'
                ELSE status
            END,
            updated_at = ?,
            updated_by = COALESCE(NULLIF(updated_by, ''), 'tracker-status-migration')
            WHERE status IN (
                'passed', 'closed', 'in_progress', 'pending_verification',
                'failed', 'needs_rework', 'ready_for_review'
            )
            """,
            (migration_time,),
        )
        conn.execute(
            """
            UPDATE feedback_tracker_state
            SET review_status='failed',
                updated_at=?,
                updated_by=COALESCE(NULLIF(updated_by, ''), 'tracker-status-migration')
            WHERE review_status='needs_adjustment'
            """,
            (migration_time,),
        )
        conn.commit()
    finally:
        conn.close()
    migrate_legacy_tracker_data()


def table_exists(conn: sqlite3.Connection, table_name: str) -> bool:
    return bool(
        conn.execute(
            "SELECT 1 FROM sqlite_master WHERE type='table' AND name=?", (table_name,)
        ).fetchone()
    )


def migrate_legacy_tracker_data() -> None:
    """Move the first local tracker dataset out of the deployable application folder."""
    if not Path(DB_PATH).exists() or Path(DB_PATH).resolve() == TRACKER_DB_FILE.resolve():
        return
    tracker_conn = tracker_db_conn()
    try:
        migrated = tracker_conn.execute(
            "SELECT 1 FROM tracker_sync_meta WHERE meta_key='legacy_migration_v1'"
        ).fetchone()
        if migrated:
            return

        legacy_conn = sqlite3.connect(DB_PATH)
        legacy_conn.row_factory = sqlite3.Row
        try:
            for table_name, columns in (
                (
                    "feedback_records",
                    "feedback_id, user_id, company_code, company, tv_cable, feedback_type, "
                    "suggestion, user_message, ai_response, conversation_json, memory_json, csv_path, created_at",
                ),
                (
                    "feedback_tracker_state",
                    "feedback_id, status, priority, owner, review_status, review_note, updated_at, updated_by",
                ),
                (
                    "regression_cases",
                    "case_id, suite, priority, group_name, company_codes, case_kind, script, "
                    "expected_behavior, must_include, must_not_include, dependency, focus, "
                    "source_feedback_ids_json, test_results_json, status, owner, review_status, "
                    "review_note, updated_at, updated_by",
                ),
            ):
                if not table_exists(legacy_conn, table_name):
                    continue
                source_columns = {
                    row[1] for row in legacy_conn.execute(f"PRAGMA table_info({table_name})")
                }
                selected_columns = [column.strip() for column in columns.split(",")]
                if not set(selected_columns).issubset(source_columns):
                    continue
                rows = legacy_conn.execute(f"SELECT {columns} FROM {table_name}").fetchall()
                if not rows:
                    continue
                placeholders = ", ".join("?" for _ in selected_columns)
                tracker_conn.executemany(
                    f"INSERT OR IGNORE INTO {table_name} ({columns}) VALUES ({placeholders})",
                    [tuple(row[column] for column in selected_columns) for row in rows],
                )
        finally:
            legacy_conn.close()
        tracker_conn.execute(
            "INSERT OR REPLACE INTO tracker_sync_meta (meta_key, meta_value) VALUES (?, ?)",
            ("legacy_migration_v1", now_iso()),
        )
        tracker_conn.commit()
    finally:
        tracker_conn.close()


def parse_date(value: Any) -> str:
    text = str(value or "").strip()
    return text[:10] if len(text) >= 10 else ""


def tracker_meta_value(meta_key: str) -> str | None:
    init_tracker_schema()
    conn = tracker_db_conn()
    try:
        row = conn.execute(
            "SELECT meta_value FROM tracker_sync_meta WHERE meta_key=?", (meta_key,)
        ).fetchone()
    finally:
        conn.close()
    return str(row["meta_value"]) if row else None


def set_tracker_meta_value(meta_key: str, meta_value: str) -> None:
    conn = tracker_db_conn()
    try:
        conn.execute(
            "INSERT OR REPLACE INTO tracker_sync_meta (meta_key, meta_value) VALUES (?, ?)",
            (meta_key, meta_value),
        )
        conn.commit()
    finally:
        conn.close()


def sync_overlap_cursor(cursor: str | None) -> str | None:
    """Replay a short window so same-time writes cannot be missed."""
    if not cursor:
        return None
    try:
        return (datetime.fromisoformat(cursor) - timedelta(minutes=5)).isoformat()
    except ValueError:
        return cursor


def load_feedback_csv_records(
    start_date: str = DEFAULT_START_DATE,
    include_feedback_ids: set[str] | None = None,
    source_dir: Path | None = None,
    after: str | None = None,
    limit: int | None = None,
) -> list[dict[str, str]]:
    records: list[dict[str, str]] = []
    for path in sorted((source_dir or Path(FEEDBACK_DIR)).glob("ai_feedback_*.csv")):
        try:
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                for row in csv.DictReader(handle):
                    feedback_id = str(row.get("feedback_id") or "").strip()
                    created_at = str(row.get("created_at") or "")
                    is_regression_source = feedback_id in (include_feedback_ids or set())
                    if (
                        (parse_date(created_at) >= start_date and (not after or created_at > after))
                        or is_regression_source
                    ):
                        row["csv_path"] = str(path)
                        records.append(row)
        except OSError:
            continue
    records = sorted(records, key=lambda row: str(row.get("created_at") or ""))
    return records[:max(1, limit)] if limit else records


def import_feedback_csvs(
    start_date: str = DEFAULT_START_DATE,
    include_feedback_ids: set[str] | None = None,
    source_dir: Path | None = None,
) -> int:
    return import_feedback_records(
        load_feedback_csv_records(start_date, include_feedback_ids, source_dir)
    )


def import_feedback_records(records: list[dict[str, Any]]) -> int:
    init_tracker_schema()
    if not records:
        return 0
    conn = tracker_db_conn()
    try:
        inserted = 0
        for row in records:
            cursor = conn.execute(
                """
                INSERT OR IGNORE INTO feedback_records (
                    feedback_id, user_id, company_code, company, tv_cable,
                    feedback_type, suggestion, user_message, ai_response,
                    conversation_json, memory_json, csv_path, created_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    row.get("feedback_id"),
                    row.get("user_id") or "legacy-feedback",
                    row.get("company_code"),
                    row.get("company"),
                    row.get("tv_cable"),
                    row.get("feedback_type"),
                    row.get("suggestion") or "未填寫",
                    row.get("user_message"),
                    row.get("ai_response"),
                    row.get("conversation_json"),
                    row.get("memory_json") or "{}",
                    row.get("csv_path"),
                    row.get("created_at") or now_iso(),
                ),
            )
            inserted += int(cursor.rowcount > 0)
        conn.commit()
        return inserted
    finally:
        conn.close()


def conversation_log_id(record: dict[str, Any]) -> str:
    """Stable ID avoids duplicate imports when the manager pulls more than once."""
    identity = "\x1f".join(
        str(record.get(key) or "")
        for key in ("id", "user_id", "customer_key", "role", "message", "created_at")
    )
    return hashlib.sha256(identity.encode("utf-8")).hexdigest()


def load_conversation_log_records(
    start_date: str = DEFAULT_START_DATE,
    limit: int = 20000,
    after: str | None = None,
) -> list[dict[str, Any]]:
    return list_chat_logs_since(start_date, limit, after)


def import_conversation_logs(records: list[dict[str, Any]]) -> int:
    init_tracker_schema()
    if not records:
        return 0
    conn = tracker_db_conn()
    try:
        imported = 0
        for record in records:
            message = str(record.get("message") or "").strip()
            user_id = str(record.get("user_id") or "").strip()
            created_at = str(record.get("created_at") or "").strip()
            if not (message and user_id and created_at):
                continue
            cursor = conn.execute(
                """
                INSERT OR IGNORE INTO tracker_conversation_logs (
                    log_id, user_id, customer_key, role, message, created_at, imported_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    conversation_log_id(record), user_id, record.get("customer_key"),
                    str(record.get("role") or "assistant"), message, created_at, now_iso(),
                ),
            )
            imported += int(cursor.rowcount > 0)
        conn.commit()
        return imported
    finally:
        conn.close()


def load_latency_log_records(
    start_date: str = DEFAULT_START_DATE,
    limit: int = 20000,
    after: str | None = None,
) -> list[dict[str, Any]]:
    base_path = Path(CHAT_LATENCY_LOG_PATH)
    paths = sorted(base_path.parent.glob(f"{base_path.stem}_*{base_path.suffix}*"))
    if base_path.exists():
        paths.append(base_path)

    records: list[dict[str, Any]] = []
    for path in paths:
        try:
            with path.open("r", encoding="utf-8") as handle:
                for line in handle:
                    try:
                        record = json.loads(line)
                    except json.JSONDecodeError:
                        continue
                    timestamp = str(record.get("timestamp") or "")
                    if (
                        record.get("event") != "chat_latency"
                        or timestamp[:10] < start_date
                        or (after and timestamp <= after)
                    ):
                        continue
                    records.append({**record, "_source_file": path.name})
                    if len(records) >= limit:
                        return sorted(records, key=lambda item: str(item.get("timestamp") or ""))
        except OSError:
            continue
    return sorted(records, key=lambda item: str(item.get("timestamp") or ""))


def latency_log_id(record: dict[str, Any]) -> str:
    identity = json.dumps(
        {key: value for key, value in record.items() if key != "_source_file"},
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )
    return hashlib.sha256(identity.encode("utf-8")).hexdigest()


def import_latency_logs(records: list[dict[str, Any]]) -> int:
    init_tracker_schema()
    if not records:
        return 0
    conn = tracker_db_conn()
    try:
        imported = 0
        for record in records:
            timestamp = str(record.get("timestamp") or "").strip()
            if not timestamp:
                continue
            cursor = conn.execute(
                """
                INSERT OR IGNORE INTO tracker_latency_logs (
                    log_id, timestamp, request_id, user_id, channel, company_code,
                    source_file, payload_json, imported_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    latency_log_id(record), timestamp, record.get("request_id"),
                    record.get("user_id"), record.get("channel"), record.get("company_code"),
                    str(record.get("_source_file") or "chat_latency.log"),
                    json.dumps(record, ensure_ascii=False, sort_keys=True), now_iso(),
                ),
            )
            imported += int(cursor.rowcount > 0)
        conn.commit()
        return imported
    finally:
        conn.close()


def load_regression_test_results(results_path: Path = REGRESSION_RESULTS_FILE) -> dict[str, list[dict[str, Any]]]:
    if not results_path.exists():
        return {}
    try:
        raw_results = json.loads(results_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}

    results_by_case: dict[str, list[dict[str, Any]]] = {}
    for raw in raw_results if isinstance(raw_results, list) else []:
        case_id = str(raw.get("caseId") or "").strip()
        if case_id:
            results_by_case.setdefault(case_id, []).append(raw)
    return results_by_case


def regression_source_feedback_ids(workbook_path: Path = REGRESSION_WORKBOOK) -> set[str]:
    if not workbook_path.exists():
        return set()
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook["回歸題庫"]
        values = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(values)]
        return {
            feedback_id.strip()
            for raw in values
            for feedback_id in str(
                raw[headers.index("來源回饋ID")] if headers.index("來源回饋ID") < len(raw) else ""
            ).splitlines()
            if feedback_id.strip()
        }
    finally:
        workbook.close()


def import_august_regression_cases(workbook_path: Path = REGRESSION_WORKBOOK) -> int:
    init_tracker_schema()
    if not workbook_path.exists():
        return 0

    results_by_case = load_regression_test_results()
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook["回歸題庫"]
        values = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(values)]
        rows = []
        for raw in values:
            row = {headers[index]: raw[index] if index < len(raw) else "" for index in range(len(headers))}
            case_id = str(row.get("題號") or "").strip()
            if not case_id:
                continue
            source_ids = [item.strip() for item in str(row.get("來源回饋ID") or "").splitlines() if item.strip()]
            rows.append((
                case_id,
                str(row.get("測試套件") or ""),
                str(row.get("優先級") or ""),
                str(row.get("功能群組") or ""),
                str(row.get("公司代碼") or ""),
                str(row.get("題型") or ""),
                str(row.get("對話腳本（依序輸入）") or ""),
                str(row.get("期望AI行為") or ""),
                str(row.get("必須包含") or ""),
                str(row.get("禁止出現/禁止行為") or ""),
                str(row.get("外部依賴") or ""),
                str(row.get("判定重點") or ""),
                json.dumps(source_ids, ensure_ascii=False),
                json.dumps(results_by_case.get(case_id, []), ensure_ascii=False),
            ))
    finally:
        workbook.close()

    conn = tracker_db_conn()
    try:
        for row in rows:
            conn.execute(
                """
                INSERT INTO regression_cases (
                    case_id, suite, priority, group_name, company_codes, case_kind,
                    script, expected_behavior, must_include, must_not_include,
                    dependency, focus, source_feedback_ids_json, test_results_json,
                    status, review_status, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(case_id) DO UPDATE SET
                    suite=excluded.suite,
                    priority=excluded.priority,
                    group_name=excluded.group_name,
                    company_codes=excluded.company_codes,
                    case_kind=excluded.case_kind,
                    script=excluded.script,
                    expected_behavior=excluded.expected_behavior,
                    must_include=excluded.must_include,
                    must_not_include=excluded.must_not_include,
                    dependency=excluded.dependency,
                    focus=excluded.focus,
                    source_feedback_ids_json=excluded.source_feedback_ids_json,
                    test_results_json=excluded.test_results_json
                """,
                (*row, DEFAULT_CASE_STATUS, DEFAULT_REVIEW_STATUS, now_iso()),
            )
        conn.commit()
        return len(rows)
    finally:
        conn.close()


def online_runtime_dir() -> Path:
    configured = str(FEEDBACK_TRACKER_ONLINE_RUNTIME_DIR or "").strip()
    if not configured:
        raise ValueError("尚未設定線上 runtime 資料夾。")
    return Path(configured)


def apply_tracking_update_bundle(payload: dict[str, Any]) -> int:
    """Apply case definitions and only newer manual tracking states."""
    if payload.get("kind") != "feedback_tracker_updates":
        raise ValueError("不是可辨識的追蹤同步檔。")

    init_tracker_schema()
    applied = 0
    conn = tracker_db_conn()
    try:
        for row in payload.get("case_definitions") or []:
            case_id = str(row.get("case_id") or "").strip()
            if not case_id:
                continue
            cursor = conn.execute(
                """
                INSERT INTO regression_cases (
                    case_id, suite, priority, group_name, company_codes, case_kind,
                    script, expected_behavior, must_include, must_not_include,
                    dependency, focus, source_feedback_ids_json, test_results_json,
                    status, review_status, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(case_id) DO UPDATE SET
                    suite=excluded.suite,
                    priority=excluded.priority,
                    group_name=excluded.group_name,
                    company_codes=excluded.company_codes,
                    case_kind=excluded.case_kind,
                    script=excluded.script,
                    expected_behavior=excluded.expected_behavior,
                    must_include=excluded.must_include,
                    must_not_include=excluded.must_not_include,
                    dependency=excluded.dependency,
                    focus=excluded.focus,
                    source_feedback_ids_json=excluded.source_feedback_ids_json,
                    test_results_json=excluded.test_results_json
                """,
                (
                    case_id, row.get("suite"), row.get("priority"), row.get("group_name"),
                    row.get("company_codes"), row.get("case_kind"), row.get("script"),
                    row.get("expected_behavior"), row.get("must_include"), row.get("must_not_include"),
                    row.get("dependency"), row.get("focus"),
                    row.get("source_feedback_ids_json") or "[]",
                    row.get("test_results_json") or "[]",
                    DEFAULT_CASE_STATUS, DEFAULT_REVIEW_STATUS, now_iso(),
                ),
            )
            applied += int(cursor.rowcount > 0)

        for row in payload.get("feedback_states") or []:
            if not row.get("feedback_id") or not row.get("updated_at"):
                continue
            cursor = conn.execute(
                """
                INSERT INTO feedback_tracker_state (
                    feedback_id, status, priority, owner, review_status, review_note,
                    acceptance_issue, acceptance_feedback, adjusted_conversation_json,
                    updated_at, updated_by
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(feedback_id) DO UPDATE SET
                    status=excluded.status, priority=excluded.priority, owner=excluded.owner,
                    review_status=excluded.review_status, review_note=excluded.review_note,
                    acceptance_issue=COALESCE(
                        excluded.acceptance_issue,
                        feedback_tracker_state.acceptance_issue
                    ),
                    acceptance_feedback=COALESCE(
                        excluded.acceptance_feedback,
                        feedback_tracker_state.acceptance_feedback
                    ),
                    adjusted_conversation_json=COALESCE(
                        excluded.adjusted_conversation_json,
                        feedback_tracker_state.adjusted_conversation_json
                    ),
                    updated_at=excluded.updated_at, updated_by=excluded.updated_by
                WHERE excluded.updated_at > feedback_tracker_state.updated_at
                """,
                (
                    row.get("feedback_id"), normalize_feedback_status(row.get("status")),
                    row.get("priority"), row.get("owner"),
                    normalize_review_status(row.get("review_status")), row.get("review_note"),
                    row.get("acceptance_issue"), row.get("acceptance_feedback"),
                    row.get("adjusted_conversation_json"),
                    row.get("updated_at"), row.get("updated_by"),
                ),
            )
            applied += int(cursor.rowcount > 0)

        for row in payload.get("case_states") or []:
            if not row.get("case_id") or not row.get("updated_at"):
                continue
            cursor = conn.execute(
                """
                UPDATE regression_cases
                SET status=?, owner=?, review_status=?, review_note=?, updated_at=?, updated_by=?
                WHERE case_id=? AND updated_at < ?
                """,
                (
                    row.get("status") or DEFAULT_CASE_STATUS, row.get("owner"),
                    row.get("review_status") or DEFAULT_REVIEW_STATUS, row.get("review_note"),
                    row.get("updated_at"), row.get("updated_by"), row.get("case_id"), row.get("updated_at"),
                ),
            )
            applied += int(cursor.rowcount > 0)
        conn.commit()
    finally:
        conn.close()
    return applied


def apply_pending_tracking_updates() -> int:
    if not TRACKER_SYNC_INBOX.exists():
        return 0
    applied = 0
    processed_dir = TRACKER_SYNC_INBOX / "processed"
    rejected_dir = TRACKER_SYNC_INBOX / "rejected"
    for path in sorted(TRACKER_SYNC_INBOX.glob("tracker_updates_*.json")):
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
            applied += apply_tracking_update_bundle(payload)
            processed_dir.mkdir(parents=True, exist_ok=True)
            path.replace(processed_dir / path.name)
        except (OSError, ValueError, json.JSONDecodeError):
            rejected_dir.mkdir(parents=True, exist_ok=True)
            path.replace(rejected_dir / path.name)
    return applied


def pull_online_feedback(start_date: str = DEFAULT_START_DATE) -> dict[str, Any]:
    if FEEDBACK_TRACKER_ONLINE_API_BASE_URL:
        if not FEEDBACK_TRACKER_SYNC_TOKEN:
            raise ValueError("尚未設定追蹤同步權杖。")
        feedback_after = sync_overlap_cursor(tracker_meta_value(SYNC_CURSOR_KEYS["feedback"]))
        conversation_after = sync_overlap_cursor(tracker_meta_value(SYNC_CURSOR_KEYS["conversation"]))
        latency_after = sync_overlap_cursor(tracker_meta_value(SYNC_CURSOR_KEYS["latency"]))
        try:
            response = requests.get(
                f"{FEEDBACK_TRACKER_ONLINE_API_BASE_URL}/api/feedback-tracker/sync/remote/feedback",
                params={
                    "start_date": start_date,
                    "feedback_after": feedback_after,
                    "conversation_after": conversation_after,
                    "latency_after": latency_after,
                },
                headers={TRACKER_SYNC_HEADER: FEEDBACK_TRACKER_SYNC_TOKEN},
                timeout=30,
            )
            response.raise_for_status()
            payload = response.json()
            records = payload.get("feedback_records") or []
            conversation_logs = payload.get("conversation_logs") or []
            latency_logs = payload.get("latency_logs") or []
        except requests.RequestException as exc:
            raise ValueError(f"線上回饋拉取失敗：{exc}") from exc
        imported = import_feedback_records(records)
        logs_imported = import_conversation_logs(conversation_logs)
        latency_logs_imported = import_latency_logs(latency_logs)
        tracking_updates = payload.get("tracking_updates")
        tracking_updates_applied = (
            apply_tracking_update_bundle(tracking_updates)
            if isinstance(tracking_updates, dict)
            else 0
        )
        for cursor_key, cursor in (
            (SYNC_CURSOR_KEYS["feedback"], payload.get("feedback_cursor")),
            (SYNC_CURSOR_KEYS["conversation"], payload.get("conversation_cursor")),
            (SYNC_CURSOR_KEYS["latency"], payload.get("latency_cursor")),
        ):
            if cursor:
                set_tracker_meta_value(cursor_key, str(cursor))
        return {
            "status": "success",
            "imported": imported,
            "conversation_logs_imported": logs_imported,
            "conversation_logs_total": payload.get("conversation_logs_total", len(conversation_logs)),
            "conversation_logs_truncated": bool(payload.get("conversation_logs_truncated")),
            "feedback_truncated": bool(payload.get("feedback_truncated")),
            "latency_logs_imported": latency_logs_imported,
            "latency_logs_total": payload.get("latency_logs_total", len(latency_logs)),
            "latency_logs_truncated": bool(payload.get("latency_logs_truncated")),
            "tracking_updates_applied": tracking_updates_applied,
            "tracking_updates_available": isinstance(tracking_updates, dict),
            "incremental": bool(feedback_after or conversation_after or latency_after),
            "source": "online_api",
        }

    remote_runtime = online_runtime_dir()
    source_dir = remote_runtime / "feedback"
    if not source_dir.exists():
        raise ValueError(f"找不到線上回饋資料夾：{source_dir}")
    imported = import_feedback_csvs(
        start_date,
        include_feedback_ids=regression_source_feedback_ids(),
        source_dir=source_dir,
    )
    return {"status": "success", "imported": imported, "source_dir": str(source_dir)}


def build_tracking_update_bundle() -> dict[str, Any]:
    init_tracker_schema()
    conn = tracker_db_conn()
    try:
        feedback_states = [
            dict(row) for row in conn.execute(
                "SELECT feedback_id, status, priority, owner, review_status, review_note, "
                "acceptance_issue, acceptance_feedback, "
                "adjusted_conversation_json, updated_at, updated_by "
                "FROM feedback_tracker_state"
            ).fetchall()
        ]
        case_states = [
            dict(row) for row in conn.execute(
                "SELECT case_id, status, owner, review_status, review_note, updated_at, updated_by "
                "FROM regression_cases WHERE updated_by IS NOT NULL"
            ).fetchall()
        ]
        case_definitions = [
            dict(row) for row in conn.execute(
                """
                SELECT case_id, suite, priority, group_name, company_codes, case_kind,
                       script, expected_behavior, must_include, must_not_include,
                       dependency, focus, source_feedback_ids_json, test_results_json
                FROM regression_cases
                ORDER BY case_id
                """
            ).fetchall()
        ]
    finally:
        conn.close()
    return {
        "kind": "feedback_tracker_updates",
        "created_at": now_iso(),
        "source_runtime": str(RUNTIME_DIR),
        "feedback_states": feedback_states,
        "case_states": case_states,
        "case_definitions": case_definitions,
    }


def build_tracking_state_bundle() -> dict[str, Any]:
    """Return mutable tracker states without republishing case definitions."""
    payload = build_tracking_update_bundle()
    payload["case_definitions"] = []
    return payload


def publish_tracking_updates() -> dict[str, Any]:
    payload = build_tracking_update_bundle()
    if FEEDBACK_TRACKER_ONLINE_API_BASE_URL:
        if not FEEDBACK_TRACKER_SYNC_TOKEN:
            raise ValueError("尚未設定追蹤同步權杖。")
        try:
            response = requests.post(
                f"{FEEDBACK_TRACKER_ONLINE_API_BASE_URL}/api/feedback-tracker/sync/remote/states",
                json=payload,
                headers={TRACKER_SYNC_HEADER: FEEDBACK_TRACKER_SYNC_TOKEN},
                timeout=30,
            )
            response.raise_for_status()
            result = response.json()
        except requests.RequestException as exc:
            raise ValueError(f"線上追蹤頁更新失敗：{exc}") from exc
        return {
            "status": "success",
            "feedback_states": len(payload["feedback_states"]),
            "case_states": len(payload["case_states"]),
            "case_definitions": len(payload["case_definitions"]),
            "applied": result.get("applied", 0),
            "target": "online_api",
        }

    remote_runtime = online_runtime_dir()
    target_inbox = remote_runtime / "feedback_tracker_sync" / "inbox"
    target_inbox.mkdir(parents=True, exist_ok=True)
    filename = f"tracker_updates_{datetime.now().strftime('%Y%m%d%H%M%S')}_{uuid.uuid4().hex[:8]}.json"
    target_path = target_inbox / filename
    temporary_path = target_path.with_suffix(".tmp")
    temporary_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary_path.replace(target_path)
    return {
        "status": "success",
        "feedback_states": len(payload["feedback_states"]),
        "case_states": len(payload["case_states"]),
        "case_definitions": len(payload["case_definitions"]),
        "target_path": str(target_path),
    }


def bootstrap_tracker_data(start_date: str = DEFAULT_START_DATE) -> dict[str, int]:
    source_ids = regression_source_feedback_ids()
    return {
        "feedback_imported": import_feedback_csvs(start_date, include_feedback_ids=source_ids),
        "cases_imported": import_august_regression_cases(),
        "tracking_updates_applied": apply_pending_tracking_updates(),
    }


def decode_json(value: Any, fallback: Any) -> Any:
    try:
        return json.loads(value) if value else fallback
    except (TypeError, json.JSONDecodeError):
        return fallback


def collapse_recent_identical_feedback_rows(rows: list[sqlite3.Row]) -> list[sqlite3.Row]:
    """Hide accidental repeated submissions while retaining every raw record."""
    kept: list[sqlite3.Row] = []
    latest_by_content: dict[str, datetime] = {}
    for row in rows:
        item = dict(row)
        fingerprint_fields = (
            item.get("company_code"),
            item.get("feedback_type"),
            item.get("suggestion"),
            item.get("conversation_json"),
            item.get("status"),
            item.get("review_status"),
            item.get("adjusted_conversation_json"),
        )
        fingerprint = hashlib.sha256(
            "\x1f".join(str(value or "").strip() for value in fingerprint_fields).encode("utf-8")
        ).hexdigest()
        try:
            created_at = datetime.fromisoformat(str(item.get("created_at") or ""))
        except ValueError:
            kept.append(row)
            continue
        previous = latest_by_content.get(fingerprint)
        if previous and previous - created_at <= timedelta(minutes=10):
            continue
        latest_by_content[fingerprint] = created_at
        kept.append(row)
    return kept


def list_feedback_items(
    *,
    start_date: str = DEFAULT_START_DATE,
    end_date: str | None = None,
    status: str | None = None,
    review_status: str | None = None,
    company_code: str | None = None,
    feedback_type: str | None = None,
    include_merged: bool = True,
) -> list[dict[str, Any]]:
    init_tracker_schema()
    clauses = ["substr(f.created_at, 1, 10) >= ?"]
    params: list[Any] = [start_date]
    if end_date:
        clauses.append("substr(f.created_at, 1, 10) <= ?")
        params.append(end_date)
    if status:
        clauses.append("COALESCE(t.status, ?) = ?")
        params.extend([DEFAULT_FEEDBACK_STATUS, status])
    if review_status:
        clauses.append("COALESCE(t.review_status, ?) = ?")
        params.extend([DEFAULT_REVIEW_STATUS, review_status])
    if company_code:
        clauses.append("f.company_code = ?")
        params.append(company_code)
    if feedback_type:
        clauses.append("f.feedback_type = ?")
        params.append(feedback_type)
    if not include_merged:
        clauses.append("COALESCE(t.review_status, ?) <> 'merged'")
        params.append(DEFAULT_REVIEW_STATUS)

    conn = tracker_db_conn()
    try:
        rows = conn.execute(
            f"""
            SELECT f.*, COALESCE(t.status, ?) AS status, t.priority, t.owner,
                   COALESCE(t.review_status, ?) AS review_status, t.review_note,
                   t.acceptance_issue, t.acceptance_feedback,
                   t.adjusted_conversation_json,
                   t.updated_at, t.updated_by
            FROM feedback_records f
            LEFT JOIN feedback_tracker_state t ON t.feedback_id = f.feedback_id
            WHERE {' AND '.join(clauses)}
            ORDER BY f.created_at DESC, f.feedback_id DESC
            """,
            [DEFAULT_FEEDBACK_STATUS, DEFAULT_REVIEW_STATUS, *params],
        ).fetchall()
    finally:
        conn.close()

    rows = collapse_recent_identical_feedback_rows(rows)
    return [
        {
            **dict(row),
            "conversation": decode_json(row["conversation_json"], []),
            "adjusted_conversation": decode_json(row["adjusted_conversation_json"], []),
        }
        for row in rows
    ]


def list_regression_cases(
    *,
    status: str | None = None,
    review_status: str | None = None,
    priority: str | None = None,
) -> list[dict[str, Any]]:
    init_tracker_schema()
    clauses = ["1 = 1"]
    params: list[Any] = []
    if status:
        clauses.append("status = ?")
        params.append(status)
    if review_status:
        clauses.append("review_status = ?")
        params.append(review_status)
    if priority:
        clauses.append("priority = ?")
        params.append(priority)
    conn = tracker_db_conn()
    try:
        rows = conn.execute(
            f"SELECT * FROM regression_cases WHERE {' AND '.join(clauses)} ORDER BY priority, case_id",
            params,
        ).fetchall()
        source_ids = {
            source_id
            for row in rows
            for source_id in decode_json(row["source_feedback_ids_json"], [])
        }
        sources_by_id: dict[str, dict[str, Any]] = {}
        if source_ids:
            placeholders = ", ".join("?" for _ in source_ids)
            source_rows = conn.execute(
                f"SELECT feedback_id, company_code, suggestion, conversation_json, created_at "
                f"FROM feedback_records WHERE feedback_id IN ({placeholders})",
                list(source_ids),
            ).fetchall()
            sources_by_id = {
                row["feedback_id"]: {
                    **dict(row),
                    "conversation": decode_json(row["conversation_json"], []),
                }
                for row in source_rows
            }
    finally:
        conn.close()
    return [
        {
            **dict(row),
            "source_feedback_ids": decode_json(row["source_feedback_ids_json"], []),
            "test_results": decode_json(row["test_results_json"], []),
            "source_feedback": [
                sources_by_id[source_id]
                for source_id in decode_json(row["source_feedback_ids_json"], [])
                if source_id in sources_by_id
            ],
        }
        for row in rows
    ]


def update_feedback_item(feedback_id: str, changes: dict[str, Any], actor: str) -> dict[str, Any]:
    init_tracker_schema()
    allowed = {
        "status", "priority", "owner", "review_status", "review_note",
        "acceptance_issue", "acceptance_feedback",
    }
    values = {key: str(value or "").strip() for key, value in changes.items() if key in allowed}
    adjusted_conversation_json = None
    if changes.get("adjusted_conversation") is not None:
        adjusted_conversation_json = json.dumps(
            changes.get("adjusted_conversation") or [],
            ensure_ascii=False,
        )
    conn = tracker_db_conn()
    try:
        conn.execute(
            """
            INSERT INTO feedback_tracker_state (
                feedback_id, status, priority, owner, review_status, review_note,
                acceptance_issue, acceptance_feedback, adjusted_conversation_json,
                updated_at, updated_by
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(feedback_id) DO UPDATE SET
                status=excluded.status, priority=excluded.priority, owner=excluded.owner,
                review_status=excluded.review_status, review_note=excluded.review_note,
                acceptance_issue=COALESCE(
                    excluded.acceptance_issue,
                    feedback_tracker_state.acceptance_issue
                ),
                acceptance_feedback=COALESCE(
                    excluded.acceptance_feedback,
                    feedback_tracker_state.acceptance_feedback
                ),
                adjusted_conversation_json=COALESCE(
                    excluded.adjusted_conversation_json,
                    feedback_tracker_state.adjusted_conversation_json
                ),
                updated_at=excluded.updated_at, updated_by=excluded.updated_by
            """,
            (
                feedback_id,
                normalize_feedback_status(values.get("status")),
                values.get("priority"),
                values.get("owner"),
                normalize_review_status(values.get("review_status")),
                values.get("review_note"),
                values.get("acceptance_issue"),
                values.get("acceptance_feedback"),
                adjusted_conversation_json,
                now_iso(),
                actor,
            ),
        )
        conn.commit()
    finally:
        conn.close()
    return {"status": "success", "feedback_id": feedback_id}


def update_regression_case(case_id: str, changes: dict[str, Any], actor: str) -> dict[str, Any]:
    allowed = {"status", "owner", "review_status", "review_note"}
    values = {key: str(value or "").strip() for key, value in changes.items() if key in allowed}
    conn = tracker_db_conn()
    try:
        conn.execute(
            """
            UPDATE regression_cases
            SET status=?, owner=?, review_status=?, review_note=?, updated_at=?, updated_by=?
            WHERE case_id=?
            """,
            (
                values.get("status") or DEFAULT_CASE_STATUS,
                values.get("owner"),
                values.get("review_status") or DEFAULT_REVIEW_STATUS,
                values.get("review_note"),
                now_iso(),
                actor,
                case_id,
            ),
        )
        conn.commit()
    finally:
        conn.close()
    return {"status": "success", "case_id": case_id}
